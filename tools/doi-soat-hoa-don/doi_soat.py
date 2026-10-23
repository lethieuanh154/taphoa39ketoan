#!/usr/bin/env python3
"""
SCRIPT 1 — DOI SOAT + QUY DOI DON VI hoa don dau vao / dau ra (XML hoa don dien tu)

Cach dung: dat file nay trong THU MUC GOC, ben trong co:
  - cac thu muc chua XML/ZIP (vd "Dau vao", "Dau ra" — ten gi cung duoc, script tu nhan chieu theo MST)
  - file bang_quy_doi_kiotviet_*.json  (TapHoa39BanHang > Quan ly hang hoa > nut "Xuat bang quy doi")
Bam dup file (hoac: python doi_soat.py). Ket qua nam ngay trong thu muc goc:
  doi_soat_<ngay_gio>.xlsx  - bao cao (danh sach san pham da ghep + quy doi ve DVT master, doi soat, ...)
  anh_xa_hang_hoa.xlsx      - ke toan sua Ma master / He so cho dong can xem, roi chay lai
  ton_dau_ky.xlsx           - (tuy chon) ton dau ky theo ma master
  BK_MUAVAO_BANRA_*, TK_GTGT_* - bang ke + to khai 01/GTGT (ban nhap)
  TK_TNCN_*                 - to khai 05/KK-TNCN (ban nhap) neu co file bang luong *luong*.xlsx o thu muc goc
  XML to khai (ban nhap)/   - XML 2 to khai theo dinh dang HTKK (script khong doc lai thu muc nay)
  Dau vao (quy doi)/, Dau ra (quy doi)/ - BAN SAO XML hoa don da quy doi DVT ve DVT master (tao lai moi lan chay)
File XML goc KHONG bi sua. Chi chay tren may, khong gui du lieu di dau.
Dung nhu module (KeToanBackEnd): read_xml -> reconcile -> report_sheets / tax_figures / tncn_figures — README muc 9b.
"""
import argparse
import bisect
import fnmatch
import glob
import io
import json
import math
import os
import re
import subprocess
import sys
import time
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict, namedtuple
from xml.sax.saxutils import escape as xml_escape
from datetime import date, datetime, timedelta

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("Lần đầu chạy: đang cài thư viện Excel (openpyxl)...", flush=True)
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--disable-pip-version-check", "-q", "openpyxl"])
    from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

VERSION = "2.0"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ======================= CẤU HÌNH — sửa ở đây nếu cần =======================
CONFIG = {
    "mst_cong_ty": "",                  # để trống = tự nhận (MST bán ra nhiều nhất)
    "tu_ngay": "",                      # "2026-07-01" — để trống = không lọc theo ngày
    "den_ngay": "",                     # "2026-09-30"
    "sai_so_tien": 0.5,                 # lệch tiền lớn hơn số này (đồng) thì báo; 0.5 = báo cả lệch 1đ
    "nguong_tin_cay_goi_y": 0.6,        # ghép theo tên đạt mức này thì tự điền Mã master
    "so_ngay_hd_vao_lap_sau": 0,        # cho phép HĐ đầu vào lập sau ngày bán bao nhiêu ngày
    "nguong_duoi_gia_von_phan_tram": 0, # giá bán thấp hơn giá vốn quá bao nhiêu % thì báo
    "phuong_phap_gia_von": "fifo",      # "fifo" = nhập trước xuất trước (mặc định CLI); "binh_quan" = bình quân gia quyền tính
                                        # lại sau MỖI lần nhập (KeToan web — kế toán: hết hàng mới nhập). Không ảnh hưởng số thuế
    "file_bang_quy_doi": ["bang_quy_doi_kiotviet*.json", "kiotvietproducts*.json"],  # lấy file mới nhất
    "file_anh_xa": "anh_xa_hang_hoa.xlsx",
    "file_ton_dau_ky": "ton_dau_ky.xlsx",
    # --- Bảng kê + tờ khai GTGT (01/GTGT TT80) + phụ lục giảm thuế NQ 204/2025/QH15
    "xuat_bang_ke_to_khai": True,
    "xuat_pdf_ban_nhap": True,
    "thue_con_khau_tru_ky_truoc_22": 0,  # chỉ tiêu [22] = [43] tờ khai kỳ trước (Q3/2026: 30542593)
    "ky_tinh_thue": "",                 # để trống = tự nhận quý có nhiều HĐ nhất, vd "Quý 3 năm 2026"
    "nguoi_ky": "",                     # họ tên người ký in cuối bảng kê / tờ khai
    # --- Tờ khai TNCN (05/KK-TNCN TT80) — từ file bảng lương Excel đặt ở thư mục gốc
    "xuat_to_khai_tncn": True,
    "file_bang_luong": ["*luong*.xlsx"],  # vd "LUONG SM 2026.xlsx": mỗi tháng 1 khối "Tháng MM/YYYY" + dòng tiêu đề
    "tncn_dem_so_nguoi": "ca_quy",      # [16]: "ca_quy" = số người có lương trong quý; "thang_cuoi" = chỉ tháng cuối quý
    "tncn_tru_bh_nld": True,            # [21] = Tổng lương thực tế + Thưởng − BH bắt buộc phần NLĐ (như tờ khai đã nộp)
    "tncn_khong_hdld": [],              # họ tên NV KHÔNG ký HĐLĐ từ 3 tháng trở lên → không tính vào [17]
    # --- XML tờ khai (định dạng HTKK) cho 01/GTGT và 05/KK-TNCN, ghi vào thư mục con XML_OUT_DIR
    "xuat_xml_to_khai": True,
    # --- Bản sao XML hóa đơn đã quy đổi ĐVT về ĐVT master → thư mục CLONE_DIRS (file gốc giữ nguyên)
    "xuat_xml_quy_doi": True,
    "xml_quy_doi_doi_ma": True,         # MHHDVu → Mã Vietda (nếu có), không thì Mã master (False = giữ mã trên HĐ)
    "xml_quy_doi_doi_ten": True,        # THHDVu → Tên Vietda / Tên master (False = giữ tên trên HĐ)
    "xml_quy_doi_chi_khi_chac_chan": True,  # chỉ sửa dòng có ánh xạ CHÍNH XÁC / kế toán đã xác nhận (cột 'Bản sao XML')
    "file_danh_muc_vietda": ["QUYDOI*.xlsx"],  # danh mục vật tư - hàng hóa xuất từ Vietda (Mã hàng kế toán tự đặt)
}
# ============================================================================

INVOICE_STATUS = {"": "Gốc", "1": "Thay thế", "2": "Điều chỉnh"}
XML_OUT_DIR = "XML to khai (ban nhap)"  # script ghi XML tờ khai vào đây và KHÔNG đọc lại thư mục này
CLONE_DIRS = {"VAO": "Dau vao (quy doi)", "RA": "Dau ra (quy doi)"}  # bản sao XML quy đổi — cũng KHÔNG đọc lại
SKIP_DIRS = {XML_OUT_DIR, *CLONE_DIRS.values()}


_LOG_HOOK = None       # None = in ra console (CLI); KeToanBackEnd gắn hàm riêng bằng set_hooks()
_PROGRESS_HOOK = None


def set_hooks(log_fn=None, progress_fn=None):
    """KeToanBackEnd: log() → log_fn(msg); thanh tiến trình → progress_fn(bước, đã xong, tổng). None = console như CLI."""
    global _LOG_HOOK, _PROGRESS_HOOK
    _LOG_HOOK, _PROGRESS_HOOK = log_fn, progress_fn


def log(msg=""):
    if _LOG_HOOK is not None:
        _LOG_HOOK(msg)
    else:
        print(msg, flush=True)


# ============================================================ CHUAN HOA

def strip_accents(s):
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return s.replace("đ", "d").replace("Đ", "D")


def norm_text(s):
    s = strip_accents(str(s or "")).lower()
    s = re.sub(r"(\d),(\d)", r"\1.\2", s)
    s = re.sub(r"[^a-z0-9.]+", " ", s)
    s = re.sub(r"(?<!\d)\.|\.(?!\d)", " ", s)
    s = re.sub(r"(\d+(?:\.\d+)?)\s*(ml|lit|l|kg|gr|g|cc)\b",
               lambda m: m.group(1) + {"lit": "l", "gr": "g", "cc": "ml"}.get(m.group(2), m.group(2)), s)
    return re.sub(r"\s+", " ", s).strip()


UNIT_SYNONYMS = {"carton": "thung", "ctn": "thung", "cartoon": "thung"}


def norm_unit(s):
    u = norm_text(s)
    return UNIT_SYNONYMS.get(u, u)


def norm_code(s):
    return re.sub(r"\s+", "", str(s or "")).upper()


COUNT_UNITS = r"lon|chai|goi|hop|cai|bich|tui|vi|cuon|qua|cay|vien|chiec|hu|lo|ong|tuyp|mieng|thanh|can|binh|bao|thoi"
PACK_WORDS = r"thung|loc|ket|kien|khay|hop|bao|goi|bich|tui|vi|day|block|set"
STRONG_PACK = {"thung", "loc", "ket", "kien", "khay", "block", "nua thung"}
UNIT_WORDS = set(PACK_WORDS.split("|")) | set(COUNT_UNITS.split("|")) | {"cai", "chiec", "con", "dvt", "nua"}

RE_PACK1 = re.compile(r"\b(?:" + PACK_WORDS + r")\s*(?:x\s*)?(\d{1,3})\s*(?:" + COUNT_UNITS + r")?\b")
RE_PACK2 = re.compile(r"(?<![\d.])(\d{1,3})\s*x\s*(?=\d)")
RE_PACK3 = re.compile(r"(?<![a-z\d])x\s*(\d{1,3})\b")
RE_PACK4 = re.compile(r"\b(\d{1,3})\s*(?:" + COUNT_UNITS + r")\b")


def base_name(norm_name):
    """Tên gốc: bỏ phần quy cách (thùng 48 hộp, x24, 6 chai) và chữ chỉ đơn vị."""
    cleaned = norm_name
    for rx in (RE_PACK1, RE_PACK2, RE_PACK3, RE_PACK4):
        cleaned = rx.sub(" ", cleaned)
    tokens = [t for t in cleaned.split() if t not in UNIT_WORDS]
    return " ".join(tokens).strip() or norm_name


def pack_count_from_name(norm_name, unit_n):
    """Số đơn vị lẻ đọc được từ tên hàng — CHỈ khi ĐVT là đơn vị đóng gói (thùng/lốc/két...)."""
    if unit_n not in STRONG_PACK:
        return None
    for rx in (RE_PACK1, RE_PACK2, RE_PACK3, RE_PACK4):
        m = rx.search(norm_name + " | " + unit_n)
        if m and int(m.group(1)) >= 2:
            return int(m.group(1))
    return None


def parse_num(s):
    if s is None:
        return 0.0
    if isinstance(s, (int, float)):
        return float(s)
    s = str(s).strip().replace(" ", "").replace(" ", "")
    if not s:
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        v = float(s)
        return -v if neg else v
    except ValueError:
        pass
    t = re.sub(r"[^\d,.\-]", "", s)
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", t) else t.replace(",", ".")
    elif t.count(".") > 1:
        t = t.replace(".", "")
    try:
        v = float(t)
        return -v if neg else v
    except ValueError:
        return 0.0


def parse_rate(s):
    """'10%' -> (0.10, '10%'); 'KCT' -> (0.0, 'KCT'); '' -> (None, '')."""
    raw = (s or "").strip().upper().replace(" ", "")
    if not raw:
        return None, ""
    if raw in ("KCT", "KKKNT"):
        return 0.0, raw
    m = re.search(r"-?\d+(?:[.,]\d+)?", raw)
    if not m:
        return None, raw
    v = float(m.group(0).replace(",", "."))
    if "%" in raw or v > 1:
        v = v / 100
    return v, f"{v * 100:g}%"


def parse_date(s):
    s = (s or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def round_half_up(x):
    return math.floor(abs(x) + 0.5) * (1 if x >= 0 else -1)


# ============================================================ DOC FILE XML

def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _child(el, name):
    if el is None:
        return None
    for c in el:
        if _local(c.tag) == name:
            return c
    return None


def _text(el, name):
    c = _child(el, name)
    return (c.text or "").strip() if c is not None and c.text else ""


def _desc(el, name):
    if el is None:
        return None
    for d in el.iter():
        if _local(d.tag) == name:
            return d
    return None


def iter_sources(root_dir):
    for dirpath, dirnames, files in os.walk(root_dir):
        dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS]
        for fn in sorted(files):
            path = os.path.join(dirpath, fn)
            low = fn.lower()
            rel = os.path.relpath(path, root_dir)
            try:
                if low.endswith(".xml"):
                    with open(path, "rb") as f:
                        yield rel, f.read()
                elif low.endswith(".zip"):
                    with open(path, "rb") as f:
                        yield from _iter_zip(rel, f.read(), 0)
            except OSError as e:
                log(f"  ! Không đọc được {rel}: {e}")


def _iter_zip(label, data, depth):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        log(f"  ! File ZIP hỏng: {label}")
        return
    for info in z.infolist():
        if info.is_dir():
            continue
        name = info.filename.lower()
        if name.endswith(".xml"):
            yield f"{label}!{info.filename}", z.read(info)
        elif name.endswith(".zip") and depth < 3:
            yield from _iter_zip(f"{label}!{info.filename}", z.read(info), depth + 1)


def parse_submitted_gtgt(data):
    """XML tờ khai 01/GTGT đã nộp (HTKK, maTKhai 842) → {mã chỉ tiêu: số} + tổng phụ lục 142; None nếu không phải."""
    root = ET.fromstring(data)
    if _local(root.tag) != "HSoThueDTu":
        return None
    ma = _desc(root, "maTKhai")
    if ma is None or (ma.text or "").strip() != "842":
        return {"_khac": True}
    out = {"_ky": (_desc(root, "kyKKhai").text or "").strip() if _desc(root, "kyKKhai") is not None else ""}
    ct = _desc(root, "CTieuTKhaiChinh")
    for e in ct.iter():
        tag = _local(e.tag)
        if tag.startswith("ct") and not len(list(e)) and e.text and re.fullmatch(r"-?\d+(\.\d+)?", e.text.strip()):
            out[tag[2:]] = float(e.text)
    for tag, key in (("tongCongGiaTriHHDVMuaVao", "pl1_gt"), ("tongCongThueGTGTHHDV", "pl1_thue"),
                     ("tongCongGiaTriHHDV", "pl2_gt"), ("tongCongThueGTGTDuocGiam", "pl2_giam"), ("ct9", "pl_09")):
        e = _desc(root, tag)
        if e is not None and e.text:
            out[key] = float(e.text)
    for tag, key in (("HH_DV_MuaVaoTrongKy", "pl1_dong"), ("HH_DV_BanRaTrongKy", "pl2_dong")):
        e = _desc(root, tag)
        if e is not None:
            out[key] = sum(1 for c in e if _local(c.tag) == "BangKeTenHHDV")
    return out


def parse_invoices(data, source):
    root = ET.fromstring(data)
    if _local(root.tag) == "HDon":
        blocks = [root]
    else:
        blocks = [e for e in root.iter() if _local(e.tag) == "HDon"]
    if not blocks:
        dl = _desc(root, "DLHDon")
        return [_parse_dl(dl, "", source)] if dl is not None else []
    out = []
    for h in blocks:
        dl = _child(h, "DLHDon")
        if dl is not None:
            out.append(_parse_dl(dl, _text(h, "MCCQT"), source))
    return out


def _parse_dl(dl, mccqt, source):
    tt = _child(dl, "TTChung")
    nd = _child(dl, "NDHDon")
    nban = _child(nd, "NBan")
    nmua = _child(nd, "NMua")
    toan = _child(nd, "TToan")
    lq = _child(tt, "TTHDLQuan")

    khms = _text(tt, "KHMSHDon")
    khhd = _text(tt, "KHHDon")
    try:
        so_int = int(float(_text(tt, "SHDon")))
    except ValueError:
        so_int = 0
    mst_ban = _text(nban, "MST")
    inv = {
        "file": source, "mccqt": mccqt, "khms": khms, "khhd": khhd, "ky_hieu": khms + khhd, "so": so_int,
        "ngay": parse_date(_text(tt, "NLap")),
        "dvtte": _text(tt, "DVTTe") or "VND",
        "tgia": parse_num(_text(tt, "TGia")) or 1.0,
        "mst_ban": mst_ban, "ten_ban": _text(nban, "Ten"), "dc_ban": _text(nban, "DChi"),
        "mst_mua": _text(nmua, "MST"), "ten_mua": _text(nmua, "Ten") or _text(nmua, "HVTNMHang"),
        "tc_hd": _text(lq, "TCHDon") if lq is not None else "", "lq_key": "",
        "lq_ngay": parse_date(_text(lq, "NLHDCLQuan")) if lq is not None else None,
        "eff_date": None, "in_period": True, "period_note": "",
        "tong_chua_thue": parse_num(_text(toan, "TgTCThue")),
        "tong_thue": parse_num(_text(toan, "TgTThue")),
        "tong_tt": parse_num(_text(toan, "TgTTTBSo")),
        "ck_tm": parse_num(_text(toan, "TTCKTMai")),
        "phi": 0.0, "lt_suat": [], "lines": [], "status": "", "direction": "", "issues": [], "rounding_lines": 0,
    }
    inv["key"] = f"{mst_ban}|{inv['ky_hieu']}|{so_int}"
    if lq is not None:
        lq_so = _text(lq, "SHDCLQuan")
        try:
            lq_so = int(float(lq_so))
        except ValueError:
            pass
        inv["lq_key"] = f"{mst_ban}|{_text(lq, 'KHMSHDCLQuan')}{_text(lq, 'KHHDCLQuan')}|{lq_so}"

    dsphi = _child(toan, "DSLPhi")
    if dsphi is not None:
        inv["phi"] = sum(parse_num(_text(p, "TPhi")) for p in dsphi)
    thtt = _child(toan, "THTTLTSuat")
    if thtt is not None:
        for lt in thtt:
            rate, label = parse_rate(_text(lt, "TSuat"))
            inv["lt_suat"].append((rate, label, parse_num(_text(lt, "ThTien")), parse_num(_text(lt, "TThue"))))

    ds = _child(nd, "DSHHDVu")
    header_rate = inv["lt_suat"][0][:2] if len(inv["lt_suat"]) == 1 else (None, "")
    if ds is not None:
        for h in ds:
            if _local(h.tag) != "HHDVu":
                continue
            rate, label = parse_rate(_text(h, "TSuat"))
            if rate is None and header_rate[0] is not None:
                rate, label = header_rate
            inv["lines"].append({
                "tchat": _text(h, "TChat") or "1", "stt": _text(h, "STT"), "ma": _text(h, "MHHDVu"),
                "ten": _text(h, "THHDVu"), "dvt": _text(h, "DVTinh"),
                "sl": parse_num(_text(h, "SLuong")), "dgia": parse_num(_text(h, "DGia")),
                "tlck": parse_num(_text(h, "TLCKhau")), "stck": parse_num(_text(h, "STCKhau")),
                "thtien": parse_num(_text(h, "ThTien")), "rate": rate, "rate_label": label,
                "tthue": parse_num(_text(h, "TThue")),
            })

    if inv["dvtte"].upper() not in ("VND", "VNĐ") and inv["tgia"] not in (0, 1):
        k = inv["tgia"]
        for f in ("tong_chua_thue", "tong_thue", "tong_tt", "ck_tm", "phi"):
            inv[f] *= k
        inv["lt_suat"] = [(r, lb, a * k, t * k) for r, lb, a, t in inv["lt_suat"]]
        for ln in inv["lines"]:
            for f in ("dgia", "stck", "thtien", "tthue"):
                ln[f] *= k
    return inv


# ============================================================ KIEM TRA SO HOC

def check_arithmetic(inv, tol):
    """Trả danh sách lỗi số học. Lệch SL×ĐG do ĐƠN GIÁ IN RA ĐÃ LÀM TRÒN (≤ SL×0.5đ) không tính là lỗi."""
    issues = []
    goods = 0.0
    by_rate = defaultdict(float)
    for ln in inv["lines"]:
        t = ln["tchat"]
        if t == "4":
            continue
        amount = ln["thtien"]
        if t == "3":
            amount = -abs(amount)
        elif t == "1" and ln["sl"] and ln["dgia"]:
            gross = ln["sl"] * ln["dgia"]
            if ln["stck"]:
                expected = gross - ln["stck"]
            elif ln["tlck"]:
                expected = gross * (1 - ln["tlck"] / 100)
            else:
                expected = gross
            diff = ln["thtien"] - expected
            if abs(diff) > tol and abs(diff) >= 0.995:
                if abs(diff) <= abs(ln["sl"]) * 0.5 + 0.5:
                    inv["rounding_lines"] += 1
                else:
                    issues.append(("Dòng hàng: SL × ĐG ≠ Thành tiền", diff,
                                   f"Dòng {ln['stt']} '{ln['ten']}': {ln['sl']:g} × {ln['dgia']:,.2f} → {expected:,.2f}, HĐ ghi {ln['thtien']:,.0f}"))
        goods += amount
        if ln["rate"] is not None:
            by_rate[ln["rate_label"]] += amount

    head = inv["tong_chua_thue"]
    if inv["lines"]:
        diff = head - goods
        if abs(diff) > tol and abs(diff + inv["ck_tm"]) > tol:
            issues.append(("Tổng tiền hàng ≠ cộng các dòng", diff, f"Cộng dòng {goods:,.0f}, HĐ ghi tổng chưa thuế {head:,.0f}"))

    if inv["lt_suat"]:
        for rate, label, base, tax in inv["lt_suat"]:
            if rate is None:
                continue
            expected = round_half_up(base * rate)
            diff = tax - expected
            if abs(diff) > tol:
                issues.append((f"Tiền thuế {label} ≠ tiền chưa thuế × thuế suất", diff,
                               f"{base:,.0f} × {label} = {expected:,.0f}, HĐ ghi {tax:,.0f}"))
            if label in by_rate and abs(by_rate[label] - base) > tol:
                issues.append((f"Tiền chưa thuế {label}: bảng tổng hợp ≠ cộng các dòng", base - by_rate[label],
                               f"Cộng dòng {by_rate[label]:,.0f}, bảng thuế suất ghi {base:,.0f}"))
        sum_tax = sum(t for _, _, _, t in inv["lt_suat"])
        if abs(sum_tax - inv["tong_thue"]) > tol:
            issues.append(("Tổng thuế ≠ cộng thuế các thuế suất", inv["tong_thue"] - sum_tax,
                           f"Cộng {sum_tax:,.0f}, HĐ ghi {inv['tong_thue']:,.0f}"))
    elif by_rate:
        expected = sum(round_half_up(b * (parse_rate(lb)[0] or 0)) for lb, b in by_rate.items())
        diff = inv["tong_thue"] - expected
        if abs(diff) > tol:
            issues.append(("Tổng thuế ≠ tiền chưa thuế × thuế suất", diff,
                           f"Tính lại {expected:,.0f}, HĐ ghi {inv['tong_thue']:,.0f}"))

    expected_total = head + inv["tong_thue"] + inv["phi"]
    diff = inv["tong_tt"] - expected_total
    if abs(diff) > tol and abs(diff + inv["ck_tm"]) > tol:
        issues.append(("Tổng thanh toán ≠ chưa thuế + thuế", diff,
                       f"{head:,.0f} + {inv['tong_thue']:,.0f} = {expected_total:,.0f}, HĐ ghi {inv['tong_tt']:,.0f}"))
    return issues


# ============================================================ BANG QUY DOI KIOTVIET

def unit_keyword(label):
    """'6 lốc (24h)' -> ('loc', False) ; 'lốc (4h)' -> ('loc', True) ; '10 gói' -> ('goi', False)."""
    toks = norm_text(re.sub(r"\(.*?\)", " ", str(label or ""))).split()
    for i, t in enumerate(toks):
        if re.fullmatch(r"[a-z]+", t):
            return t, i == 0
    return "", False


def kw_same(a, b):
    """Cùng loại ĐVT, chỉ khác cách ghi: 'lốc' vs 'lốc (4h)'."""
    ka, la = unit_keyword(a)
    kb, lb = unit_keyword(b)
    return bool(ka) and ka == kb and la and lb


def raw_kiotviet_to_products(items):
    """Dữ liệu gốc KiotViet (/api/kiotviet/items/all) -> danh sách products như bảng quy đổi. Bỏ hàng xóa / ngừng KD."""
    by_id = {it.get("Id"): it for it in items}
    out = []
    for it in items:
        if it.get("isDeleted") or not it.get("isActive", True):
            continue
        mid = it.get("MasterUnitId") or it.get("Id")
        m = by_id.get(mid) or {}
        out.append({"id": it.get("Id"), "code": it.get("Code"), "name": it.get("Name"), "fullName": it.get("FullName"),
                    "unit": it.get("Unit"), "masterId": mid, "masterCode": m.get("Code", ""),
                    "masterName": m.get("FullName") or m.get("Name", ""), "masterUnit": m.get("Unit", ""),
                    "conversion": 1 if mid == it.get("Id") else (it.get("ConversionValue") or 1)})
    return out


class Catalog:
    """Bảng quy đổi KiotViet: mỗi mã hàng → mã master + hệ số. Nhận file từ nút 'Xuất bảng quy đổi' (BanHang /
    BE /api/kiotviet/unit-conversion) hoặc dữ liệu gốc /api/kiotviet/items/all."""

    def __init__(self, path):
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, list) or (isinstance(data, dict) and "Data" in data and "schema" not in data):
            products = raw_kiotviet_to_products(data if isinstance(data, list) else data["Data"])
            self.exported_at = datetime.fromtimestamp(os.path.getmtime(path)).isoformat()
        elif str(data.get("schema", "")).startswith("taphoa39-unit-conversion/"):
            products = [p for p in data.get("products", []) if p.get("isActive", True) and not p.get("isDeleted")]
            self.exported_at = data.get("exportedAt", "")
        else:
            raise ValueError("Không phải bảng quy đổi (BanHang 'Xuất bảng quy đổi') hay dữ liệu gốc KiotViet items/all")
        self.path = path
        self.by_code = {}
        self.masters = {}          # mã master -> {"code","name","unit","units": {đvt chuẩn hóa: hệ số}, "unit_list"}
        by_master_id = defaultdict(list)
        for p in products:
            if p.get("isClone"):
                continue
            by_master_id[p.get("masterId")].append(p)
        for mid, prods in by_master_id.items():
            master = next((p for p in prods if p.get("id") == mid), None)
            mcode = (master or prods[0]).get("masterCode") or (master or prods[0]).get("code")
            if not mcode:
                continue
            g = self.masters.setdefault(norm_code(mcode), {
                "code": mcode,
                "name": (master or {}).get("fullName") or (master or {}).get("name") or prods[0].get("masterName") or prods[0].get("fullName"),
                "unit": (master or {}).get("unit") or prods[0].get("masterUnit") or "",
                "units": {}, "unit_list": [], "names": set(), "tax": (master or {}).get("tax"),
            })
            for p in prods:
                conv = float(p.get("conversion") or 1)
                un = norm_unit(p.get("unit"))
                if un not in g["units"]:
                    g["units"][un] = conv
                    kw, leading = unit_keyword(p.get("unit"))
                    g["unit_list"].append((p.get("unit") or "", conv, kw, leading))
                for nm in (p.get("fullName"), p.get("name")):
                    if nm:
                        g["names"].add(base_name(norm_text(nm)))
                code = norm_code(p.get("code"))
                if code:
                    self.by_code[code] = (norm_code(mcode), conv, p)
        # chỉ mục tên → master
        df = Counter()
        for mk, g in self.masters.items():
            for t in {t for nm in g["names"] for t in nm.split()}:
                df[t] += 1
        n = max(1, len(self.masters))
        self.idf = {t: math.log(1 + n / c) for t, c in df.items()}
        self.inverted = defaultdict(set)
        for mk, g in self.masters.items():
            for t in {t for nm in g["names"] for t in nm.split()}:
                self.inverted[t].add(mk)

    def by_name(self, base, top=3):
        toks = set(base.split())
        cands = set()
        for t in toks:
            if self.idf.get(t, 0) > 0.3:
                cands |= self.inverted.get(t, set())
        wa = sum(self.idf.get(t, math.log(2)) for t in toks)
        scored = []
        for mk in cands:
            best = 0.0
            for nm in self.masters[mk]["names"]:
                gt = set(nm.split())
                inter = sum(self.idf.get(t, 0) for t in toks & gt)
                wb = sum(self.idf.get(t, 0) for t in gt)
                if wa + wb:
                    best = max(best, 2 * inter / (wa + wb))
            scored.append((best, mk))
        scored.sort(reverse=True)
        return scored[:top]

    def name_of(self, mk):
        g = self.masters.get(norm_code(mk))
        return (g["name"], g["unit"]) if g else (None, None)


def match_unit(g, dvt, named):
    """Hệ số cho ĐVT trên HĐ trong các ĐVT KiotViet của 1 mặt hàng → (hệ số, cách, nhãn ĐVT KiotViet, [ứng viên])."""
    units = g["unit_list"]
    un = norm_unit(dvt)
    if un in g["units"]:
        return g["units"][un], "exact", dvt, []
    kw, _ = unit_keyword(dvt)
    cands = [u for u in units if kw and u[2] == kw]
    pref = [u for u in cands if u[3]] or cands
    if pref and len({u[1] for u in pref}) == 1:
        return pref[0][1], "keyword", pref[0][0], cands
    if named:
        same = [u for u in units if abs(u[1] - named) < 1e-9]
        if same:
            return float(named), "count", same[0][0], cands
    return None, None, None, cands


def _units_text(g):
    return ", ".join(f"{lb}={cv:g}" for lb, cv, _, _ in sorted(g["unit_list"], key=lambda u: u[1]))


def find_catalog(data_dir, patterns):
    files = []
    for pattern in ([patterns] if isinstance(patterns, str) else patterns):
        files += glob.glob(os.path.join(data_dir, pattern))
    files = sorted(set(files), key=os.path.getmtime)
    return files[-1] if files else None


# ============================================================ ANH XA MAT HANG

MAP_HEADERS = ["Chiều", "MST người bán", "Tên người bán", "Mã hàng (HĐ)", "Tên hàng (HĐ)", "ĐVT (HĐ)", "Số dòng",
               "Tổng SL (HĐ)", "Mã master", "Tên master", "ĐVT master", "Hệ số quy đổi", "Tổng SL quy đổi",
               "Mã Vietda", "Tên Vietda", "ĐVT Vietda", "Cách ghép Vietda",
               "Cách ghép", "Độ tin cậy", "Cần xem", "Bản sao XML", "Gợi ý khác", "Đã kiểm (x)",
               "Máy gợi ý mã", "Máy gợi ý hệ số", "Máy gợi ý Vietda", "Khóa (không sửa)"]
XML_PLAN_KEEP = "Giữ nguyên — kiểm rồi đánh x"
COL = {h: i for i, h in enumerate(MAP_HEADERS)}
EDITABLE = ("Mã master", "Hệ số quy đổi", "Mã Vietda", "Đã kiểm (x)")
HIDDEN = ("Máy gợi ý mã", "Máy gợi ý hệ số", "Máy gợi ý Vietda")


class VietdaCatalog:
    """Danh mục vật tư - hàng hóa xuất từ Vietda (file QUYDOI*.xlsx): Mã hàng (kế toán tự đặt), Tên, ĐVT."""

    def __init__(self, path):
        self.path = path
        self.rows, self.by_code, self.by_name, self.by_base = [], {}, defaultdict(list), defaultdict(list)
        wb = load_workbook(path, data_only=True, read_only=True)
        try:
            all_rows = list(wb.worksheets[0].iter_rows(values_only=True))
        finally:
            wb.close()  # read_only giữ file mở tới khi close → Windows không xóa được file (KeToanBackEnd dọn thư mục tạm)
        cols = None
        for vals in all_rows:
            head = [norm_text(v) if isinstance(v, str) else "" for v in vals]
            if cols is None:
                if any(h.startswith("ma hang") for h in head) and any(h.startswith("ten hang") for h in head):
                    cols = {"ma": next(i for i, h in enumerate(head) if h.startswith("ma hang")),
                            "ten": next(i for i, h in enumerate(head) if h.startswith("ten hang")),
                            "dvt": next((i for i, h in enumerate(head) if h in ("dvt", "don vi tinh")), None)}
                continue
            ma, ten = vals[cols["ma"]], vals[cols["ten"]]
            if not (isinstance(ma, str) and ma.strip() and isinstance(ten, str) and ten.strip()):
                continue
            if len(ma.strip()) == 1 and len(ten.strip()) == 1:
                continue  # dòng ký hiệu cột "A | B | C" ngay dưới tiêu đề
            r = {"ma": ma.strip(), "ten": " ".join(ten.split()),
                 "dvt": str(vals[cols["dvt"]] or "").strip() if cols["dvt"] is not None else ""}
            self.rows.append(r)
            self.by_code[norm_code(r["ma"])] = r
            self.by_name[norm_text(r["ten"])].append(r)
            self.by_base[base_name(norm_text(r["ten"]))].append(r)
        if cols is None:
            raise ValueError("không thấy dòng tiêu đề có 'Mã hàng' và 'Tên hàng'")
        df = Counter(t for r in self.rows for t in set(base_name(norm_text(r["ten"])).split()))
        n = max(1, len(self.rows))
        self.idf = {t: math.log(1 + n / c) for t, c in df.items()}
        self.inverted = defaultdict(set)
        for i, r in enumerate(self.rows):
            for t in set(base_name(norm_text(r["ten"])).split()):
                self.inverted[t].add(i)

    def search(self, base, top=3):
        """Tên gần giống (trọng số IDF, giống Catalog.by_name) → [(điểm, dòng Vietda)]."""
        toks = set(base.split())
        cands = set()
        for t in toks:
            if self.idf.get(t, 0) > 0.3:
                cands |= self.inverted.get(t, set())
        wa = sum(self.idf.get(t, math.log(2)) for t in toks)
        scored = []
        for i in cands:
            gt = set(base_name(norm_text(self.rows[i]["ten"])).split())
            inter = sum(self.idf.get(t, 0) for t in toks & gt)
            wb = sum(self.idf.get(t, 0) for t in gt)
            if wa + wb:
                scored.append((2 * inter / (wa + wb), i))
        scored.sort(reverse=True)
        return [(s, self.rows[i]) for s, i in scored[:top]]


def suggest_vietda(it, master_code, master_unit, vd, master_vd, threshold):
    """Mã Vietda cho 1 mặt hàng → (dòng Vietda | None, cách ghép, [cờ])."""
    flags = []
    units = {norm_unit(it["dvt"]), norm_unit(master_unit)}
    hits = vd.by_name.get(norm_text(it["ten"]))
    if hits:
        pick = next((h for h in hits if norm_unit(h["dvt"]) in units), hits[0])
        if len(hits) > 1:
            flags.append("Tên trùng nhiều mã Vietda: " + ", ".join(f"{h['ma']} ({h['dvt']})" for h in hits))
        main = master_vd.get(norm_code(master_code)) if master_code else None
        if main and norm_code(main) != norm_code(pick["ma"]):
            flags.append(f"Mã Vietda khác mã Vietda của các dòng cùng mã master ({main})")
        return pick, "Tên HĐ = tên Vietda", flags
    if master_code and norm_code(master_code) in master_vd:
        return vd.by_code.get(norm_code(master_vd[norm_code(master_code)])), "Theo mã master", flags
    ranked = vd.search(base_name(norm_text(it["ten"])))
    if ranked and ranked[0][0] >= threshold:
        flags.append(f"Mã Vietda gợi ý theo tên gần giống ({ranked[0][0]:.2f}) — kiểm")
        return ranked[0][1], "Tên gần giống", flags
    flags.append("Chưa có mã Vietda" + (f" (gần nhất: {ranked[0][1]['ma']} {ranked[0][1]['ten']} {ranked[0][0]:.2f})" if ranked else ""))
    return None, "", flags


def find_vietda(data_dir, patterns):
    files = []
    for fn in os.listdir(data_dir):
        low = fn.lower()
        if low.endswith(".xlsx") and not low.startswith("~$") and any(fnmatch.fnmatch(low, p.lower()) for p in patterns):
            files.append(os.path.join(data_dir, fn))
    return max(files, key=os.path.getmtime) if files else None


def item_key(direction, mst, ma, ten, dvt):
    return "|".join([direction, mst if direction == "VAO" else "", norm_code(ma), norm_text(ten), norm_unit(dvt)])


def suggest(it, cat, threshold):
    """Trả (mã master, hệ số, cách ghép, độ tin cậy, [cờ], gợi ý khác)."""
    flags = []
    unit_n = norm_unit(it["dvt"])
    hit = cat.by_code.get(norm_code(it["ma"])) if it["ma"] else None
    if hit:
        mk, conv, p = hit
        if norm_unit(p.get("unit")) != unit_n and unit_n:
            alt, how, _, _ = match_unit(cat.masters[mk], it["dvt"], None)
            if alt and how == "exact":
                flags.append(f"ĐVT HĐ '{it['dvt']}' khác ĐVT của mã ({p.get('unit')}) — dùng hệ số của '{it['dvt']}'")
                conv = alt
            elif kw_same(it["dvt"], p.get("unit")):
                pass   # 'lốc' vs 'lốc (4h)': cùng ĐVT, chỉ khác cách ghi
            else:
                flags.append(f"ĐVT HĐ '{it['dvt']}' khác ĐVT KiotViet '{p.get('unit')}'")
        return mk, conv, "Mã hàng", 1.0, flags, ""

    base = base_name(norm_text(it["ten"]))
    ranked = cat.by_name(base)
    others = " | ".join(f"{cat.masters[m]['code']} {cat.masters[m]['name']} ({s:.2f})" for s, m in ranked[1:])
    if not ranked or ranked[0][0] < threshold:
        if ranked:
            others = " | ".join(f"{cat.masters[m]['code']} {cat.masters[m]['name']} ({s:.2f})" for s, m in ranked)
        return "", None, "", round(ranked[0][0], 2) if ranked else 0.0, ["Chưa ghép được — điền Mã master"], others
    score, mk = ranked[0]
    g = cat.masters[mk]
    method = "Tên + ĐVT KiotViet"
    named = pack_count_from_name(norm_text(it["ten"]), unit_n)
    conv, how, label, cands = match_unit(g, it["dvt"], named)
    if how == "keyword" and len(cands) > 1:
        flags.append(f"ĐVT '{it['dvt']}' → chọn '{label}' = {conv:g} trong: {', '.join(f'{c[0]}={c[1]:g}' for c in cands)}")
    elif how == "count":
        method = "Tên + số trong tên hàng"
        flags.append(f"ĐVT HĐ '{it['dvt']}' không có trên KiotViet — số {named} trong tên hàng khớp ĐVT KiotViet '{label}'")
    if conv is None:
        if named:
            conv = float(named)
            method = "Tên + số trong tên hàng"
            flags.append(f"KiotViet không có ĐVT '{it['dvt']}' — hệ số {named} đọc từ tên hàng (KiotViet có: {_units_text(g)})")
        else:
            flags.append(f"Thiếu hệ số: KiotViet không có ĐVT '{it['dvt']}' (có: {_units_text(g)})")
    elif named and how != "count" and abs(named - conv) > 1e-9:
        flags.append(f"Tên hàng ghi {named}/{it['dvt']} nhưng KiotViet '{label}' = {conv:g} {g['unit']}")
    if score < 0.85:
        flags.append("Độ tin cậy thấp")
    return g["code"], conv, method, round(score, 2), flags, others


def load_mapping(path):
    rows = {}
    if not os.path.exists(path):
        return rows
    wb = load_workbook(path)
    for ws in wb.worksheets:
        if ws.title not in ("Dau vao", "Dau ra"):
            continue
        # đọc theo TÊN cột → file của bản trước (ít cột hơn) vẫn giữ được phần kế toán đã sửa
        idx = {str(c.value).strip(): i for i, c in enumerate(ws[1]) if c.value not in (None, "")}
        if "Khóa (không sửa)" not in idx or "Mã master" not in idx:
            log(f"  ! {os.path.basename(path)} không đúng mẫu (thiếu cột 'Mã master' / 'Khóa') — tạo lại từ đầu")
            return {}
        k = idx["Khóa (không sửa)"]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r and len(r) > k and r[k]:
                rows[r[k]] = [r[idx[h]] if h in idx and idx[h] < len(r) else None for h in MAP_HEADERS]
    return rows


def _same(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) or isinstance(b, (int, float)):
        return abs(parse_num(a) - parse_num(b)) < 1e-9
    return norm_code(a) == norm_code(b)


def build_mapping(items, existing, cat, threshold, vd=None):
    result, self_flags, master_ok = {}, {}, {}
    for k, it in items.items():
        old = existing.get(k)
        keep = False
        if old:
            checked = str(old[COL["Đã kiểm (x)"]] or "").strip().lower() == "x"
            edited = not (_same(old[COL["Mã master"]], old[COL["Máy gợi ý mã"]])
                          and _same(old[COL["Hệ số quy đổi"]], old[COL["Máy gợi ý hệ số"]]))
            keep = checked or edited
        mk, conv, method, conf, flags, others = suggest(it, cat, threshold)
        row = [None] * len(MAP_HEADERS)
        row[COL["Chiều"]] = "Vào" if it["direction"] == "VAO" else "Ra"
        row[COL["MST người bán"]] = it["mst"] if it["direction"] == "VAO" else ""
        row[COL["Tên người bán"]] = it["ten_ban"] if it["direction"] == "VAO" else ""
        row[COL["Mã hàng (HĐ)"]] = it["ma"]
        row[COL["Tên hàng (HĐ)"]] = it["ten"]
        row[COL["ĐVT (HĐ)"]] = it["dvt"]
        row[COL["Số dòng"]] = it["lines"]
        row[COL["Tổng SL (HĐ)"]] = round(it["qty"], 3)
        row[COL["Máy gợi ý mã"]] = mk
        row[COL["Máy gợi ý hệ số"]] = conv
        row[COL["Khóa (không sửa)"]] = k
        row[COL["Gợi ý khác"]] = others
        if keep:
            row[COL["Mã master"]] = old[COL["Mã master"]]
            row[COL["Hệ số quy đổi"]] = old[COL["Hệ số quy đổi"]]
            row[COL["Đã kiểm (x)"]] = old[COL["Đã kiểm (x)"]]
            edited = not (_same(old[COL["Mã master"]], mk) and _same(old[COL["Hệ số quy đổi"]], conv))
            row[COL["Cách ghép"]] = "Kế toán sửa" if edited else method
            row[COL["Độ tin cậy"]] = 1.0 if edited else conf
            row[COL["Cần xem"]] = ""
        else:
            row[COL["Mã master"]] = mk
            row[COL["Hệ số quy đổi"]] = conv
            row[COL["Đã kiểm (x)"]] = ""
            row[COL["Cách ghép"]] = method
            row[COL["Độ tin cậy"]] = conf
            row[COL["Cần xem"]] = "; ".join(flags)
        mname, munit = cat.name_of(row[COL["Mã master"]]) if row[COL["Mã master"]] else (None, None)
        if row[COL["Mã master"]] and mname is None:
            mname, munit = f"(ngoài KiotViet) {row[COL['Mã master']]}", it["dvt"]
        row[COL["Tên master"]] = mname or ""
        row[COL["ĐVT master"]] = munit or ""
        factor = parse_num(row[COL["Hệ số quy đổi"]]) if row[COL["Hệ số quy đổi"]] not in (None, "") else 0
        row[COL["Tổng SL quy đổi"]] = round(it["qty"] * factor, 3) if factor else None
        result[k] = row
        self_flags[k] = (old, checked if old else False)
        # chắc chắn = kế toán đã sửa / đánh x, hoặc ghép CHÍNH XÁC theo mã vạch mà không có cờ nào
        master_ok[k] = keep or (method == "Mã hàng" and not flags)
    vd_ok = _apply_vietda(items, result, self_flags, vd, cat, threshold, master_ok) if vd is not None else {}
    for k, row in result.items():
        factor = parse_num(row[COL["Hệ số quy đổi"]]) if row[COL["Hệ số quy đổi"]] not in (None, "") else 0
        unit_ok = master_ok[k] and row[COL["Mã master"]] and factor > 0
        if unit_ok:
            plan = "Đổi ĐVT + " + ("mã Vietda" if vd_ok.get(k) else "mã master")
        elif vd_ok.get(k):
            plan = "Chỉ đổi mã Vietda" + (" (quy đổi ĐVT chờ xác nhận)" if row[COL["Mã master"]] else "")
        else:
            plan = XML_PLAN_KEEP
        row[COL["Bản sao XML"]] = plan
    return result


def _apply_vietda(items, result, olds, vd, cat, threshold, master_ok):
    """Lượt 2 của build_mapping: điền Mã Vietda (tên HĐ = tên Vietda → theo mã master → tên gần giống); giữ mã kế toán sửa.
    → {khóa: True nếu mã Vietda CHẮC CHẮN}. Chỉ dòng có Mã master chắc chắn mới được 'bầu' mã Vietda cho mã master đó."""
    votes = defaultdict(Counter)   # mã master → mã Vietda của các dòng có tên trùng đúng tên Vietda
    for k, it in items.items():
        hits = vd.by_name.get(norm_text(it["ten"]))
        mk = result[k][COL["Mã master"]]
        if hits and len(hits) == 1 and mk and master_ok[k]:
            votes[norm_code(mk)][hits[0]["ma"]] += it["lines"]
    master_vd = {mk: c.most_common(1)[0][0] for mk, c in votes.items()}
    for mk, g in cat.masters.items():   # chưa có phiếu bầu: tên KiotViet của mã master trùng tên Vietda
        if mk not in master_vd:
            hit = next((vd.by_base[nm][0] for nm in g["names"] if vd.by_base.get(nm)), None)
            if hit:
                master_vd[mk] = hit["ma"]
    ok = {}
    for k, it in items.items():
        row = result[k]
        old, checked = olds[k]
        pick, method, flags = suggest_vietda(it, row[COL["Mã master"]], row[COL["ĐVT master"]], vd, master_vd, threshold)
        row[COL["Máy gợi ý Vietda"]] = pick["ma"] if pick else ""
        vd_edited = bool(old) and not _same(old[COL["Mã Vietda"]], old[COL["Máy gợi ý Vietda"]])
        confirmed = bool(old) and (vd_edited or (checked and old[COL["Mã Vietda"]] not in (None, "")))
        if confirmed:
            code = old[COL["Mã Vietda"]]
            row[COL["Mã Vietda"]] = code
            changed = not _same(code, row[COL["Máy gợi ý Vietda"]])
            row[COL["Cách ghép Vietda"]] = "Kế toán sửa" if changed else method
            flags = []
        else:
            code = row[COL["Máy gợi ý Vietda"]]
            row[COL["Mã Vietda"]] = code
            row[COL["Cách ghép Vietda"]] = method
        r = vd.by_code.get(norm_code(code)) if code else None
        if code and r is None:
            flags.append(f"Mã Vietda '{code}' không có trong {os.path.basename(vd.path)}")
        row[COL["Tên Vietda"]] = r["ten"] if r else ""
        row[COL["ĐVT Vietda"]] = r["dvt"] if r else ""
        unit_after = row[COL["ĐVT master"]] or it["dvt"]
        if r and r["dvt"] and unit_after and norm_unit(r["dvt"]) != norm_unit(unit_after) and not kw_same(r["dvt"], unit_after):
            if not confirmed:
                flags.append(f"ĐVT Vietda '{r['dvt']}' khác ĐVT sau quy đổi '{unit_after}'")
        if flags:
            row[COL["Cần xem"]] = "; ".join(x for x in (row[COL["Cần xem"]], "; ".join(flags)) if x)
        # chắc chắn = kế toán xác nhận, hoặc ghép CHÍNH XÁC mà không có cờ nào: tên HĐ trùng đúng tên Vietda, hay theo mã
        # master khi chính Mã master của dòng đã chắc chắn (mã vạch / kế toán xác nhận)
        exact = method == "Tên HĐ = tên Vietda" or (method == "Theo mã master" and master_ok[k])
        ok[k] = bool(r) and (confirmed or (exact and not flags))
    return ok


def save_mapping(path, mapping):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Huong dan"
    guide = [
        "ÁNH XẠ MẶT HÀNG → MÃ MASTER KIOTVIET — kế toán chỉ sửa 2 cột tô vàng: 'Mã master' và 'Hệ số quy đổi'",
        "",
        "• Mã master: mã hàng KiotViet của ĐVT NHỎ NHẤT (ĐVT master). Dòng vào và dòng ra cùng 1 mặt hàng phải cùng Mã master.",
        "• Hệ số quy đổi: 1 ĐVT trên HĐ = bao nhiêu ĐVT master. VD HĐ vào 'Thùng' sữa chua 48 hộp → 48.",
        "• Máy đã tự điền: đầu ra ghép theo mã vạch (chính xác), đầu vào ghép theo tên + ĐVT trong KiotViet.",
        "• Cột 'Cần xem' có nội dung = nên kiểm. Các dòng đó nằm TRÊN CÙNG.",
        "• Mã master trống hoặc hệ số trống = chưa ánh xạ → dòng hàng không vào đối soát, báo trong 'Ngoai le'.",
        "• Hàng không có trên KiotViet (VPP, hàng không bán...): tự đặt 1 mã bất kỳ (vd NB-KHANGIAY) cho cùng nhóm.",
        "• Mã Vietda: mã hàng kế toán tự đặt trong Vietda (file QUYDOI*.xlsx ở thư mục gốc). Bản sao XML quy đổi đổi mã hàng",
        "  sang mã này (không có thì dùng Mã master). Máy gợi ý: tên trên HĐ trùng tên Vietda → mã Vietda của các dòng cùng",
        "  Mã master → tên gần giống (có cờ 'Cần xem'). Sai thì sửa cột 'Mã Vietda'.",
        "• Đánh 'x' cột 'Đã kiểm' để khóa dòng. Dòng đã sửa tay cũng tự được giữ ở lần chạy sau.",
        "• KHÔNG sửa cột 'Khóa' và các cột 'Máy gợi ý' (đang ẩn). Đóng file Excel trước khi chạy lại script.",
    ]
    for line in guide:
        ws0.append([line])
    ws0.column_dimensions["A"].width = 140
    ws0["A1"].font = Font(bold=True, size=13)
    widths = [6, 13, 26, 15, 45, 8, 7, 9, 15, 38, 9, 8, 11, 12, 38, 9, 18, 18, 8, 50, 26, 50, 8, 14, 9, 12, 30]
    for title, direction in (("Dau vao", "Vào"), ("Dau ra", "Ra")):
        ws = wb.create_sheet(title)
        ws.append(MAP_HEADERS)
        rows = [r for r in mapping.values() if r[COL["Chiều"]] == direction]
        rows.sort(key=lambda r: (0 if r[COL["Cần xem"]] else 1, r[COL["Độ tin cậy"]] or 0, str(r[COL["Tên hàng (HĐ)"]])))
        for r in rows:
            ws.append(r)
        _style_sheet(ws, widths, editable_cols=tuple(COL[h] for h in EDITABLE))
        for h in HIDDEN:
            ws.column_dimensions[_col_letter(COL[h] + 1)].hidden = True
    wb.save(path)


def _style_sheet(ws, widths, editable_cols=()):
    head_fill = PatternFill("solid", fgColor="1F4E78")
    edit_fill = PatternFill("solid", fgColor="FFF2CC")
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for ci in editable_cols:
        ws.cell(row=1, column=ci + 1).fill = PatternFill("solid", fgColor="BF8F00")
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=ci + 1).fill = edit_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ============================================================ TON DAU KY

OPEN_HEADERS = ["Mã master", "Tên master", "ĐVT master", "SL tồn đầu kỳ (ĐVT master)", "Đơn giá vốn (chưa thuế)"]


def load_opening(path):
    data = {}
    if not os.path.exists(path):
        return data
    ws = load_workbook(path, data_only=True).worksheets[0]
    if [c.value for c in ws[1]][:len(OPEN_HEADERS)] != OPEN_HEADERS:
        log(f"  ! {os.path.basename(path)} là bản cũ (khác cột) — bỏ qua, hãy xóa file này để tạo lại")
        return data
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0]:
            qty = parse_num(r[3]) if len(r) > 3 and r[3] not in (None, "") else 0.0
            cost = parse_num(r[4]) if len(r) > 4 and r[4] not in (None, "") else 0.0
            data[norm_code(r[0])] = (qty, cost)
    return data


def save_opening_template(path, groups):
    """groups: mã master chuẩn hóa -> (mã, tên, đvt). Thêm nhóm mới, không đụng số đã điền."""
    if os.path.exists(path):
        wb = load_workbook(path)
        ws = wb.worksheets[0]
        if [c.value for c in ws[1]][:len(OPEN_HEADERS)] != OPEN_HEADERS:
            return 0
        known = {norm_code(r[0]) for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}
        added = 0
        for gk, (code, name, unit) in sorted(groups.items(), key=lambda x: str(x[1][1])):
            if gk not in known:
                ws.append([code, name, unit, None, None])
                added += 1
        if added:
            wb.save(path)
        return added
    wb = Workbook()
    ws = wb.active
    ws.title = "Ton dau ky"
    ws.append(OPEN_HEADERS)
    for gk, (code, name, unit) in sorted(groups.items(), key=lambda x: str(x[1][1])):
        ws.append([code, name, unit, None, None])
    _style_sheet(ws, [16, 50, 10, 18, 18], editable_cols=(3, 4))
    wb.save(path)
    return len(groups)


# ============================================================ BAO CAO EXCEL

class Report:
    def __init__(self):
        self.wb = Workbook(write_only=True)
        self.head_font = Font(bold=True, color="FFFFFF")
        self.head_fill = PatternFill("solid", fgColor="1F4E78")

    def sheet(self, title, headers, rows, widths, formats=None, wrap=False):
        ws = self.wb.create_sheet(title)
        top_wrap = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[_col_letter(i)].width = w
        ws.freeze_panes = "A2"
        head = []
        for h in headers:
            c = WriteOnlyCell(ws, value=h)
            c.font = self.head_font
            c.fill = self.head_fill
            c.alignment = Alignment(wrap_text=True, vertical="center")
            head.append(c)
        ws.append(head)
        formats = formats or {}
        n = 0
        for r in rows:
            cells = []
            for i, v in enumerate(r):
                c = WriteOnlyCell(ws, value=v)
                if isinstance(v, date) and not isinstance(v, datetime):
                    c.number_format = "dd/mm/yyyy"
                elif i in formats and isinstance(v, (int, float)):
                    c.number_format = formats[i]
                if wrap:
                    c.alignment = top_wrap
                    if i == 0:
                        c.font = Font(bold=True)
                cells.append(c)
            ws.append(cells)
            n += 1
        if not wrap:
            ws.auto_filter.ref = f"A1:{_col_letter(len(headers))}{n + 1}"
        return n

    def save(self, path):
        self.wb.save(path)


def _col_letter(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


MONEY = "#,##0"
MONEY2 = "#,##0.00"
QTY = "#,##0.###"
PCT = "0.0%"

# Sheet "Huong dan" — (sheet, mục đích, cách đọc / cột chính, khi nào cần làm gì)
REPORT_GUIDE = [
    ("Tong quan",
     "Tóm tắt lần chạy: thư mục, bảng quy đổi KiotViet đã dùng, MST công ty, số HĐ vào/ra được tính, số mặt hàng theo cách ghép, số ngoại lệ theo mức độ và theo loại.",
     "Kiểm 3 dòng trước: 'MST công ty' đúng chưa; 'File không phải hóa đơn / lỗi' nên = 0; 'Có tồn đầu kỳ'. "
     "'Dòng lệch do làm tròn đơn giá' chỉ là thông tin, không phải lỗi.",
     "MST sai → sửa CONFIG['mst_cong_ty'] đầu file doi_soat.py. 'Mặt hàng cần xem' > 0 → mở anh_xa_hang_hoa.xlsx."),
    ("Danh sach san pham",
     "KẾT QUẢ QUY ĐỔI: mọi mặt hàng xuất hiện trên HĐ (mỗi tổ hợp Chiều + NCC + Mã + Tên + ĐVT trên HĐ) đã ghép với mã master KiotViet và hệ số quy đổi về ĐVT nhỏ nhất.",
     "'Tổng SL (HĐ)' tính theo 'ĐVT (HĐ)'; × 'Hệ số quy đổi' = 'Tổng SL quy đổi' tính theo 'ĐVT master'. "
     "'Cách ghép': Mã hàng = khớp mã vạch (chính xác) · Tên + ĐVT KiotViet = khớp tên, hệ số lấy từ KiotViet · "
     "Tên + số trong tên hàng = KiotViet không có ĐVT này, hệ số đọc từ tên (nên kiểm) · Kế toán sửa · trống = chưa ghép. "
     "'Độ tin cậy' 0–1 (mức giống tên). 'Cần xem' ghi lý do.",
     "Dòng có 'Cần xem' hoặc 'Mã master' trống → sửa trong anh_xa_hang_hoa.xlsx (KHÔNG sửa ở báo cáo này — báo cáo tạo lại mỗi lần chạy), rồi chạy lại."),
    ("Doi soat theo master",
     "Đối soát NHẬP – XUẤT từng mặt hàng sau khi đã quy về ĐVT master: hàng bán ra có đủ hóa đơn đầu vào không, giá bán so với giá vốn, lãi gộp.",
     "Tồn đầu (từ ton_dau_ky.xlsx) + Nhập SL − Xuất SL = 'Tồn cuối (theo HĐ)'. Giá nhập BQ / Giá bán BQ = tiền chưa thuế ÷ SL master. "
     "'Bán vượt HĐ vào (SL)' = lượng bán ra KHÔNG có HĐ vào nào bù được → rủi ro thuế cao. "
     "'HĐ vào lập sau (SL)' = bán trước ngày lập HĐ vào nhưng về sau có HĐ vào đủ bù. 'Ghi chú' tóm tắt vấn đề của dòng.",
     "Tồn cuối âm / Bán vượt > 0 → kiểm theo thứ tự: hệ số quy đổi sai? thiếu file HĐ vào? chưa điền tồn đầu kỳ? "
     "Thuế suất ra ≠ vào → kiểm mặt hàng có được giảm thuế (8%) không."),
    ("Ngoai le",
     "DANH SÁCH VIỆC CẦN XỬ LÝ — mọi vấn đề phát hiện được, xếp Cao → Trung bình → Thấp.",
     "Loại: Bán vượt SL HĐ vào · Bán dưới giá vốn · HĐ vào lập sau ngày bán · Nhiều thuế suất ra · Thuế suất ra ≠ vào · "
     "Mặt hàng chưa ánh xạ · Ánh xạ cần xem · Lệch số học · Trùng số HĐ khác nội dung · HĐ thay thế/điều chỉnh thiếu HĐ gốc · HĐ không liên quan MST. "
     "'Giá trị / SL lệch' là SL theo ĐVT master hoặc số tiền tùy loại.",
     "Xử lý từ trên xuống; lọc cột 'Loại' để làm từng nhóm. Mỗi loại có sheet chi tiết tương ứng (xem các dòng dưới)."),
    ("HD thay the - dieu chinh",
     "Mọi HĐ thay thế / điều chỉnh / bị thay thế / bị điều chỉnh (vào và ra) và việc có tính vào KỲ TÍNH THUẾ này hay không.",
     "Quy tắc (theo kế toán): HĐ THAY THẾ tính vào kỳ của HĐ GỐC → lập trong kỳ này để thay HĐ kỳ trước thì TRỪ ra; lập ở kỳ sau để thay HĐ kỳ này thì CỘNG vào "
     "(cần chép XML HĐ thay thế lập đầu kỳ sau vào thư mục). HĐ bị thay thế không tính. HĐ ĐIỀU CHỈNH tính vào kỳ lập HĐ điều chỉnh (cộng/trừ theo dấu). "
     "'Ngày tính kỳ' = ngày dùng để xếp kỳ; 'Tính vào kỳ này' = Có/Không.",
     "Kiểm từng dòng, nhất là HĐ điều chỉnh (dấu +/− của số tiền) và dòng 'Không rõ ngày HĐ gốc'. Tổng TRỪ/CỘNG được ghi ở sheet tờ khai."),
    ("Dong hang quy doi",
     "Dữ liệu CHI TIẾT từng dòng hàng của từng HĐ (vào và ra) kèm số lượng/đơn giá đã quy về ĐVT master. Bảng nguồn để lọc, Pivot.",
     "Bên trái là dòng gốc trên HĐ (Mã, Tên, ĐVT, SL, Đơn giá, Thành tiền, Thuế suất); 'Loại dòng' = Khuyến mại nếu là hàng tặng. "
     "Bên phải là quy đổi: Mã master, ĐVT master, Hệ số, SL quy đổi, Đơn giá quy đổi (= thành tiền chưa thuế sau chiết khấu ÷ SL quy đổi). "
     "Dòng chưa ánh xạ có 'SL quy đổi' trống.",
     "Muốn xem 1 mặt hàng: lọc 'Mã master'. Muốn tổng hợp: Insert → PivotTable, hàng = Mã master, cột = Chiều, giá trị = SL quy đổi."),
    ("XML quy doi",
     "Nhật ký tạo BẢN SAO XML đã quy đổi ĐVT ở thư mục 'Dau vao (quy doi)' / 'Dau ra (quy doi)' (file gốc giữ nguyên). "
     "Mỗi dòng = 1 dòng hàng đã sửa (hoặc chưa sửa được) trong bản sao.",
     "Cũ → mới: Mã (→ Mã Vietda, chưa có thì Mã master), Tên, ĐVT, SL (× hệ số), Đơn giá (÷ hệ số). Thành tiền, thuế, "
     "tổng tiền trên HĐ KHÔNG đổi. Chỉ sửa dòng có ánh xạ chắc chắn (cột 'Bản sao XML' trong anh_xa_hang_hoa.xlsx); dòng "
     "giữ nguyên có ghi lý do.",
     "Kiểm dòng 'Giữ nguyên' / 'chờ xác nhận' ở anh_xa_hang_hoa.xlsx, sửa nếu sai, đánh x rồi chạy lại — bản sao được tạo lại "
     "từ đầu mỗi lần chạy."),
    ("Truy vet HD vao",
     "Mỗi dòng hàng MUA VÀO (lô nhập) đã được bán ra qua những HĐ nào — trả lời câu hỏi '1 thùng sữa chua 48 hộp đã bán qua những hóa đơn nào'. Phân bổ theo FIFO.",
     "'SL quy đổi' = số lượng lô theo ĐVT master; 'Đã bán' + 'Còn lại' = SL quy đổi; 'Số HĐ ra' và 'Các HĐ ra' (tối đa 30 số, đủ danh sách ở sheet Truy vet chi tiet). "
     "'Giá vốn/ĐVT master' = giá chưa thuế sau chiết khấu.",
     "Lô 'Còn lại' lớn bất thường → có thể HĐ ra ghép sai mã master, hoặc hệ số HĐ vào quá lớn."),
    ("Truy vet chi tiet",
     "Bản đầy đủ của truy vết: từng cặp (lô nhập → dòng bán) với số lượng phân bổ, giá vốn và giá bán.",
     "Mỗi dòng = 1 phần của 1 dòng bán lấy từ 1 lô nhập. Một dòng bán có thể chia cho nhiều lô.",
     "Lọc 'Số HĐ vào' hoặc 'Số HĐ ra' để xem 1 hóa đơn cụ thể."),
    ("Ban duoi gia von",
     "Các dòng BÁN RA có giá bán/ĐVT master (chưa thuế, sau chiết khấu) thấp hơn giá vốn FIFO của các lô đã phân bổ.",
     "'Chênh %' âm = bán thấp hơn giá vốn bao nhiêu %; 'Thiệt (đ)' = (giá vốn − giá bán) × SL master.",
     "Nguyên nhân thường gặp: HỆ SỐ QUY ĐỔI SAI (vd thùng ghi 24 thay vì 48 → giá vốn/hộp bị gấp đôi) — kiểm hệ số trước; "
     "sau đó mới xét giá khuyến mại, nhập giá cao."),
    ("Thieu HD vao",
     "Các dòng bán ra mà TẠI NGÀY BÁN chưa có đủ HĐ đầu vào (tính theo ngày lập HĐ vào), và các lần trả hàng NCC vượt tồn.",
     "'SL thiếu' theo ĐVT master. Phần thiếu này nếu về sau có HĐ vào bù đủ → cộng vào cột 'HĐ vào lập sau' ở sheet Doi soat theo master; "
     "nếu không → cộng vào 'Bán vượt HĐ vào'.",
     "Nhiều dòng đầu kỳ → thường do thiếu tồn đầu kỳ. NCC hay xuất HĐ trễ → tăng CONFIG['so_ngay_hd_vao_lap_sau']."),
    ("Lech so hoc HD",
     "Hóa đơn có SAI SỐ HỌC THẬT: tiền thuế ≠ tiền chưa thuế × thuế suất, tổng ≠ cộng các dòng, SL × đơn giá ≠ thành tiền (vượt mức làm tròn).",
     "'Chênh (đ)' = số trên HĐ − số tính lại; chênh ≤ 1đ ở mức Thấp. Lệch do làm tròn đơn giá KHÔNG nằm ở đây.",
     "Đối chiếu HĐ trên cổng thuế; lệch lớn → liên hệ bên bán (HĐ vào) hoặc lập HĐ điều chỉnh (HĐ ra)."),
    ("Thue theo thue suat",
     "Tổng tiền chưa thuế và tiền thuế theo từng thuế suất, tách đầu vào / đầu ra — chỉ HĐ được tính (bỏ HĐ bị thay thế, ngoài kỳ).",
     "Lấy từ bảng tổng hợp thuế suất in trên từng HĐ (THTTLTSuat).",
     "Đối chiếu với chỉ tiêu thuế đầu vào / đầu ra trên tờ khai GTGT của kỳ."),
    ("HD dau vao / HD dau ra",
     "Danh sách hóa đơn theo ngày, mỗi HĐ 1 dòng.",
     "'Trạng thái': Gốc / Thay thế / Điều chỉnh / Bị thay thế (KHÔNG được tính) / Bị điều chỉnh; thêm '· ngoài kỳ' = không thuộc kỳ tính thuế (không tính). "
     "'Số lỗi số học', 'Dòng làm tròn ĐG', 'Có mã CQT', 'File' = file XML nguồn (đường dẫn trong thư mục, 'a.zip!b.xml' = file nằm trong zip).",
     "Tìm nhanh 1 HĐ theo ký hiệu + số; kiểm HĐ thiếu mã CQT."),
    ("HD bo qua",
     "HĐ không đưa vào tính: không liên quan MST công ty (MST bán và mua đều khác) và file trùng (cùng MST bán + ký hiệu + số HĐ — chỉ giữ bản đầu).",
     "Cột 'File' cho biết file nào bị bỏ / trùng với file nào.",
     "Nhiều HĐ 'không liên quan' → có thể MST công ty tự nhận sai, hoặc lẫn file của đơn vị khác."),
    ("File loi",
     "File XML không đọc được (hỏng, tải thiếu, sai định dạng).",
     "Cột 'Lỗi' ghi lý do.",
     "Tải lại file đó từ cổng thuế / email."),
    ("", "", "", ""),
    ("THỨ TỰ NÊN XEM",
     "1) Tong quan → 2) Ngoai le → 3) Doi soat theo master → 4) Danh sach san pham → các sheet chi tiết khi cần.",
     "Sửa ánh xạ / tồn đầu kỳ ở anh_xa_hang_hoa.xlsx, ton_dau_ky.xlsx — rồi chạy lại để có báo cáo mới.", ""),
    ("", "", "", ""),
    ("THUẬT NGỮ", "", "", ""),
    ("ĐVT master", "Đơn vị nhỏ nhất của mặt hàng trong KiotViet (vd hộp). Mọi số lượng được quy về đơn vị này để so vào/ra.", "", ""),
    ("Mã master", "Mã hàng KiotViet của ĐVT master — dùng làm khóa ghép đầu vào với đầu ra.", "", ""),
    ("Hệ số quy đổi", "1 ĐVT trên HĐ = bao nhiêu ĐVT master (thùng sữa chua = 48, lốc = 4, nửa thùng = 24, hộp = 1). Lấy từ ConversionValue của KiotViet.", "", ""),
    ("FIFO", "Nhập trước – xuất trước: mỗi lần bán lấy hàng (và giá vốn) từ lô nhập sớm nhất còn hàng, chỉ lô có ngày lập HĐ ≤ ngày bán.", "", ""),
    ("Tồn cuối (theo HĐ)", "Tồn tính theo hóa đơn = tồn đầu + nhập − xuất. Không phải tồn kho trên KiotViet.", "", ""),
    ("Giá / tiền", "Luôn CHƯA thuế GTGT, đã trừ chiết khấu thương mại (phân bổ theo tỷ lệ tiền từng dòng).", "", ""),
    ("Mức độ", "Cao = rủi ro thuế/tiền lớn, cần xử lý · Trung bình = nên kiểm · Thấp = thông tin.", "", ""),
    ("Làm tròn đơn giá", "Đơn giá in trên HĐ đã làm tròn nên SL × đơn giá lệch thành tiền vài đồng (vd 2 × 31.481 = 62.962, HĐ ghi 62.963) — bình thường, không phải lỗi.", "", ""),
    ("Hàng khuyến mại", "Dòng hàng tặng: tính số lượng nhập/xuất, giá = 0, không xét bán dưới giá vốn.", "", ""),
]


# ============================================================ BẢN SAO XML ĐÃ QUY ĐỔI ĐVT

_HHDVU_RE = re.compile(r"<((?:[\w.-]+:)?)HHDVu\b[^>]*>(.*?)</\1HHDVu\s*>", re.S)
_XML_TAG_RE = {}
CLONE_HEADERS = ["Chiều", "File gốc", "File quy đổi", "Ký hiệu", "Số HĐ", "STT", "Mã cũ", "Mã mới", "Tên cũ", "Tên mới",
                 "ĐVT cũ", "ĐVT mới", "SL cũ", "SL mới", "Đơn giá cũ", "Đơn giá mới", "Hệ số", "Ghi chú"]


def _xml_tag_re(tag):
    r = _XML_TAG_RE.get(tag)
    if r is None:  # nhóm 1 = tiền tố ns, 2 = thuộc tính, 4 = nội dung (None nếu thẻ rỗng <x/>)
        r = _XML_TAG_RE[tag] = re.compile(rf"<((?:[\w.-]+:)?){tag}((?:\s[^>]*?)?)\s*(/>|>(.*?)</\1{tag}\s*>)", re.S)
    return r


def _xml_get(body, tag):
    m = _xml_tag_re(tag).search(body)
    return m.group(4) if m else None


def _xml_set(body, tag, value, before=None):
    """Thay nội dung thẻ con của 1 dòng HHDVu (giữ nguyên phần còn lại); chưa có thẻ thì chèn trước thẻ `before`."""
    val = xml_escape(value)
    m = _xml_tag_re(tag).search(body)
    if m:
        p = m.group(1)
        return f"{body[:m.start()]}<{p}{tag}{m.group(2)}>{val}</{p}{tag}>{body[m.end():]}"
    mb = re.search(rf"<((?:[\w.-]+:)?){before}\b", body) if before else None
    if mb:
        return f"{body[:mb.start()]}<{mb.group(1)}{tag}>{val}</{mb.group(1)}{tag}>{body[mb.start():]}"
    return body


def _clone_note(note):
    return ("",) * (len(CLONE_HEADERS) - 4) + (note,)


def _xml_num(x):
    s = f"{x:.6f}".rstrip("0").rstrip(".")
    return "0" if s in ("", "-0") else s


def read_source(data_dir, src):
    """Đọc lại 1 nguồn XML theo tên iter_sources trả về ('a.zip!b.xml' = file trong ZIP, có thể lồng)."""
    parts = src.split("!")
    with open(os.path.join(data_dir, parts[0]), "rb") as f:
        data = f.read()
    for inner in parts[1:]:
        data = zipfile.ZipFile(io.BytesIO(data)).read(inner)
    return data


def _clone_name(src, used):
    parts = src.split("!")
    name = os.path.basename(parts[-1].replace("\\", "/"))
    if len(parts) > 1:  # XML trong ZIP: thêm tên ZIP để không trùng (nhiều ZIP cùng chứa 'invoice.xml')
        name = f"{os.path.splitext(os.path.basename(parts[-2].replace(chr(92), '/')))[0]}__{name}"
    stem, ext = os.path.splitext(name)
    ext = ext if ext.lower() == ".xml" else ".xml"
    cand, n = stem + ext, 1
    while cand.lower() in used:
        n += 1
        cand = f"{stem} ({n}){ext}"
    used.add(cand.lower())
    return cand


def _convert_xml_bytes(data, src, company, mapping, doi_ma, doi_ten, stamp, strict=True):
    """→ (bytes mới, các dòng nhật ký, trạng thái 'sua' | 'giu' | 'loi'). Chỉ sửa SLuong/DGia/DVTinh (+ MHHDVu/THHDVu)
    của dòng hàng có Mã master + Hệ số (+ Mã Vietda); thành tiền, thuế, tổng tiền giữ nguyên. strict: chỉ sửa theo cột
    'Bản sao XML' của anh_xa_hang_hoa.xlsx (ánh xạ chính xác / kế toán đã xác nhận)."""
    bom = data.startswith(b"\xef\xbb\xbf")
    try:
        text = (data[3:] if bom else data).decode("utf-8")
        invs = parse_invoices(data, src)
    except (UnicodeDecodeError, ET.ParseError) as e:
        return data, [_clone_note(f"Không đọc được ({type(e).__name__}) → chép nguyên file gốc")], "loi"
    lines = [(inv, ln) for inv in invs for ln in inv["lines"]]
    blocks = list(_HHDVU_RE.finditer(text))
    if len(blocks) != len(lines):
        return data, [_clone_note(f"Không khớp số dòng hàng ({len(blocks)} thẻ HHDVu / {len(lines)} dòng đọc được) "
                                  "→ chép nguyên file gốc, cần sửa tay")], "loi"
    out, pos, rows = [], 0, []
    for m, (inv, ln) in zip(blocks, lines):
        if ln["tchat"] not in ("1", "2"):
            continue
        direction = "RA" if inv["mst_ban"] == company else ("VAO" if inv["mst_mua"] == company else "")
        body = m.group(2)
        sl = parse_num(_xml_get(body, "SLuong"))
        if not direction or not sl:
            continue
        dg_raw = _xml_get(body, "DGia")
        dg = parse_num(dg_raw) if dg_raw and dg_raw.strip() else None
        ref = (inv["ky_hieu"], inv["so"], ln["stt"])
        row = mapping.get(item_key(direction, inv["mst_ban"], ln["ma"], ln["ten"], ln["dvt"]))
        mk = str(row[COL["Mã master"]] or "").strip() if row else ""
        factor = parse_num(row[COL["Hệ số quy đổi"]]) if row and row[COL["Hệ số quy đổi"]] not in (None, "") else 0.0
        vd_code = str(row[COL["Mã Vietda"]] or "").strip() if row else ""
        plan = (row[COL["Bản sao XML"]] or "") if row else ""
        do_unit = bool(mk) and factor > 0 and (not strict or plan.startswith("Đổi ĐVT"))
        do_vd = bool(vd_code) and (not strict or "mã Vietda" in plan)
        if not do_unit and not do_vd:
            why = ("Chưa có Mã master / Hệ số / Mã Vietda trong anh_xa_hang_hoa.xlsx" if not (mk and factor > 0) and not vd_code
                   else "Ánh xạ chưa chắc chắn — kiểm ở anh_xa_hang_hoa.xlsx rồi đánh x")
            rows.append((*ref, ln["ma"], "", ln["ten"], "", ln["dvt"], "", sl, None, dg, None, None, why + " → giữ nguyên dòng này"))
            continue
        new_unit = str(row[COL["ĐVT master"]] or "").strip() if do_unit else ""
        if do_vd:
            code, new_name = vd_code, str(row[COL["Tên Vietda"]] or "").strip()
        else:
            code, new_name = mk, str(row[COL["Tên master"]] or "").strip()
            if factor == 1 and norm_code(ln["ma"]) == norm_code(mk) and norm_unit(ln["dvt"]) == norm_unit(new_unit or ln["dvt"]):
                continue  # dòng đã ở mã + ĐVT master, không có mã Vietda → giữ nguyên (không đổi tên cho đẹp)
        nb = body
        new_sl, new_dg, f = sl, dg, None
        if do_unit:
            f = factor
            new_sl, new_dg = sl * factor, (dg / factor if dg is not None else None)
            nb = _xml_set(nb, "SLuong", _xml_num(new_sl))
            if new_dg is not None:
                nb = _xml_set(nb, "DGia", _xml_num(new_dg))
            if new_unit:
                nb = _xml_set(nb, "DVTinh", new_unit, before="SLuong")
        if doi_ma and code:
            nb = _xml_set(nb, "MHHDVu", code, before="THHDVu")
        if doi_ten and new_name:
            nb = _xml_set(nb, "THHDVu", new_name)
        if nb == body:
            continue
        out += [text[pos:m.start(2)], nb]
        pos = m.end(2)
        note = "" if do_unit or not mk else "Chỉ đổi mã Vietda — quy đổi ĐVT chờ xác nhận"
        if not strict and row[COL["Cần xem"]] and not row[COL["Đã kiểm (x)"]]:
            note = "Ánh xạ là GỢI Ý CHƯA KIỂM ('Cần xem' trong anh_xa_hang_hoa.xlsx)"
        rows.append((*ref, ln["ma"], code if doi_ma else ln["ma"], ln["ten"], new_name if doi_ten and new_name else ln["ten"],
                     ln["dvt"], new_unit or ln["dvt"], sl, new_sl, dg, new_dg, f, note))
    if not out:
        return data, rows, "giu"
    new_text = "".join(out) + text[pos:]
    nl = "\r\n" if "\r\n" in text[:500] else "\n"
    note = (f" Bản QUY ĐỔI ĐVT về ĐVT master KiotViet — doi_soat.py tạo {stamp} từ file gốc: {src}. Chỉ dùng tổng hợp "
            "nội bộ; KHÔNG phải hóa đơn gốc (chữ ký số không còn khớp nội dung). ").replace("--", "-")
    decl = re.match(r"\s*<\?xml[^>]*\?>", new_text)
    at = decl.end() if decl else 0
    new_text = f"{new_text[:at]}{nl if decl else ''}<!--{note}-->{'' if decl else nl}{new_text[at:]}"
    return (b"\xef\xbb\xbf" if bom else b"") + new_text.encode("utf-8"), rows, "sua"


def export_converted_xml(data_dir, invoices, company, mapping, cfg):
    """Bản sao XML hóa đơn vào/ra đã quy đổi ĐVT → CLONE_DIRS. File gốc KHÔNG đổi. → (dòng nhật ký, thống kê)."""
    files = {}
    for inv in invoices:
        files.setdefault(inv["file"], inv["direction"])
    out_dirs = {d: os.path.join(data_dir, name) for d, name in CLONE_DIRS.items()}
    for p in out_dirs.values():  # xóa bản sao của lần chạy trước (chỉ file .xml trong 2 thư mục này)
        os.makedirs(p, exist_ok=True)
        for fn in os.listdir(p):
            if fn.lower().endswith(".xml"):
                try:
                    os.remove(os.path.join(p, fn))
                except OSError:
                    pass
    doi_ma, doi_ten = cfg.get("xml_quy_doi_doi_ma", True), cfg.get("xml_quy_doi_doi_ten", True)
    stamp = datetime.now().strftime("%d/%m/%Y %H:%M")
    used = {d: set() for d in out_dirs}
    rows, stats = [], Counter()
    n, last = len(files), 0.0
    for i, (src, direction) in enumerate(files.items(), start=1):
        name = _clone_name(src, used[direction])
        try:
            data = read_source(data_dir, src)
            new_data, file_rows, status = _convert_xml_bytes(data, src, company, mapping, doi_ma, doi_ten, stamp,
                                                             cfg.get("xml_quy_doi_chi_khi_chac_chan", True))
            with open(os.path.join(out_dirs[direction], name), "wb") as f:
                f.write(new_data)
        except (OSError, KeyError, zipfile.BadZipFile) as e:
            file_rows, status = [_clone_note(f"Lỗi: {e} — không tạo được bản sao")], "loi"
        stats[status] += 1
        stats[f"dir_{direction}"] += 1
        stats["dong_sua"] += sum(1 for r in file_rows if isinstance(r[10], (int, float)))
        stats["dong_chua_anh_xa"] += sum(1 for r in file_rows if r[10] is None)
        label = "Vào" if direction == "VAO" else "Ra"
        rows += [(label, src, f"{CLONE_DIRS[direction]}\\{name}", *r) for r in file_rows]
        now = time.time()
        if _PROGRESS_HOOK is not None:
            if i == n or now - last >= 0.5:
                last = now
                _PROGRESS_HOOK("xml_quy_doi", i, n)
        elif i == n or now - last >= 0.1:
            last = now
            pct = i * 100 // n
            shown = name if len(name) <= 34 else name[:31] + "..."
            sys.stdout.write(f"\r  Sửa XML quy đổi [{'#' * (pct // 5)}{'-' * (20 - pct // 5)}] {i:,}/{n:,} ({pct}%) "
                             f"· đã sửa {stats['sua']:,} file · {shown:<34}")
            sys.stdout.flush()
    if n and _PROGRESS_HOOK is None:
        sys.stdout.write("\n")
    return rows, stats


# ============================================================ CHUONG TRINH CHINH

def parse_args():
    ap = argparse.ArgumentParser(description="Đối soát + quy đổi đơn vị hóa đơn đầu vào - đầu ra từ file XML")
    ap.add_argument("--du-lieu", help="Thư mục gốc chứa XML/ZIP + bảng quy đổi (mặc định: thư mục chứa script)")
    ap.add_argument("--mst", help="MST công ty (mặc định tự nhận)")
    ap.add_argument("--khong-dung", action="store_true", help="Không chờ Enter khi chạy xong")
    return ap.parse_args()


def read_xml(sources):
    """Bước 1 — đọc XML: hóa đơn + tờ khai đã nộp (01/GTGT, 05/KK-TNCN). sources = các cặp (tên, bytes) như iter_sources().
    Không ghi file. Dùng chung cho CLI và KeToanBackEnd."""
    invoices, read_errors, skipped, n_files = [], [], 0, 0
    submitted_list, submitted_tncn = [], []
    for src, data in sources:
        n_files += 1
        if n_files % 2000 == 0:
            log(f"  ... đã đọc {n_files:,} file")
        if b"HSoThueDTu" in data[:800]:
            try:
                tkx = parse_submitted_gtgt(data)
                tnx = parse_submitted_tncn(data) if tkx and tkx.get("_khac") else None
            except ET.ParseError:
                tkx = tnx = None
            if tkx and not tkx.get("_khac"):
                submitted_list.append(dict(tkx, _file=src, _raw=data))
                log(f"  Thấy tờ khai 01/GTGT đã nộp: {src} (kỳ {tkx.get('_ky')})")
            elif tnx:
                submitted_tncn.append(dict(tnx, _file=src, _raw=data))
                log(f"  Thấy tờ khai 05/KK-TNCN đã nộp: {src} (kỳ {tnx.get('_ky')})")
            continue
        try:
            found = parse_invoices(data, src)
        except ET.ParseError as e:
            read_errors.append((src, f"XML lỗi: {e}"))
            continue
        except Exception as e:
            read_errors.append((src, f"{type(e).__name__}: {e}"))
            continue
        if not found:
            skipped += 1
        invoices.extend(found)
    log(f"  {n_files:,} file XML, {len(invoices):,} hóa đơn, {skipped:,} file không phải hóa đơn, {len(read_errors)} file lỗi")
    return {"invoices": invoices, "read_errors": read_errors, "skipped": skipped, "n_files": n_files,
            "submitted_gtgt": submitted_list, "submitted_tncn": submitted_tncn}


def reconcile(invoices, cat, cfg, company="", existing=None, opening=None, vd=None):
    """Bước 2–7 — chiều HĐ, trùng / thay thế / điều chỉnh / kỳ tính thuế, kiểm số học, ghép mã master, quy đổi + FIFO,
    ngoại lệ. KHÔNG đọc / ghi file: ánh xạ đã sửa (existing, dạng load_mapping), tồn đầu kỳ (opening, dạng load_opening),
    danh mục Vietda (vd) do nơi gọi nạp sẵn. Trả dict các bảng → report_sheets(), tax_figures(), save_mapping()...
    Dùng chung cho CLI và KeToanBackEnd."""
    tol = float(cfg["sai_so_tien"])
    existing = existing or {}
    opening = opening or {}
    map_name = os.path.basename(cfg["file_anh_xa"])

    # ---------- 2. MST công ty + chiều HĐ
    company = (company or "").strip()
    if not company:
        company = Counter(i["mst_ban"] for i in invoices if i["mst_ban"]).most_common(1)[0][0]
        log(f"  MST công ty (tự nhận): {company}  — sai thì sửa CONFIG['mst_cong_ty'] ở đầu file doi_soat.py")
    unrelated = []
    for inv in invoices:
        if inv["mst_ban"] == company:
            inv["direction"] = "RA"
        elif inv["mst_mua"] == company:
            inv["direction"] = "VAO"
        else:
            unrelated.append(inv)
    invoices = [i for i in invoices if i["direction"]]

    # ---------- 3. Trùng, thay thế / điều chỉnh, kỳ tính thuế
    by_key, duplicates = {}, []
    for inv in invoices:
        prev = by_key.get(inv["key"])
        if prev is None:
            by_key[inv["key"]] = inv
        else:
            same = abs(prev["tong_tt"] - inv["tong_tt"]) <= tol and len(prev["lines"]) == len(inv["lines"])
            duplicates.append((inv, prev, same))
    invoices = list(by_key.values())
    missing_ref = []
    for inv in invoices:
        inv["eff_date"] = inv["ngay"]
    for inv in invoices:
        if inv["tc_hd"] in ("1", "2") and inv["lq_key"]:
            ref = by_key.get(inv["lq_key"])
            if ref is None:
                missing_ref.append(inv)
            if inv["tc_hd"] == "1":
                # HĐ thay thế tính vào kỳ của HĐ GỐC (kế toán: trừ HĐ thay thế cho kỳ trước, cộng HĐ lập kỳ sau thay cho kỳ này)
                orig_date = ref["ngay"] if ref is not None else inv["lq_ngay"]
                if orig_date:
                    inv["eff_date"] = orig_date
                if ref is not None:
                    ref["status"] = "Bị thay thế"
            elif ref is not None and not ref["status"]:
                ref["status"] = "Bị điều chỉnh"
    for inv in invoices:
        if not inv["status"]:
            inv["status"] = INVOICE_STATUS.get(inv["tc_hd"], "Gốc")
    d_from, d_to, ky_label, ky_short = resolve_period(invoices, cfg)
    log(f"  Kỳ tính thuế: {ky_label} ({d_from:%d/%m/%Y} → {d_to:%d/%m/%Y}) — HĐ ngoài kỳ không tính")
    out_of_period = []
    for inv in invoices:
        ed = inv["eff_date"]
        inv["in_period"] = ed is None or d_from <= ed <= d_to
        own_in = inv["ngay"] is None or d_from <= inv["ngay"] <= d_to
        if inv["tc_hd"] == "1" and own_in and not inv["in_period"]:
            inv["period_note"] = f"Thay thế HĐ ngày {ed:%d/%m/%Y} (kỳ khác) → TRỪ khỏi kỳ này"
        elif inv["tc_hd"] == "1" and not own_in and inv["in_period"]:
            inv["period_note"] = f"Lập {inv['ngay']:%d/%m/%Y} thay thế HĐ của kỳ này → CỘNG vào kỳ này"
        elif inv["tc_hd"] == "1" and inv["lq_ngay"] is None and by_key.get(inv["lq_key"]) is None:
            inv["period_note"] = "Không rõ ngày HĐ gốc — tạm tính theo ngày lập, cần kiểm"
        if not inv["in_period"] and inv["status"] != "Bị thay thế":
            out_of_period.append(inv)
    active = [i for i in invoices if i["in_period"] and i["status"] != "Bị thay thế"]
    active_ids = {id(i) for i in active}

    # HĐ thay thế / điều chỉnh: danh sách cho kế toán kiểm + số trừ/cộng do khác kỳ
    repl_rows = []
    repl_adj = defaultdict(lambda: [0, 0.0, 0.0])
    for inv in invoices:
        if inv["tc_hd"] not in ("1", "2") and inv["status"] not in ("Bị thay thế", "Bị điều chỉnh"):
            continue
        counted = id(inv) in active_ids
        note = inv["period_note"] or {
            "Thay thế": "Thay thế HĐ cùng kỳ → tính HĐ thay thế, bỏ HĐ bị thay thế",
            "Điều chỉnh": "Tính vào kỳ lập HĐ điều chỉnh — kiểm dấu (+/−) số tiền điều chỉnh",
            "Bị thay thế": "Có HĐ thay thế trong dữ liệu → không tính HĐ này",
            "Bị điều chỉnh": "Vẫn tính; phần điều chỉnh nằm ở HĐ điều chỉnh",
        }.get(inv["status"], "")
        if not inv["in_period"] and inv["status"] != "Bị thay thế" and not inv["period_note"]:
            note = "Ngoài kỳ tính thuế → không tính"
        lq_parts = inv["lq_key"].split("|") if inv["lq_key"] else []
        repl_rows.append((_dir_label(inv), inv["status"], inv["ky_hieu"], inv["so"], inv["ngay"],
                          f"{lq_parts[1]}-{lq_parts[2]}" if len(lq_parts) == 3 else "", inv["lq_ngay"], inv["eff_date"],
                          "Có" if counted else "Không", inv["tong_chua_thue"], inv["tong_thue"], note, inv["file"]))
        if inv["tc_hd"] == "1" and inv["period_note"]:
            kind = "tru" if "TRỪ" in inv["period_note"] else ("cong" if "CỘNG" in inv["period_note"] else "")
            if kind:
                parts = invoice_rate_parts(inv)
                a = repl_adj[f"{inv['direction']}_{kind}"]
                a[0] += 1
                a[1] += sum(p[3] for p in parts)
                a[2] += sum(p[4] for p in parts)
    if repl_adj:
        for k, (n, b, t) in sorted(repl_adj.items()):
            d, kind = k.split("_")
            log(f"  HĐ {'ra' if d == 'RA' else 'vào'} thay thế khác kỳ — {'TRỪ' if kind == 'tru' else 'CỘNG'}: {n} HĐ, {b:,.0f} / thuế {t:,.0f}")

    # ---------- 4. Kiểm tra số học
    n_arith = 0
    for inv in active:
        inv["issues"] = check_arithmetic(inv, tol)
        n_arith += 1 if inv["issues"] else 0
    n_round = sum(i["rounding_lines"] for i in active)
    log(f"  Hóa đơn tính: {len(active):,} (vào {sum(1 for i in active if i['direction'] == 'VAO'):,}, "
        f"ra {sum(1 for i in active if i['direction'] == 'RA'):,}); lệch số học: {n_arith:,}; "
        f"dòng lệch do làm tròn đơn giá: {n_round:,}")

    # ---------- 5. Dòng hàng -> mặt hàng -> mã master
    stock_lines, items = [], {}
    for inv in active:
        goods = [ln for ln in inv["lines"] if ln["tchat"] == "1"]
        disc = sum(abs(ln["thtien"]) for ln in inv["lines"] if ln["tchat"] == "3")
        base_sum = sum(ln["thtien"] for ln in goods)
        for ln in inv["lines"]:
            if ln["tchat"] not in ("1", "2") or not ln["sl"]:
                continue
            net = 0.0
            if ln["tchat"] == "1":
                net = ln["thtien"] - (disc * ln["thtien"] / base_sum if disc and base_sum else 0.0)
            k = item_key(inv["direction"], inv["mst_ban"], ln["ma"], ln["ten"], ln["dvt"])
            it = items.get(k)
            if it is None:
                it = items[k] = {"direction": inv["direction"], "mst": inv["mst_ban"], "ten_ban": inv["ten_ban"],
                                 "ma": ln["ma"], "ten": ln["ten"], "dvt": ln["dvt"], "lines": 0, "qty": 0.0}
            it["lines"] += 1
            it["qty"] += ln["sl"]
            stock_lines.append({"inv": inv, "ln": ln, "key": k, "net": net, "km": ln["tchat"] == "2"})

    if vd is not None:
        log(f"  Danh mục Vietda: {os.path.basename(vd.path)} — {len(vd.rows):,} mã hàng")
    mapping = build_mapping(items, existing, cat, float(cfg["nguong_tin_cay_goi_y"]), vd)
    if vd is not None:
        n_vd = sum(1 for r in mapping.values() if r[COL["Mã Vietda"]])
        log(f"  Mã Vietda: {n_vd:,}/{len(mapping):,} mặt hàng — "
            + ", ".join(f"{k}: {v:,}" for k, v in Counter(r[COL['Cách ghép Vietda']] or 'Chưa có' for r in mapping.values()).most_common()))
    method_count = Counter(r[COL["Cách ghép"]] or "Chưa ghép" for r in mapping.values())
    n_review = sum(1 for r in mapping.values() if r[COL["Cần xem"]])
    log(f"  Mặt hàng trên HĐ: {len(mapping):,} — " + ", ".join(f"{k}: {v:,}" for k, v in method_count.most_common())
        + f"; cần xem: {n_review:,}")

    # Nhóm = mã master
    group_info = {}
    for row in mapping.values():
        mk = row[COL["Mã master"]]
        if mk:
            group_info.setdefault(norm_code(mk), (str(mk).strip(), row[COL["Tên master"]], row[COL["ĐVT master"]]))

    # ---------- 6. Quy đổi + FIFO theo mã master
    lag = timedelta(days=int(cfg["so_ngay_hd_vao_lap_sau"]))
    below_pct = float(cfg["nguong_duoi_gia_von_phan_tram"]) / 100
    lots_by_group, outs_by_group = defaultdict(list), defaultdict(list)
    unmapped = defaultdict(lambda: {"lines": 0, "qty": 0.0, "amount": 0.0})
    converted_rows = []
    for sl in stock_lines:
        row = mapping[sl["key"]]
        inv, ln = sl["inv"], sl["ln"]
        mk = row[COL["Mã master"]]
        factor = parse_num(row[COL["Hệ số quy đổi"]]) if row[COL["Hệ số quy đổi"]] not in (None, "") else 0.0
        q = ln["sl"] * factor if mk and factor > 0 else None
        converted_rows.append((
            _dir_label(inv), inv["ky_hieu"], inv["so"], inv["ngay"], inv["mst_ban"], inv["ten_ban"], inv["mst_mua"],
            ln["stt"], ln["ma"], ln["ten"], ln["dvt"], ln["sl"], ln["dgia"], ln["thtien"], ln["rate_label"],
            "Khuyến mại" if sl["km"] else "", mk or "", row[COL["Tên master"]], row[COL["ĐVT master"]],
            row[COL["Mã Vietda"]] or "", factor or None, q, (sl["net"] / q) if q else None,
            row[COL["Cách ghép"]] or "Chưa ghép", inv["file"]))
        if q is None:
            u = unmapped[sl["key"]]
            u["lines"] += 1
            u["qty"] += ln["sl"]
            u["amount"] += sl["net"]
            continue
        gid = norm_code(mk)
        rec = {"date": inv["ngay"] or date(1900, 1, 1), "inv": inv, "ln": ln, "qty": q, "net": sl["net"],
               "km": sl["km"], "factor": factor}
        if inv["direction"] == "VAO":
            if q > 0:
                rec.update(remaining=q, unit_cost=sl["net"] / q, alloc=[])
                lots_by_group[gid].append(rec)
            else:
                rec.update(kind="giam_vao", qty=-q)
                outs_by_group[gid].append(rec)
        else:
            rec["kind"] = "ban" if q > 0 else "tra_lai"
            if q < 0:
                rec["qty"] = -q
            outs_by_group[gid].append(rec)

    nxt_rows, trace_rows, lot_rows, below_rows, shortage_rows, exceptions = [], [], [], [], [], []
    for gid in sorted(set(lots_by_group) | set(outs_by_group) | set(opening), key=lambda g: str(group_info.get(g, (g, g, ""))[1])):
        gcode, gname, gunit = group_info.get(gid, (gid, gid, ""))
        label = f"{gcode} {gname}"
        lots = sorted(lots_by_group.get(gid, []), key=lambda r: (r["date"], r["inv"]["ky_hieu"], r["inv"]["so"]))
        op_qty, op_cost = opening.get(gid, (0.0, 0.0))
        if op_qty:
            lots.insert(0, {"date": date(1900, 1, 1), "inv": None, "ln": None, "qty": op_qty, "remaining": op_qty,
                            "unit_cost": op_cost, "alloc": [], "net": op_qty * op_cost, "km": False})
        outs = sorted(outs_by_group.get(gid, []),
                      key=lambda r: (r["date"], {"tra_lai": 0, "ban": 1, "giam_vao": 2}[r["kind"]],
                                     r["inv"]["ky_hieu"], r["inv"]["so"]))
        lot_dates = [l["date"] for l in lots]
        p = 0
        sold_qty = revenue = cogs = short_total = loss_below = 0.0
        out_rates, in_rates = Counter(), Counter()
        n_below = 0
        avg_cost = cfg.get("phuong_phap_gia_von", "fifo") == "binh_quan"
        last_cost = lots[0]["unit_cost"] if lots else 0.0
        stock_qty = stock_val = 0.0          # bình quân: tồn đã nhận (lô có ngày lập ≤ ngày bán + lag) và giá trị
        for o in outs:
            if avg_cost:
                # nhận các lô nhập tới ngày bán (+ lag) → giá bình quân tính lại sau MỖI lần nhập
                while p < len(lots) and lots[p]["date"] <= o["date"] + lag:
                    stock_qty += lots[p]["qty"]
                    stock_val += lots[p]["qty"] * lots[p]["unit_cost"]
                    p += 1
                cur = stock_val / stock_qty if stock_qty > 1e-9 else last_cost
                last_cost = cur
                if o["kind"] == "tra_lai":       # khách trả lại: nhập lại kho theo giá bình quân hiện tại
                    stock_qty += o["qty"]
                    stock_val += o["qty"] * cur
                    sold_qty -= o["qty"]
                    revenue += o["net"]
                    cogs -= o["qty"] * cur
                    continue
                need = o["qty"]
                got_qty = min(need, max(stock_qty, 0.0))
                got_cost = got_qty * cur
                if got_qty > 0:
                    stock_qty -= got_qty
                    stock_val -= got_cost
                    if stock_qty <= 1e-9:
                        stock_qty = stock_val = 0.0
                need -= got_qty
            else:
                if o["kind"] == "tra_lai":
                    ret = {"date": o["date"], "inv": o["inv"], "ln": o["ln"], "qty": o["qty"], "remaining": o["qty"],
                           "unit_cost": last_cost, "alloc": [], "net": 0.0, "km": False, "is_return": True}
                    idx = bisect.bisect_right(lot_dates, o["date"])
                    lots.insert(idx, ret)
                    lot_dates.insert(idx, o["date"])
                    p = min(p, idx)
                    sold_qty -= o["qty"]
                    revenue += o["net"]
                    cogs -= o["qty"] * last_cost
                    continue
                need = o["qty"]
                got_qty = got_cost = 0.0
                while need > 1e-9 and p < len(lots) and lots[p]["date"] <= o["date"] + lag:
                    lot = lots[p]
                    take = min(need, lot["remaining"])
                    if take > 0:
                        lot["remaining"] -= take
                        need -= take
                        got_qty += take
                        got_cost += take * lot["unit_cost"]
                        last_cost = lot["unit_cost"]
                        if o["kind"] == "ban":
                            lot["alloc"].append((o, take))
                            trace_rows.append((gcode, gname, *_inv_ref(lot["inv"]), _ln_name(lot), lot["qty"],
                                               o["inv"]["ky_hieu"], o["inv"]["so"], o["inv"]["ngay"], o["ln"]["ten"],
                                               take, lot["unit_cost"], (o["net"] / o["qty"]) if o["qty"] and not o["km"] else 0.0))
                    if lot["remaining"] <= 1e-9:
                        p += 1
            if o["kind"] == "giam_vao":
                if need > 1e-9:
                    shortage_rows.append((gcode, gname, "Giảm/trả hàng NCC vượt tồn", o["inv"]["ky_hieu"], o["inv"]["so"],
                                          o["date"], o["ln"]["ten"], o["qty"], need))
                continue
            sold_qty += o["qty"]
            revenue += o["net"]
            cogs += got_cost
            if o["ln"]["rate_label"]:
                out_rates[o["ln"]["rate_label"]] += 1
            if need > 1e-9:
                short_total += need
                shortage_rows.append((gcode, gname, "Thiếu HĐ đầu vào tại ngày bán", o["inv"]["ky_hieu"], o["inv"]["so"],
                                      o["date"], o["ln"]["ten"], o["qty"], need))
            if not o["km"] and got_qty > 1e-9 and o["qty"] > 0:
                unit_cost = got_cost / got_qty
                unit_price = o["net"] / o["qty"]
                if unit_price < unit_cost * (1 - below_pct) - 1e-6:
                    n_below += 1
                    loss = (unit_cost - unit_price) * o["qty"]
                    loss_below += loss
                    below_rows.append((gcode, gname, o["inv"]["ky_hieu"], o["inv"]["so"], o["date"], o["ln"]["ten"],
                                       o["ln"]["dvt"], o["qty"], unit_price, unit_cost,
                                       unit_price / unit_cost - 1 if unit_cost else 0, loss))

        real_lots = [l for l in lots if l["inv"] is not None and not l.get("is_return")]
        in_qty = sum(l["qty"] for l in real_lots)
        in_amount = sum(l["net"] for l in real_lots)
        for l in real_lots:
            if l["ln"]["rate_label"]:
                in_rates[l["ln"]["rate_label"]] += 1
        if avg_cost:
            while p < len(lots):             # lô nhập sau lần bán cuối
                stock_qty += lots[p]["qty"]
                stock_val += lots[p]["qty"] * lots[p]["unit_cost"]
                p += 1
        remaining_later = stock_qty if avg_cost else sum(l["remaining"] for l in lots)
        covered_later = min(short_total, remaining_later)
        uncovered = short_total - covered_later
        closing = op_qty + in_qty - sold_qty - sum(o["qty"] for o in outs if o["kind"] == "giam_vao")
        notes = []
        if uncovered > 1e-9:
            exceptions.append(("Cao", "Bán vượt số lượng HĐ đầu vào", "Ra", label, "", "", None, uncovered,
                               f"Bán {sold_qty:,.3g} > tồn đầu {op_qty:,.3g} + nhập {in_qty:,.3g}; thiếu {uncovered:,.3g} {gunit}"))
            notes.append("Bán vượt")
        if covered_later > 1e-9:
            exceptions.append(("Trung bình", "HĐ đầu vào lập sau ngày bán", "Ra", label, "", "", None, covered_later,
                               f"{covered_later:,.3g} {gunit} bán trước ngày lập HĐ vào (xem sheet 'Thieu HD vao')"))
            notes.append("HĐ vào lập sau")
        if n_below:
            exceptions.append(("Cao", "Bán dưới giá vốn", "Ra", label, "", "", None, loss_below,
                               f"{n_below:,} dòng bán dưới giá vốn, tổng chênh {loss_below:,.0f}đ (xem sheet 'Ban duoi gia von')"))
            notes.append("Dưới giá vốn")
        if len(out_rates) > 1:
            exceptions.append(("Trung bình", "Nhiều thuế suất đầu ra cho cùng mặt hàng", "Ra", label, "", "", None, None,
                               ", ".join(f"{k}: {v} dòng" for k, v in out_rates.most_common())))
            notes.append("Nhiều thuế suất ra")
        if out_rates and in_rates and set(out_rates) != set(in_rates):
            exceptions.append(("Trung bình", "Thuế suất đầu ra khác đầu vào", "", label, "", "", None, None,
                               f"Ra: {', '.join(out_rates)} — Vào: {', '.join(in_rates)}"))
            notes.append("Thuế suất ra ≠ vào")
        if real_lots and not outs:
            notes.append("Chưa bán")
        margin = (revenue - cogs) / revenue if revenue else None
        nxt_rows.append((gcode, gname, gunit, op_qty, in_qty, in_amount, in_amount / in_qty if in_qty else None,
                         sold_qty, revenue, revenue / sold_qty if sold_qty else None, cogs, revenue - cogs, margin,
                         closing, uncovered, covered_later, ", ".join(out_rates), ", ".join(in_rates),
                         len({l["inv"]["key"] for l in real_lots}),
                         len({o["inv"]["key"] for o in outs if o["kind"] == "ban"}), "; ".join(notes)))
        for l in real_lots:
            if avg_cost:                     # bình quân: không phân bổ theo lô → không có SL đã bán / còn lại từng lô
                lot_rows.append((gcode, gname, *_inv_ref(l["inv"]), _ln_name(l), l["ln"]["dvt"], l["ln"]["sl"], l["factor"],
                                 l["qty"], l["unit_cost"], None, None, None, "(giá vốn bình quân — không phân bổ theo lô)"))
                continue
            outs_list = sorted({f"{o['inv']['ky_hieu']}-{o['inv']['so']}" for o, _ in l["alloc"]})
            lot_rows.append((gcode, gname, *_inv_ref(l["inv"]), _ln_name(l), l["ln"]["dvt"], l["ln"]["sl"], l["factor"],
                             l["qty"], l["unit_cost"], l["qty"] - l["remaining"], l["remaining"], len(outs_list),
                             ", ".join(outs_list[:30]) + (" …" if len(outs_list) > 30 else "")))

    # ---------- 7. Ngoại lệ cấp hóa đơn / ánh xạ
    for inv in active:
        for kind, diff, desc in inv["issues"]:
            sev = "Thấp" if abs(diff) < 1.5 else ("Trung bình" if abs(diff) < 1000 else "Cao")
            exceptions.append((sev, "Lệch số học: " + kind, _dir_label(inv), "", inv["ky_hieu"], inv["so"], inv["ngay"], diff, desc))
    for inv, prev, same in duplicates:
        if not same:
            exceptions.append(("Cao", "Trùng số HĐ nhưng khác nội dung", _dir_label(inv), "", inv["ky_hieu"], inv["so"],
                               inv["ngay"], None, f"{prev['file']} ↔ {inv['file']}"))
    for inv in missing_ref:
        exceptions.append(("Thấp", f"HĐ {INVOICE_STATUS.get(inv['tc_hd'], '')} — không có HĐ gốc trong dữ liệu",
                           _dir_label(inv), "", inv["ky_hieu"], inv["so"], inv["ngay"], None, f"HĐ gốc: {inv['lq_key']}"))
    for k, u in unmapped.items():
        it = items[k]
        exceptions.append(("Trung bình", "Mặt hàng chưa ánh xạ", "Vào" if it["direction"] == "VAO" else "Ra",
                           f"{it['ma']} {it['ten']}".strip(), "", "", None, u["amount"],
                           f"{u['lines']:,} dòng, SL {u['qty']:,.3g} {it['dvt']} — điền Mã master + Hệ số trong {map_name}"))
    for row in mapping.values():
        if row[COL["Cần xem"]] and row[COL["Mã master"]]:
            exceptions.append(("Thấp", "Ánh xạ cần xem", row[COL["Chiều"]], f"{row[COL['Mã hàng (HĐ)']] or ''} {row[COL['Tên hàng (HĐ)']]}".strip(),
                               "", "", None, None, row[COL["Cần xem"]]))
    if unrelated:
        exceptions.append(("Thấp", "HĐ không liên quan MST công ty (đã bỏ qua)", "", "", "", "", None, len(unrelated),
                           f"{len(unrelated):,} HĐ — xem sheet 'HD bo qua'"))
    sev_order = {"Cao": 0, "Trung bình": 1, "Thấp": 2}
    exceptions.sort(key=lambda e: (sev_order[e[0]], e[1], str(e[3]), str(e[4]), e[5] or 0))

    return {"company": company, "invoices": invoices, "active": active, "unrelated": unrelated, "duplicates": duplicates,
            "out_of_period": out_of_period, "repl_rows": repl_rows, "repl_adj": repl_adj,
            "period": (d_from, d_to, ky_label, ky_short), "n_arith": n_arith, "n_round": n_round,
            "items": items, "mapping": mapping, "method_count": method_count, "n_review": n_review,
            "group_info": group_info, "opening": opening, "unmapped": unmapped, "converted_rows": converted_rows,
            "nxt_rows": nxt_rows, "trace_rows": trace_rows, "lot_rows": lot_rows, "below_rows": below_rows,
            "shortage_rows": shortage_rows, "exceptions": exceptions, "sev_count": Counter(e[0] for e in exceptions),
            "cost_method": cfg.get("phuong_phap_gia_von", "fifo")}


def summary_rows(res, read, cfg, data_dir, cat_label, clone_stats):
    """Sheet 'Tong quan': (mục, giá trị). clone_stats = thống kê bản sao XML quy đổi (Counter rỗng nếu không xuất)."""
    active, invoices, exceptions = res["active"], res["invoices"], res["exceptions"]
    dates = [i["ngay"] for i in active if i["ngay"]]
    sev_count = res["sev_count"]
    type_count = Counter(e[1].split(":")[0] if e[1].startswith("Lệch số học") else e[1] for e in exceptions)
    return [
        ("Thời điểm chạy", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Phiên bản script", VERSION),
        ("Thư mục", data_dir),
        ("Bảng quy đổi KiotViet", cat_label),
        ("MST công ty", res["company"]),
        ("Khoảng ngày HĐ", f"{min(dates):%d/%m/%Y} → {max(dates):%d/%m/%Y}" if dates else ""),
        ("Số file XML đã đọc", read["n_files"]),
        ("File không phải hóa đơn / lỗi", f"{read['skipped']} / {len(read['read_errors'])}"),
        ("HĐ đầu vào được tính", sum(1 for i in active if i["direction"] == "VAO")),
        ("HĐ đầu ra được tính", sum(1 for i in active if i["direction"] == "RA")),
        ("HĐ bị thay thế (loại khỏi tính)", sum(1 for i in invoices if i["status"] == "Bị thay thế")),
        ("HĐ ngoài kỳ (loại)", len(res["out_of_period"])),
        ("File trùng HĐ (bỏ bản sau)", len(res["duplicates"])),
        ("HĐ không liên quan MST (bỏ qua)", len(res["unrelated"])),
        ("HĐ có lệch số học", res["n_arith"]),
        ("Dòng lệch do làm tròn đơn giá (không tính lỗi)", res["n_round"]),
        ("Mặt hàng trên HĐ", len(res["mapping"])),
        *[(f"  ghép bằng: {k}", v) for k, v in res["method_count"].most_common()],
        ("Mặt hàng cần xem", res["n_review"]),
        ("Mặt hàng chưa ánh xạ (có phát sinh)", len(res["unmapped"])),
        ("Mã master đối soát", len(res["nxt_rows"])),
        ("Có tồn đầu kỳ", "Có" if res["opening"] else "KHÔNG — mọi hàng bán ra phải có HĐ vào trong dữ liệu; điền ton_dau_ky.xlsx để chính xác hơn"),
        ("Ngoại lệ — Cao / Trung bình / Thấp", f"{sev_count['Cao']:,} / {sev_count['Trung bình']:,} / {sev_count['Thấp']:,}"),
        ("", ""),
        ("Ngoại lệ theo loại", ""),
        *[(f"  {k}", v) for k, v in type_count.most_common()],
        ("", ""),
        ("XML quy đổi (bản sao)", f"{clone_stats['sua']:,} file sửa / {clone_stats['dong_sua']:,} dòng hàng · "
                                  f"{clone_stats['giu']:,} file chép nguyên · {clone_stats['loi']:,} file lỗi · "
                                  f"{clone_stats['dong_chua_anh_xa']:,} dòng giữ nguyên (chưa ánh xạ / chưa chắc chắn)"
         if cfg.get("xuat_xml_quy_doi", True) else "Tắt (CONFIG xuat_xml_quy_doi)"),
        ("Lưu ý", f"File XML gốc KHÔNG bị sửa. Bản sao đã quy đổi ĐVT nằm ở '{CLONE_DIRS['VAO']}' / '{CLONE_DIRS['RA']}'; "
                  "số quy đổi cũng có ở sheet 'Danh sach san pham' và 'Dong hang quy doi'."),
        ("Lưu ý", "HĐ bị HỦY không nhận ra được từ XML — loại file của HĐ đã hủy khỏi thư mục dữ liệu."),
        ("Lưu ý", "Giá vốn/giá bán tính theo tiền CHƯA thuế, đã trừ chiết khấu thương mại phân bổ theo dòng, quy về ĐVT master."),
    ]


Sheet = namedtuple("Sheet", "title headers rows widths formats wrap")


def report_sheets(res, read, summary, clone_rows=None):
    """Các sheet của doi_soat_<ngày>.xlsx theo thứ tự — web dùng lại làm các tab (cùng tiêu đề cột / định dạng).
    clone_rows = None → bỏ sheet 'XML quy doi'."""
    active, invoices, mapping = res["active"], res["invoices"], res["mapping"]
    product_rows = sorted(mapping.values(), key=lambda r: (r[COL["Chiều"]] != "Vào", str(r[COL["Tên master"]] or "~"), str(r[COL["Tên hàng (HĐ)"]])))
    sheets = [
        Sheet("Huong dan", ["Sheet", "Mục đích", "Cách đọc / cột chính", "Khi nào cần làm gì"], REPORT_GUIDE,
              [22, 55, 80, 55], None, True),
        Sheet("Tong quan", ["Mục", "Giá trị"], summary, [46, 110], {1: "#,##0"}, False),
        Sheet("Danh sach san pham", MAP_HEADERS[:COL["Đã kiểm (x)"]],
              [r[:COL["Đã kiểm (x)"]] for r in product_rows],
              [6, 13, 26, 15, 45, 8, 7, 10, 15, 38, 9, 8, 12, 12, 38, 9, 18, 18, 8, 50, 26, 50],
              {7: QTY, 11: QTY, 12: QTY, COL["Độ tin cậy"]: "0.00"}, False),
        Sheet("Doi soat theo master", ["Mã master", "Tên master", "ĐVT master", "Tồn đầu", "Nhập SL", "Nhập tiền",
                                       "Giá nhập BQ/ĐVT", "Xuất SL", "Doanh thu", "Giá bán BQ/ĐVT",
                                       "Giá vốn bình quân" if res.get("cost_method") == "binh_quan" else "Giá vốn FIFO",
                                       "Lãi gộp", "Biên LN", "Tồn cuối (theo HĐ)", "Bán vượt HĐ vào (SL)",
                                       "HĐ vào lập sau (SL)", "Thuế suất ra", "Thuế suất vào", "Số HĐ vào", "Số HĐ ra",
                                       "Ghi chú"], res["nxt_rows"],
              [15, 40, 9, 9, 10, 14, 12, 10, 14, 12, 14, 13, 8, 12, 12, 12, 10, 10, 8, 8, 30],
              {3: QTY, 4: QTY, 5: MONEY, 6: MONEY2, 7: QTY, 8: MONEY, 9: MONEY2, 10: MONEY, 11: MONEY, 12: PCT,
               13: QTY, 14: QTY, 15: QTY}, False),
        Sheet("Ngoai le", ["Mức độ", "Loại", "Chiều", "Mã/Tên hàng", "Ký hiệu", "Số HĐ", "Ngày",
                           "Giá trị / SL lệch", "Mô tả"], res["exceptions"], [10, 42, 6, 45, 11, 10, 11, 14, 90], {7: MONEY2}, False),
        Sheet("HD thay the - dieu chinh", ["Chiều", "Loại", "Ký hiệu", "Số HĐ", "Ngày lập", "HĐ liên quan",
                                           "Ngày HĐ liên quan", "Ngày tính kỳ", "Tính vào kỳ này", "Tổng chưa thuế",
                                           "Tổng thuế", "Ghi chú", "File"],
              sorted(res["repl_rows"], key=lambda r: (r[0], r[4] or date(1900, 1, 1), r[2], r[3])),
              [6, 13, 11, 9, 11, 16, 11, 11, 9, 15, 13, 60, 45], {9: MONEY, 10: MONEY}, False),
        Sheet("Dong hang quy doi", ["Chiều", "Ký hiệu", "Số HĐ", "Ngày", "MST bán", "Tên bán", "MST mua", "STT",
                                    "Mã hàng (HĐ)", "Tên hàng (HĐ)", "ĐVT (HĐ)", "SL (HĐ)", "Đơn giá (HĐ)",
                                    "Thành tiền (HĐ)", "Thuế suất", "Loại dòng", "Mã master", "Tên master",
                                    "ĐVT master", "Mã Vietda", "Hệ số", "SL quy đổi", "Đơn giá quy đổi (chưa thuế, sau CK)",
                                    "Cách ghép", "File"], res["converted_rows"],
              [6, 11, 9, 11, 13, 28, 13, 5, 15, 40, 8, 8, 11, 12, 8, 10, 15, 36, 9, 12, 7, 10, 14, 16, 40],
              {11: QTY, 12: MONEY2, 13: MONEY, 20: QTY, 21: QTY, 22: MONEY2}, False),
    ]
    if clone_rows is not None:
        sheets.append(Sheet("XML quy doi", CLONE_HEADERS, clone_rows,
                            [6, 40, 40, 11, 9, 5, 15, 15, 36, 36, 9, 9, 9, 10, 12, 12, 7, 60],
                            {12: QTY, 13: QTY, 14: MONEY2, 15: MONEY2, 16: QTY}, False))
    sheets += [
        Sheet("Truy vet HD vao", ["Mã master", "Tên master", "Ký hiệu vào", "Số HĐ vào", "Ngày vào", "MST NCC",
                                  "Tên NCC", "Tên hàng vào", "ĐVT vào", "SL trên HĐ", "Hệ số", "SL quy đổi",
                                  "Giá vốn/ĐVT master", "Đã bán", "Còn lại", "Số HĐ ra", "Các HĐ ra (tối đa 30)"], res["lot_rows"],
              [15, 30, 11, 10, 11, 13, 28, 40, 8, 9, 7, 10, 12, 10, 10, 8, 80],
              {9: QTY, 10: QTY, 11: QTY, 12: MONEY2, 13: QTY, 14: QTY}, False),
        Sheet("Truy vet chi tiet", ["Mã master", "Tên master", "Ký hiệu vào", "Số HĐ vào", "Ngày vào", "MST NCC",
                                    "Tên NCC", "Tên hàng vào", "SL lô (master)", "Ký hiệu ra", "Số HĐ ra", "Ngày ra",
                                    "Tên hàng ra", "SL phân bổ", "Giá vốn/ĐVT", "Giá bán/ĐVT"], res["trace_rows"],
              [15, 30, 11, 10, 11, 13, 28, 36, 10, 11, 10, 11, 36, 10, 12, 12],
              {8: QTY, 13: QTY, 14: MONEY2, 15: MONEY2}, False),
        Sheet("Ban duoi gia von", ["Mã master", "Tên master", "Ký hiệu", "Số HĐ", "Ngày", "Tên hàng", "ĐVT (HĐ)",
                                   "SL master", "Giá bán/ĐVT master", "Giá vốn/ĐVT master", "Chênh %", "Thiệt (đ)"],
              res["below_rows"], [15, 30, 11, 10, 11, 40, 8, 10, 12, 12, 9, 12],
              {7: QTY, 8: MONEY2, 9: MONEY2, 10: PCT, 11: MONEY}, False),
        Sheet("Thieu HD vao", ["Mã master", "Tên master", "Loại", "Ký hiệu", "Số HĐ", "Ngày", "Tên hàng",
                               "SL dòng (master)", "SL thiếu"], res["shortage_rows"], [15, 30, 28, 11, 10, 11, 40, 12, 10],
              {7: QTY, 8: QTY}, False),
        Sheet("Lech so hoc HD", ["Chiều", "Ký hiệu", "Số HĐ", "Ngày", "MST bán", "Tên bán", "Loại lệch",
                                 "Chênh (đ)", "Chi tiết", "File"],
              [(_dir_label(i), i["ky_hieu"], i["so"], i["ngay"], i["mst_ban"], i["ten_ban"], k, d, desc, i["file"])
               for i in active for k, d, desc in i["issues"]],
              [6, 11, 10, 11, 13, 30, 42, 10, 70, 50], {7: MONEY2}, False),
        Sheet("Thue theo thue suat", ["Chiều", "Thuế suất", "Số HĐ", "Tiền chưa thuế", "Tiền thuế"], _tax_summary(active),
              [8, 10, 10, 18, 16], {3: MONEY, 4: MONEY}, False),
    ]
    inv_headers = ["Ký hiệu", "Số HĐ", "Ngày", "MST bán", "Tên bán", "MST mua", "Tên mua", "Tổng chưa thuế",
                   "Tổng thuế", "Tổng thanh toán", "Trạng thái", "Số lỗi số học", "Dòng làm tròn ĐG", "Có mã CQT", "File"]
    inv_widths = [11, 10, 11, 13, 32, 13, 32, 15, 13, 15, 13, 8, 8, 8, 50]
    inv_fmt = {7: MONEY, 8: MONEY, 9: MONEY}
    for title, direction in (("HD dau vao", "VAO"), ("HD dau ra", "RA")):
        rows = [(_r(i) + (i["status"] + ("" if i["in_period"] else " · ngoài kỳ"), len(i["issues"]), i["rounding_lines"],
                          "Có" if i["mccqt"] else "", i["file"]))
                for i in sorted(invoices, key=lambda x: (x["ngay"] or date(1900, 1, 1), x["ky_hieu"], x["so"]))
                if i["direction"] == direction]
        sheets.append(Sheet(title, inv_headers, rows, inv_widths, inv_fmt, False))
    sheets += [
        Sheet("HD bo qua", inv_headers[:10] + ["File"],
              [_r(i) + (i["file"],) for i in res["unrelated"]]
              + [(*_r(i), f"Trùng với {p['file']}: {i['file']}") for i, p, _ in res["duplicates"]],
              inv_widths[:10] + [80], inv_fmt, False),
        Sheet("File loi", ["File", "Lỗi"], read["read_errors"], [80, 80], None, False),
    ]
    return sheets


def write_report(path, sheets):
    rep = Report()
    for s in sheets:
        rep.sheet(s.title, s.headers, s.rows, s.widths, s.formats, wrap=s.wrap)
    rep.save(path)


def main(args):
    t0 = time.time()
    cfg = CONFIG
    data_dir = os.path.abspath(args.du_lieu or BASE_DIR)
    map_path = os.path.join(data_dir, cfg["file_anh_xa"])
    open_path = os.path.join(data_dir, cfg["file_ton_dau_ky"])
    if not os.path.isdir(data_dir):
        log(f"Không thấy thư mục: {data_dir}")
        return 1

    log(f"ĐỐI SOÁT + QUY ĐỔI ĐƠN VỊ HÓA ĐƠN v{VERSION}")
    log(f"Thư mục: {data_dir}")
    suspicious = [os.path.relpath(os.path.join(dp, fn), data_dir) for dp, _, fns in os.walk(data_dir) for fn in fns
                  if fn.lower().endswith(".xlsm") or fn.lower() == "~$cache1"]
    if suspicious:
        log("")
        log("!!! CẢNH BÁO: thấy file .xlsm / ~$cache1 — script KHÔNG tạo các file này.")
        log("!!! Máy có thể nhiễm virus macro Excel (tự đổi .xlsx → .xlsm). Đừng bật macro, đừng gửi các file này đi.")
        for p in suspicious[:10]:
            log(f"      {p}")
        log("")

    cat_path = find_catalog(data_dir, cfg["file_bang_quy_doi"])
    cat = None
    if not cat_path:
        log("")
        log("!! Thiếu bảng quy đổi đơn vị KiotViet (bang_quy_doi_kiotviet_*.json hoặc kiotvietproducts*.json).")
        log("   TapHoa39BanHang → Quản lý hàng hóa → nút 'Xuất bảng quy đổi' → chép file .json vào thư mục gốc này.")
    else:
        try:
            cat = Catalog(cat_path)
            log(f"  Bảng quy đổi: {os.path.basename(cat_path)} — {len(cat.by_code):,} mã hàng, {len(cat.masters):,} mã master")
        except (ValueError, json.JSONDecodeError) as e:
            log(f"!! Không đọc được {os.path.basename(cat_path)}: {e}")

    # ---------- 1. Đọc XML
    read = read_xml(iter_sources(data_dir))
    submitted_list, submitted_tncn = read["submitted_gtgt"], read["submitted_tncn"]
    if cat is None or not read["invoices"]:
        if not read["invoices"]:
            log("Không có hóa đơn nào → bỏ qua đối soát + tờ khai GTGT.")
        files = []
        if cfg.get("xuat_to_khai_tncn", True):
            period = resolve_period([], cfg) if cfg.get("tu_ngay") and cfg.get("den_ngay") else None
            try:
                files = export_tncn(data_dir, cfg, period, {"mst": (args.mst or cfg["mst_cong_ty"] or "").strip()}, submitted_tncn)
            except PermissionError as e:
                log(f"!! Không ghi được {e.filename} — đóng file đó (Excel/PDF) rồi chạy lại.")
        for f in files:
            log(f"  Kê khai:      {f}")
        return 0 if files else 1

    # ---------- 2–7. Đối soát (không ghi file) — dữ liệu kế toán đã sửa nạp từ thư mục
    vd = None
    vd_path = find_vietda(data_dir, cfg.get("file_danh_muc_vietda") or ["QUYDOI*.xlsx"])
    if vd_path:
        try:
            vd = VietdaCatalog(vd_path)
        except Exception as e:
            log(f"  ! Không đọc được danh mục Vietda {os.path.basename(vd_path)}: {type(e).__name__}: {e}")
    res = reconcile(read["invoices"], cat, cfg, args.mst or cfg["mst_cong_ty"],
                    existing=load_mapping(map_path), opening=load_opening(open_path), vd=vd)
    company, invoices, active, mapping = res["company"], res["invoices"], res["active"], res["mapping"]
    d_from, d_to, ky_label, ky_short = res["period"]
    try:
        save_mapping(map_path, mapping)
    except PermissionError:
        log(f"!! Không ghi được {os.path.basename(map_path)} — hãy ĐÓNG file Excel đó rồi chạy lại.")
        return 1
    try:
        added = save_opening_template(open_path, res["group_info"])
        if added:
            log(f"  Tồn đầu kỳ: thêm {added:,} mã master vào {os.path.basename(open_path)} (điền nếu có, để trống = 0)")
    except PermissionError:
        log(f"  ! Không ghi được {os.path.basename(open_path)} (đang mở?) — bỏ qua")

    # ---------- 7b. Bản sao XML đã quy đổi ĐVT (file gốc giữ nguyên)
    clone_rows, clone_stats = [], Counter()
    if cfg.get("xuat_xml_quy_doi", True):
        clone_rows, clone_stats = export_converted_xml(data_dir, invoices, company, mapping, cfg)
        log(f"  XML quy đổi: '{CLONE_DIRS['VAO']}' {clone_stats['dir_VAO']:,} file, '{CLONE_DIRS['RA']}' "
            f"{clone_stats['dir_RA']:,} file — sửa {clone_stats['sua']:,} file / {clone_stats['dong_sua']:,} dòng hàng; "
            f"chép nguyên {clone_stats['giu']:,} file")
        if clone_stats["dong_chua_anh_xa"]:
            log(f"  ! {clone_stats['dong_chua_anh_xa']:,} dòng hàng giữ nguyên trong bản sao (chưa có ánh xạ / ánh xạ chưa chắc "
                "chắn) → kiểm cột 'Bản sao XML' ở anh_xa_hang_hoa.xlsx, đánh x rồi chạy lại")
        if clone_stats["loi"]:
            log(f"  ! {clone_stats['loi']:,} file không sửa được → chép nguyên (xem sheet 'XML quy doi')")

    # ---------- 8. Xuất Excel
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_path = os.path.join(data_dir, f"doi_soat_{stamp}.xlsx")
    cat_label = f"{os.path.basename(cat_path)} (xuất lúc {cat.exported_at[:16].replace('T', ' ')})"
    summary = summary_rows(res, read, cfg, data_dir, cat_label, clone_stats)
    try:
        write_report(out_path, report_sheets(res, read, summary, clone_rows if cfg.get("xuat_xml_quy_doi", True) else None))
    except PermissionError:
        log(f"!! Không ghi được báo cáo {out_path} — đóng file Excel cùng tên rồi chạy lại.")
        return 1

    # ---------- 9. Bảng kê mua vào/bán ra + tờ khai 01/GTGT + PL NQ204
    tax_files = []
    ra_inv = [i for i in active if i["direction"] == "RA"]
    company_info = {"mst": company,
                    "name": Counter(i["ten_ban"] for i in ra_inv if i["ten_ban"]).most_common(1)[0][0] if ra_inv else "",
                    "address": Counter(i["dc_ban"] for i in ra_inv if i["dc_ban"]).most_common(1)[0][0] if any(i["dc_ban"] for i in ra_inv) else ""}
    if cfg.get("xuat_bang_ke_to_khai", True):
        log("  Đang xuất bảng kê + tờ khai GTGT...")
        try:
            tax_files, tk, pl1, pl2, pl3, ky = export_tax_files(active, company_info, data_dir, cfg, ky_label, ky_short,
                                                                d_from, d_to, res["repl_adj"], submitted_list)
            log(f"  Tờ khai {ky}: [23] {tk['23']:,.0f} · [24] {tk['24']:,.0f} · [34] {tk['34']:,.0f} · [35] {tk['35']:,.0f} · "
                f"[36] {tk['36']:,.0f} · [40] {tk['40']:,.0f} · [43] {tk['43']:,.0f}")
            log(f"  PL NQ204: mua vào 8% {len(pl1):,} mặt hàng (thuế {pl3['06']:,.0f}) · bán ra giảm thuế {len(pl2):,} mặt hàng "
                f"(giảm {pl3['08']:,.0f}) · [09] {pl3['09']:,.0f}")
        except PermissionError as e:
            log(f"!! Không ghi được {e.filename} — đóng file đó (Excel/PDF) rồi chạy lại.")

    # ---------- 10. Tờ khai TNCN 05/KK-TNCN từ bảng lương
    if cfg.get("xuat_to_khai_tncn", True):
        try:
            tax_files += export_tncn(data_dir, cfg, (d_from, d_to, ky_label, ky_short), company_info, submitted_tncn)
        except PermissionError as e:
            log(f"!! Không ghi được {e.filename} — đóng file đó (Excel/PDF) rồi chạy lại.")

    sev_count, n_review = res["sev_count"], res["n_review"]
    log("")
    log(f"XONG sau {time.time() - t0:,.1f} giây")
    log(f"  Báo cáo:      {out_path}")
    log(f"  Ánh xạ hàng:  {map_path}")
    for f in tax_files:
        log(f"  Kê khai:      {f}")
    log(f"  Ngoại lệ: Cao {sev_count['Cao']:,} · Trung bình {sev_count['Trung bình']:,} · Thấp {sev_count['Thấp']:,}")
    if n_review:
        log(f"  → {n_review:,} mặt hàng CẦN XEM (nằm trên cùng trong {os.path.basename(map_path)}): kiểm 'Mã master' + 'Hệ số', rồi chạy lại.")
    return 0


def _inv_ref(inv):
    if inv is None:
        return ("TỒN ĐẦU", "", None, "", "")
    return (inv["ky_hieu"], inv["so"], inv["ngay"], inv["mst_ban"], inv["ten_ban"])


def _ln_name(lot):
    return lot["ln"]["ten"] if lot.get("ln") else "Tồn đầu kỳ"


def _dir_label(inv):
    return "Vào" if inv["direction"] == "VAO" else "Ra"


def _r(i):
    return (i["ky_hieu"], i["so"], i["ngay"], i["mst_ban"], i["ten_ban"], i["mst_mua"], i["ten_mua"],
            i["tong_chua_thue"], i["tong_thue"], i["tong_tt"])


def _tax_summary(active):
    agg = defaultdict(lambda: [set(), 0.0, 0.0])
    for inv in active:
        d = _dir_label(inv)
        if inv["lt_suat"]:
            for _, label, base, tax in inv["lt_suat"]:
                a = agg[(d, label or "?")]
                a[0].add(inv["key"])
                a[1] += base
                a[2] += tax
        else:
            a = agg[(d, "(không có bảng thuế suất)")]
            a[0].add(inv["key"])
            a[1] += inv["tong_chua_thue"]
            a[2] += inv["tong_thue"]
    return [(d, lb, len(v[0]), v[1], v[2]) for (d, lb), v in sorted(agg.items())]


# ============================================================ BANG KE + TO KHAI GTGT + PL NQ204

RA_GROUPS = [
    ("KCT", "Hàng hóa, dịch vụ không chịu thuế GTGT"),
    ("KKKNT", "Hàng hóa, dịch vụ không kê khai nộp thuế GTGT"),
    ("0%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 0%"),
    ("5%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 5%"),
    ("3.5%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 5% giảm 30%"),
    ("7%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 7%"),
    ("8%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 8%"),
    ("10%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 10%"),
    ("10%G30", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 10% giảm 30%"),
    ("20%", "Hàng hóa, dịch vụ chịu thuế suất thuế GTGT 20%"),
    ("KHAC", "Hàng hóa, dịch vụ chịu thuế suất khác"),
]
VAO_GROUPS = [
    ("RIENG", "Hàng hóa, dịch vụ dùng riêng cho SXKD chịu thuế GTGT đủ điều kiện khấu trừ thuế"),
    ("CHUNG", "Hàng hóa, dịch vụ dùng chung cho SXKD chịu thuế và không chịu thuế đủ điều kiện khấu trừ thuế"),
    ("DUAN", "Hàng hóa, dịch vụ dùng cho dự án đầu tư đủ điều kiện được khấu trừ thuế"),
]
_RATE_KEYS = {"0%", "5%", "3.5%", "7%", "8%", "10%", "20%"}
VN_NUM = '#,##0;(#,##0);""'


def rate_group(rate, label):
    if label in ("KCT", "KKKNT"):
        return label
    if rate is None:
        return "KHAC"
    key = f"{round(rate * 100, 2):g}%"
    return key if key in _RATE_KEYS else "KHAC"


def invoice_rate_parts(inv):
    """Tách HĐ theo thuế suất: [(nhóm, thuế suất, nhãn, tiền chưa thuế, thuế, [dòng hàng])].
    Ưu tiên bảng tổng hợp thuế suất in trên HĐ (THTTLTSuat) — đúng số đã kê khai."""
    lines_by = defaultdict(list)
    for ln in inv["lines"]:
        if ln["tchat"] in ("1", "2", "3"):
            lines_by[ln["rate_label"]].append(ln)
    parts = []
    if inv["lt_suat"]:
        for rate, label, base, tax in inv["lt_suat"]:
            parts.append((rate_group(rate, label), rate, label, base, tax, lines_by.get(label, [])))
    else:
        for label, lns in lines_by.items():
            rate = parse_rate(label)[0]
            base = sum(-abs(l["thtien"]) if l["tchat"] == "3" else l["thtien"] for l in lns)
            parts.append((rate_group(rate, label), rate, label, base, round_half_up(base * (rate or 0)), lns))
    if not parts:
        parts.append(("KHAC", None, "", inv["tong_chua_thue"], inv["tong_thue"], []))
    return parts


def _quarter_bounds(y, q):
    start = date(y, 3 * q - 2, 1)
    end = (date(y + (q == 4), 1 if q == 4 else 3 * q + 1, 1)) - timedelta(days=1)
    return start, end


def resolve_period(invoices, cfg):
    """Kỳ tính thuế: CONFIG tu_ngay/den_ngay, nếu trống = quý có nhiều HĐ nhất (theo ngày tính kỳ)."""
    d_from = parse_date(cfg["tu_ngay"]) if cfg.get("tu_ngay") else None
    d_to = parse_date(cfg["den_ngay"]) if cfg.get("den_ngay") else None
    if not (d_from and d_to):
        qs = Counter((i["eff_date"].year, (i["eff_date"].month - 1) // 3 + 1) for i in invoices if i["eff_date"])
        y, q = qs.most_common(1)[0][0] if qs else (date.today().year, (date.today().month - 1) // 3 + 1)
        qf, qt = _quarter_bounds(y, q)
        d_from, d_to = d_from or qf, d_to or qt
    q = (d_from.month - 1) // 3 + 1
    if cfg.get("ky_tinh_thue"):
        label = cfg["ky_tinh_thue"]
        short = re.sub(r"[^0-9A-Za-z]+", "-", strip_accents(label)).strip("-")
    elif (d_from, d_to) == _quarter_bounds(d_from.year, q):
        label, short = f"Quý {q} năm {d_from.year}", f"Q{q}-{d_from.year}"
    else:
        label, short = f"Từ ngày {d_from:%d/%m/%Y} đến ngày {d_to:%d/%m/%Y}", f"{d_from:%Y%m%d}-{d_to:%Y%m%d}"
    return d_from, d_to, label, short


def pl_nq204_rows(invs):
    """Cộng dòng hàng thuế suất 8% theo 'Mã - Tên' (như phụ lục HTKK)."""
    agg = defaultdict(float)
    for inv in invs:
        for ln in inv["lines"]:
            if ln["tchat"] not in ("1", "3") or ln["rate"] is None or abs(ln["rate"] - 0.08) > 1e-9:
                continue
            name = f"{ln['ma']} - {ln['ten']}" if ln["ma"] else ln["ten"]
            agg[(name or "").strip()] += -abs(ln["thtien"]) if ln["tchat"] == "3" else ln["thtien"]
    return sorted(((k, v) for k, v in agg.items() if abs(v) > 0.5), key=lambda x: x[0].lower())


def compute_tax(active, v22):
    ra = [i for i in active if i["direction"] == "RA"]
    vao = [i for i in active if i["direction"] == "VAO"]
    ra_rows, vao_rows = defaultdict(list), defaultdict(list)
    sums = defaultdict(lambda: [0.0, 0.0])
    order = lambda i: (i["ngay"] or date(1900, 1, 1), i["ky_hieu"], i["so"])
    for inv in sorted(ra, key=order):
        for g, rate, label, base, tax, lns in invoice_rate_parts(inv):
            ra_rows[g].append((inv, base, tax))
            sums[g][0] += base
            sums[g][1] += tax
    v23 = v24 = 0.0
    for inv in sorted(vao, key=order):
        for g, rate, label, base, tax, lns in invoice_rate_parts(inv):
            vao_rows["RIENG"].append((inv, rate, label, base, tax, lns))
            v23 += base
            v24 += tax
    r = round_half_up
    tk = {"22": v22, "23": r(v23), "24": r(v24)}
    tk["25"] = tk["24"]
    tk["26"] = r(sums["KCT"][0])
    tk["29"] = r(sums["0%"][0])
    tk["30"] = r(sums["5%"][0] + sums["3.5%"][0])
    tk["31"] = r(sums["5%"][1] + sums["3.5%"][1])
    other = [g for g in sums if g not in ("KCT", "KKKNT", "0%", "5%", "3.5%")]
    tk["32"] = r(sum(sums[g][0] for g in other))
    tk["33"] = r(sum(sums[g][1] for g in other))
    tk["32a"] = r(sums["KKKNT"][0])
    tk["27"] = tk["29"] + tk["30"] + tk["32"] + tk["32a"]
    tk["28"] = tk["31"] + tk["33"]
    tk["34"] = tk["26"] + tk["27"]
    tk["35"] = tk["28"]
    tk["36"] = tk["35"] - tk["25"]
    tk["37"] = tk["38"] = tk["39a"] = tk["40b"] = tk["42"] = 0
    x = tk["36"] - tk["22"] + tk["37"] - tk["38"] - tk["39a"]
    tk["40a"] = x if x > 0 else 0
    tk["40"] = tk["40a"] - tk["40b"]
    tk["41"] = -x if x < 0 else 0
    tk["43"] = tk["41"] - tk["42"]
    pl1 = [(n, v, r(v * 0.08)) for n, v in pl_nq204_rows(vao)]
    pl2 = [(n, v, r(v * 0.02)) for n, v in pl_nq204_rows(ra)]
    pl3 = {"06": sum(t for _, _, t in pl1), "08": sum(t for _, _, t in pl2)}
    pl3["09"] = pl3["08"] - pl3["06"]
    return ra_rows, vao_rows, tk, pl1, pl2, pl3, sums


TK_ROWS = [  # (STT, chỉ tiêu, mã giá trị, mã thuế, in đậm)
    ("A", 'Không phát sinh hoạt động mua, bán trong kỳ (đánh dấu "X")', "21", None, True),
    ("B", "Thuế giá trị gia tăng còn được khấu trừ kỳ trước chuyển sang", None, "22", True),
    ("C", "Kê khai thuế giá trị gia tăng phải nộp ngân sách nhà nước", None, None, True),
    ("I", "Hàng hoá, dịch vụ mua vào trong kỳ", None, None, True),
    ("1", "Giá trị và thuế giá trị gia tăng của hàng hóa, dịch vụ mua vào", "23", "24", False),
    ("", "Trong đó: hàng hóa, dịch vụ nhập khẩu", "23a", "24a", False),
    ("2", "Thuế giá trị gia tăng của hàng hóa, dịch vụ mua vào được khấu trừ kỳ này", None, "25", False),
    ("II", "Hàng hoá, dịch vụ bán ra trong kỳ", None, None, True),
    ("1", "Hàng hóa, dịch vụ bán ra không chịu thuế giá trị gia tăng", "26", None, False),
    ("2", "Hàng hóa, dịch vụ bán ra chịu thuế giá trị gia tăng ([27]=[29]+[30]+[32]+[32a]; [28]=[31]+[33])", "27", "28", False),
    ("a", "Hàng hoá, dịch vụ bán ra chịu thuế suất 0%", "29", None, False),
    ("b", "Hàng hoá, dịch vụ bán ra chịu thuế suất 5%", "30", "31", False),
    ("c", "Hàng hoá, dịch vụ bán ra chịu thuế suất 10%", "32", "33", False),
    ("d", "Hàng hoá, dịch vụ bán ra không tính thuế", "32a", None, False),
    ("3", "Tổng doanh thu và thuế giá trị gia tăng của hàng hóa, dịch vụ bán ra ([34]=[26]+[27]; [35]=[28])", "34", "35", False),
    ("III", "Thuế giá trị gia tăng phát sinh trong kỳ ([36]=[35]-[25])", None, "36", True),
    ("IV", "Điều chỉnh tăng, giảm thuế giá trị gia tăng còn được khấu trừ của các kỳ trước", None, None, True),
    ("1", "Điều chỉnh giảm", None, "37", False),
    ("2", "Điều chỉnh tăng", None, "38", False),
    ("V", "Thuế giá trị gia tăng nhận bàn giao được khấu trừ trong kỳ", None, "39a", True),
    ("VI", "Xác định nghĩa vụ thuế giá trị gia tăng phải nộp trong kỳ:", None, None, True),
    ("1", "Thuế giá trị gia tăng phải nộp của hoạt động sản xuất kinh doanh trong kỳ {[40a]=([36]-[22]+[37]-[38]-[39a]) ≥ 0}", None, "40a", False),
    ("2", "Thuế giá trị gia tăng mua vào của dự án đầu tư được bù trừ với thuế GTGT còn phải nộp của hoạt động sản xuất kinh doanh cùng kỳ tính thuế ([40b]≤[40a])", None, "40b", False),
    ("3", "Thuế giá trị gia tăng còn phải nộp trong kỳ ([40]=[40a]-[40b])", None, "40", False),
    ("4", "Thuế giá trị gia tăng chưa khấu trừ hết kỳ này {[41]=([36]-[22]+[37]-[38]-[39a]) ≤ 0}", None, "41", False),
    ("4.1", "Thuế giá trị gia tăng đề nghị hoàn ([42] ≤ [41])", None, "42", False),
    ("4.2", "Thuế giá trị gia tăng còn được khấu trừ chuyển kỳ sau ([43]=[41]-[42])", None, "43", False),
]


def _xl_title(ws, row, text, last_col, size=12, bold=True, italic=False, align="center"):
    ws.cell(row=row, column=1, value=text).font = Font(bold=bold, italic=italic, size=size)
    ws.cell(row=row, column=1).alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)


def _xl_header(ws, row, headers, widths):
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[_col_letter(i)].width = w
        n = ws.cell(row=row + 1, column=i, value=f"[{i}]")
        n.alignment = Alignment(horizontal="center")
        n.font = Font(italic=True)
        n.border = border
    ws.row_dimensions[row].height = 45
    ws.freeze_panes = ws.cell(row=row + 2, column=1)


def _xl_signature(ws, row, last_col, signer):
    today = datetime.now()
    ws.cell(row=row, column=1, value="Tôi cam đoan số liệu khai trên là đúng và chịu trách nhiệm trước pháp luật về số liệu đã khai./.").font = Font(italic=True)
    sc = max(1, last_col - 4)
    ws.cell(row=row + 2, column=sc, value=f"Ngày {today:%d} tháng {today:%m} năm {today:%Y}").font = Font(italic=True)
    ws.cell(row=row + 3, column=sc, value="NGƯỜI NỘP THUẾ hoặc ĐẠI DIỆN HỢP PHÁP CỦA NGƯỜI NỘP THUẾ").font = Font(bold=True)
    ws.cell(row=row + 4, column=sc, value="(Chữ ký, ghi rõ họ tên; chức vụ và đóng dấu (nếu có)/Ký điện tử)").font = Font(italic=True)
    if signer:
        ws.cell(row=row + 8, column=sc, value=signer).font = Font(bold=True)


def write_bk(path, ky, company, ra_rows, vao_rows, v22, signer):
    wb = Workbook()
    # ---------------- BANRA (01-1/GTGT)
    ws = wb.active
    ws.title = "BANRA"
    last = 10
    ws.cell(row=1, column=8, value="Mẫu số: 01-1/GTGT\n(Kèm theo tờ khai thuế GTGT theo mẫu số 01/GTGT - TT 80/2021/TT-BTC)").alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=1, start_column=8, end_row=1, end_column=last)
    ws.row_dimensions[1].height = 42
    _xl_title(ws, 3, "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ BÁN RA", last, 14)
    _xl_title(ws, 4, ky, last, 12, bold=False, italic=True)
    _xl_title(ws, 5, "BẢN NHÁP — tính từ file XML hóa đơn (doi_soat.py)", last, 10, bold=False, italic=True)
    ws.cell(row=6, column=1, value=f"Người nộp thuế: {company['name']}")
    ws.cell(row=7, column=1, value=f"Mã số thuế: {company['mst']}")
    ws.cell(row=8, column=1, value="Đơn vị tiền: Đồng")
    _xl_header(ws, 9, ["STT", "Ký hiệu mẫu", "Ký hiệu HĐ", "Số hóa đơn", "Ngày lập", "Tên người mua",
                       "Mã số thuế người mua", "Doanh số bán chưa có thuế", "Thuế GTGT", "Ghi chú"],
               [7, 7, 11, 11, 11, 42, 16, 18, 15, 16])
    row = 11
    group_total_cells, taxable_total_cells = [], []
    for gi, (gkey, gname) in enumerate(RA_GROUPS, start=1):
        ws.cell(row=row, column=1, value=f"{gi}. {gname}").font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        row += 1
        first = row
        for n, (inv, base, tax) in enumerate(ra_rows.get(gkey, []), start=1):
            note = "" if inv["status"] in ("Gốc", "Bị điều chỉnh") else f"HĐ {inv['status'].lower()}"
            vals = [n, inv["khms"], inv["khhd"], inv["so"], inv["ngay"], inv["ten_mua"], inv["mst_mua"], base, tax, note]
            for ci, v in enumerate(vals, start=1):
                c = ws.cell(row=row, column=ci, value=v)
                if ci == 5 and v:
                    c.number_format = "dd/mm/yyyy"
                elif ci in (8, 9):
                    c.number_format = VN_NUM
            row += 1
        ws.cell(row=row, column=1, value=f"Tổng {gname}").font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        for ci in (8, 9):
            col = _col_letter(ci)
            c = ws.cell(row=row, column=ci, value=f"=SUM({col}{first}:{col}{row - 1})" if row > first else 0)
            c.font = Font(bold=True)
            c.number_format = VN_NUM
        group_total_cells.append(row)
        if gkey not in ("KCT", "KKKNT"):
            taxable_total_cells.append(row)
        row += 1
    row += 1
    for label, col, rows_ in (("Tổng giá trị hàng hoá, dịch vụ bán ra:", "H", group_total_cells),
                              ("Tổng thuế GTGT của hàng hoá, dịch vụ bán ra:", "I", group_total_cells),
                              ("Tổng doanh thu HH, DV bán ra chịu thuế GTGT:", "H", taxable_total_cells)):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=row, column=8, value="=" + "+".join(f"{col}{r}" for r in rows_))
        c.font = Font(bold=True)
        c.number_format = VN_NUM
        if label.startswith("Tổng thuế"):
            ra_tax_cell = f"BANRA!H{row}"
        row += 1
    _xl_signature(ws, row + 1, last, signer)

    # ---------------- MUAVAO (01-1/HT)
    wv = wb.create_sheet("MUAVAO")
    last = 15
    wv.cell(row=1, column=12, value="Mẫu số: 01-1/HT\n(Ban hành kèm theo Thông tư số 80/2021/TT-BTC)").alignment = Alignment(wrap_text=True)
    wv.merge_cells(start_row=1, start_column=12, end_row=1, end_column=last)
    wv.row_dimensions[1].height = 32
    _xl_title(wv, 3, "BẢNG KÊ HOÁ ĐƠN, CHỨNG TỪ HÀNG HOÁ, DỊCH VỤ MUA VÀO", last, 14)
    _xl_title(wv, 4, ky, last, 12, bold=False, italic=True)
    _xl_title(wv, 5, "BẢN NHÁP — tính từ file XML hóa đơn (doi_soat.py). HĐ nhiều thuế suất được tách 1 dòng / thuế suất.", last, 10, bold=False, italic=True)
    wv.cell(row=6, column=1, value=f"[02] Tên người nộp thuế: {company['name']}")
    wv.cell(row=7, column=1, value=f"[05] Mã số thuế: {company['mst']}")
    wv.cell(row=8, column=1, value="Đơn vị tiền: Đồng")
    _xl_header(wv, 9, ["STT", "Mẫu số", "Ký hiệu", "Số", "Ngày, tháng, năm", "Tên người bán", "Mã số thuế người bán",
                       "Tên hàng hóa, dịch vụ", "Đơn vị tính", "Số lượng", "Đơn giá",
                       "Giá trị HHDV mua vào chưa có thuế", "Thuế suất (%)", "Thuế GTGT", "Ghi chú"],
               [6, 6, 10, 10, 11, 36, 15, 42, 8, 9, 12, 17, 8, 14, 14])
    row = 11
    totals = []
    for gi, (gkey, gname) in enumerate(VAO_GROUPS, start=1):
        wv.cell(row=row, column=1, value=f"{gi}. {gname}").font = Font(bold=True)
        wv.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
        row += 1
        first = row
        for n, (inv, rate, label, base, tax, lns) in enumerate(vao_rows.get(gkey, []), start=1):
            goods = [l for l in lns if l["tchat"] in ("1", "2")]
            name = goods[0]["ten"] if goods else (lns[0]["ten"] if lns else "")
            if len(goods) > 1:
                name = f"{name} (+{len(goods) - 1} mặt hàng)"
            single = goods[0] if len(goods) == 1 else None
            rate_show = int(round(rate * 100)) if rate is not None and label not in ("KCT", "KKKNT") and float(rate * 100).is_integer() else (label or "")
            note = "" if inv["status"] in ("Gốc", "Bị điều chỉnh") else f"HĐ {inv['status'].lower()}"
            vals = [n, inv["khms"], inv["khhd"], inv["so"], inv["ngay"], inv["ten_ban"], inv["mst_ban"], name,
                    single["dvt"] if single else None, single["sl"] if single else None,
                    single["dgia"] if single else None, base, rate_show, tax, note]
            for ci, v in enumerate(vals, start=1):
                c = wv.cell(row=row, column=ci, value=v)
                if ci == 5 and v:
                    c.number_format = "dd/mm/yyyy"
                elif ci in (11, 12, 14):
                    c.number_format = VN_NUM
            row += 1
        wv.cell(row=row, column=1, value=f"Tổng {gname}").font = Font(bold=True)
        wv.merge_cells(start_row=row, start_column=1, end_row=row, end_column=11)
        for ci in (12, 14):
            col = _col_letter(ci)
            c = wv.cell(row=row, column=ci, value=f"=SUM({col}{first}:{col}{row - 1})" if row > first else 0)
            c.font = Font(bold=True)
            c.number_format = VN_NUM
        totals.append(row)
        row += 1
    wv.cell(row=row, column=1, value="Tổng cộng").font = Font(bold=True)
    for ci in (12, 14):
        col = _col_letter(ci)
        c = wv.cell(row=row, column=ci, value="=" + "+".join(f"{col}{r}" for r in totals))
        c.font = Font(bold=True)
        c.number_format = VN_NUM
    total_row = row
    _xl_signature(wv, row + 2, last, signer)
    row += 13
    wv.cell(row=row, column=1, value="THUẾ GTGT CÒN ĐƯỢC KHẤU TRỪ KỲ TRƯỚC CHUYỂN SANG [22]").font = Font(bold=True)
    c = wv.cell(row=row, column=14, value=v22)
    c.number_format = VN_NUM
    wv.cell(row=row + 1, column=1, value="THUẾ GTGT PHẢI NỘP (số âm = còn được khấu trừ chuyển kỳ sau)").font = Font(bold=True)
    c = wv.cell(row=row + 1, column=14, value=f"={ra_tax_cell}-N{row}-N{total_row}")
    c.font = Font(bold=True)
    c.number_format = VN_NUM
    wb.save(path)


def write_tk_xlsx(path, ky, company, tk, pl1, pl2, pl3, signer, repl_adj=None, compare=None, compare_file=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "To khai 01-GTGT"
    widths = [6, 80, 7, 20, 7, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[_col_letter(i)].width = w
    _xl_title(ws, 1, "TỜ KHAI THUẾ GIÁ TRỊ GIA TĂNG (GTGT) — Mẫu số 01/GTGT (TT 80/2021/TT-BTC)", 6, 14)
    _xl_title(ws, 2, "BẢN NHÁP ĐỐI CHIẾU — số liệu tính từ file XML hóa đơn; nộp chính thức qua HTKK / eTax", 6, 10, bold=False, italic=True)
    ws.cell(row=2, column=1).font = Font(italic=True, color="C00000")
    ws.cell(row=4, column=1, value=f"[01b] Kỳ tính thuế: {ky}")
    ws.cell(row=5, column=1, value=f"[04] Tên người nộp thuế: {company['name']}")
    ws.cell(row=6, column=1, value=f"[05] Mã số thuế: {company['mst']}")
    ws.cell(row=7, column=1, value="Đơn vị tiền: Đồng Việt Nam")
    _xl_header(ws, 9, ["STT", "Chỉ tiêu", "Mã", "Giá trị HHDV (chưa có thuế GTGT)", "Mã", "Thuế GTGT"], widths)
    row = 11
    for stt, text, cv, ct, bold in TK_ROWS:
        ws.cell(row=row, column=1, value=stt).font = Font(bold=bold)
        ws.cell(row=row, column=2, value=text).font = Font(bold=bold)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        for col, code in ((3, cv), (5, ct)):
            if code:
                ws.cell(row=row, column=col, value=f"[{code}]")
                c = ws.cell(row=row, column=col + 1, value=tk.get(code, "") if code != "21" else "")
                c.number_format = VN_NUM
                c.font = Font(bold=bold)
        row += 1
    notes = ["Ghi chú tính số (bản nháp):",
             "• Không có dòng thuế suất 8%: hàng giảm 10% → 8% kê ở [32]/[33] (tiền thuế tính theo 8%), chi tiết ở phụ lục 142 (NQ 204/2025/QH15).",
             "• HĐ thay thế tính vào kỳ của HĐ gốc; HĐ điều chỉnh tính vào kỳ lập HĐ điều chỉnh. Danh sách: sheet 'HD thay the - dieu chinh' trong file doi_soat_*.xlsx."]
    names = {"RA_tru": "HĐ ra thay thế cho kỳ khác (đã TRỪ khỏi kỳ này)", "RA_cong": "HĐ ra lập kỳ sau thay thế HĐ kỳ này (đã CỘNG vào kỳ này)",
             "VAO_tru": "HĐ vào thay thế cho kỳ khác (đã TRỪ khỏi kỳ này)", "VAO_cong": "HĐ vào lập kỳ sau thay thế HĐ kỳ này (đã CỘNG vào kỳ này)"}
    for k, label in names.items():
        if repl_adj and k in repl_adj:
            n, b, t = repl_adj[k]
            notes.append(f"• {label}: {n} HĐ, tiền chưa thuế {_fmt_vn(b) or 0}, thuế {_fmt_vn(t) or 0}")
    row += 1
    for line in notes:
        ws.cell(row=row, column=2, value=line).font = Font(italic=True, bold=line.endswith(":"))
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        row += 1
    _xl_signature(ws, row + 1, 6, signer)

    w2 = wb.create_sheet("PL142 NQ204")
    for i, w in enumerate([7, 70, 18, 10, 10, 16], start=1):
        w2.column_dimensions[_col_letter(i)].width = w
    _xl_title(w2, 1, "GIẢM THUẾ GIÁ TRỊ GIA TĂNG THEO NGHỊ QUYẾT SỐ 204/2025/QH15", 6, 14)
    _xl_title(w2, 2, f"(Kèm theo Tờ khai thuế GTGT Kỳ tính thuế: {ky}) — BẢN NHÁP ĐỐI CHIẾU", 6, 10, bold=False, italic=True)
    w2.cell(row=3, column=1, value=f"[01] Tên người nộp thuế: {company['name']}")
    w2.cell(row=4, column=1, value=f"[02] Mã số thuế: {company['mst']}")
    row = 6
    w2.cell(row=row, column=1, value="I. Hàng hóa, dịch vụ mua vào trong kỳ được áp dụng mức thuế suất thuế giá trị gia tăng 8% (áp dụng cho người nộp thuế kê khai theo phương pháp khấu trừ thuế)").font = Font(bold=True)
    w2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    w2.row_dimensions[row].height = 30
    w2.cell(row=row, column=1).alignment = Alignment(wrap_text=True)
    row += 1
    for ci, h in enumerate(["STT", "Tên hàng hóa, dịch vụ", "Giá trị HHDV mua vào chưa có thuế GTGT được khấu trừ trong kỳ",
                            "Thuế GTGT của HHDV mua vào được khấu trừ trong kỳ"], start=1):
        c = w2.cell(row=row, column=ci if ci < 4 else 6, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    w2.row_dimensions[row].height = 45
    row += 1
    first = row
    for n, (name, v, t) in enumerate(pl1, start=1):
        w2.cell(row=row, column=1, value=n)
        w2.cell(row=row, column=2, value=name)
        w2.cell(row=row, column=3, value=v).number_format = VN_NUM
        w2.cell(row=row, column=6, value=t).number_format = VN_NUM
        row += 1
    w2.cell(row=row, column=2, value="Tổng cộng").font = Font(bold=True)
    for ci in (3, 6):
        col = _col_letter(ci)
        c = w2.cell(row=row, column=ci, value=f"=SUM({col}{first}:{col}{row - 1})" if row > first else 0)
        c.font = Font(bold=True)
        c.number_format = VN_NUM
    row += 2
    w2.cell(row=row, column=1, value="II. Hàng hóa, dịch vụ bán ra trong kỳ").font = Font(bold=True)
    row += 1
    for ci, h in enumerate(["STT", "Tên hàng hóa, dịch vụ", "Giá trị HHDV chưa có thuế GTGT", "Thuế suất thuế GTGT theo quy định",
                            "Thuế suất thuế GTGT sau giảm", "Thuế GTGT của HHDV bán ra được giảm"], start=1):
        c = w2.cell(row=row, column=ci, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    w2.row_dimensions[row].height = 45
    row += 1
    for ci, h in enumerate(["(1)", "(2)", "(3)", "(4)", "(5)=(4)x80%", "(6)=(3)x[(4)-(5)]"], start=1):
        w2.cell(row=row, column=ci, value=h).alignment = Alignment(horizontal="center")
    row += 1
    first = row
    for n, (name, v, t) in enumerate(pl2, start=1):
        w2.cell(row=row, column=1, value=n)
        w2.cell(row=row, column=2, value=name)
        w2.cell(row=row, column=3, value=v).number_format = VN_NUM
        w2.cell(row=row, column=4, value=10)
        w2.cell(row=row, column=5, value=8)
        w2.cell(row=row, column=6, value=t).number_format = VN_NUM
        row += 1
    w2.cell(row=row, column=2, value="Tổng cộng").font = Font(bold=True)
    for ci in (3, 6):
        col = _col_letter(ci)
        c = w2.cell(row=row, column=ci, value=f"=SUM({col}{first}:{col}{row - 1})" if row > first else 0)
        c.font = Font(bold=True)
        c.number_format = VN_NUM
    row += 2
    w2.cell(row=row, column=1, value="III. Chênh lệch thuế GTGT của hàng hoá, dịch vụ bán ra và mua vào trong kỳ được áp dụng mức thuế suất thuế GTGT 8%: [09] = [08] - [06]").font = Font(bold=True)
    c = w2.cell(row=row, column=6, value=pl3["09"])
    c.number_format = VN_NUM
    c.font = Font(bold=True)
    w2.cell(row=row + 1, column=2, value=f"[06] (tổng thuế mục I) = {pl3['06']:,.0f}   ·   [08] (tổng thuế được giảm mục II) = {pl3['08']:,.0f}".replace(",", "."))

    if compare:
        w3 = wb.create_sheet("So sanh to khai da nop")
        for i, w in enumerate([40, 18, 18, 14, 18], start=1):
            w3.column_dimensions[_col_letter(i)].width = w
        _xl_title(w3, 1, f"SO SÁNH VỚI TỜ KHAI ĐÃ NỘP — {ky}", 5, 13)
        _xl_title(w3, 2, f"File: {compare_file}", 5, 9, bold=False, italic=True, align="left")
        w3.append([])
        w3.append(["Chỉ tiêu", "Đã nộp (XML)", "Script tính", "Chênh (script − đã nộp)", "Kết quả"])
        for c in w3[4]:
            c.font = Font(bold=True)
        for label, a, b, diff, verdict in compare:
            w3.append([label, a, b, diff, verdict])
            r = w3.max_row
            for ci in (2, 3, 4):
                w3.cell(row=r, column=ci).number_format = "#,##0;(#,##0);0"
            if verdict.startswith("LỆCH"):
                w3.cell(row=r, column=5).font = Font(bold=True, color="C00000")
        w3.append([])
        w3.append(["Lệch thường do: thiếu/thừa file XML (HĐ hủy, HĐ thay thế lập đầu kỳ sau chưa chép vào), [22] khác, "
                   "kế toán điều chỉnh tay, hoặc làm tròn."])
    wb.save(path)


def _fmt_vn(v):
    if v in (None, "") or (isinstance(v, (int, float)) and abs(v) < 0.5):
        return ""
    n = round_half_up(v)
    s = f"{abs(n):,}".replace(",", ".")
    return f"({s})" if n < 0 else s


def _ensure_reportlab():
    try:
        import reportlab  # noqa: F401
    except ImportError:
        log("  Lần đầu xuất PDF: đang cài thư viện reportlab...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--disable-pip-version-check", "-q", "reportlab"])


def _pdf_fonts():
    """Đăng ký font tiếng Việt của Windows (Times New Roman, không có thì Arial) tên VN / VN-B / VN-I."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_dir = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for reg, bold, ital in (("times.ttf", "timesbd.ttf", "timesi.ttf"), ("arial.ttf", "arialbd.ttf", "ariali.ttf")):
        paths = [os.path.join(font_dir, f) for f in (reg, bold, ital)]
        if all(os.path.exists(p) for p in paths):
            break
    else:
        log("  ! Không tìm thấy font Times New Roman / Arial trong Windows\\Fonts — bỏ qua PDF")
        return False
    pdfmetrics.registerFont(TTFont("VN", paths[0]))
    pdfmetrics.registerFont(TTFont("VN-B", paths[1]))
    pdfmetrics.registerFont(TTFont("VN-I", paths[2]))
    pdfmetrics.registerFontFamily("VN", normal="VN", bold="VN-B", italic="VN-I", boldItalic="VN-B")
    return True


def _pdf_draft_page(banner):
    """Hàm vẽ mỗi trang PDF: chữ mờ 'BẢN NHÁP' + dòng cảnh báo đầu trang + số trang."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.Color(0.88, 0.88, 0.88))
        canvas.setFont("VN-B", 64)
        canvas.translate(A4[0] / 2, A4[1] / 2)
        canvas.rotate(40)
        canvas.drawCentredString(0, 0, "BẢN NHÁP")
        canvas.restoreState()
        canvas.saveState()
        canvas.setFillColor(colors.Color(0.75, 0, 0))
        canvas.setFont("VN-B", 8)
        canvas.drawCentredString(A4[0] / 2, A4[1] - 10 * mm, banner)
        canvas.setFillColor(colors.black)
        canvas.setFont("VN", 8)
        canvas.drawRightString(A4[0] - 12 * mm, 8 * mm, f"Trang {doc.page}")
        canvas.restoreState()
    return on_page


def write_tk_pdf(path, ky, company, tk, pl1, pl2, pl3, signer):
    _ensure_reportlab()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if not _pdf_fonts():
        return False

    st = ParagraphStyle("n", fontName="VN", fontSize=9, leading=11)
    st_b = ParagraphStyle("b", parent=st, fontName="VN-B")
    st_c = ParagraphStyle("c", parent=st, alignment=TA_CENTER)
    st_cb = ParagraphStyle("cb", parent=st_c, fontName="VN-B")
    st_t = ParagraphStyle("t", parent=st_cb, fontSize=13, leading=16)
    st_i = ParagraphStyle("i", parent=st, fontName="VN-I")
    st_r = ParagraphStyle("r", parent=st, alignment=TA_RIGHT)
    st_rb = ParagraphStyle("rb", parent=st_r, fontName="VN-B")
    on_page = _pdf_draft_page("BẢN NHÁP ĐỐI CHIẾU — tính từ XML hóa đơn, KHÔNG DÙNG ĐỂ NỘP (nộp chính thức qua HTKK / eTax)")

    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                       ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)])
    P = Paragraph
    story = []
    head = Table([[P("<b>CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b>Độc lập - Tự do - Hạnh phúc</b>", st_c),
                   P("<b>Mẫu số: 01/GTGT</b><br/>(Ban hành kèm theo Thông tư số 80/2021/TT-BTC ngày 29 tháng 9 năm 2021 của Bộ trưởng Bộ Tài chính)", st_c)]],
                 colWidths=[110 * mm, 70 * mm])
    story += [head, Spacer(1, 4 * mm), P("TỜ KHAI THUẾ GIÁ TRỊ GIA TĂNG (GTGT)", st_t),
              P("(Áp dụng đối với người nộp thuế tính thuế theo phương pháp khấu trừ có hoạt động sản xuất kinh doanh)", ParagraphStyle("x", parent=st_c, fontName="VN-I")),
              Spacer(1, 3 * mm),
              P(f"[01b] Kỳ tính thuế: <b>{ky}</b>", st), P("[02] Lần đầu: X", st),
              P(f"[04] Tên người nộp thuế: <b>{company['name']}</b>", st), P(f"[05] Mã số thuế: <b>{company['mst']}</b>", st),
              P("Đơn vị tiền: Đồng Việt Nam", st_r), Spacer(1, 2 * mm)]
    data = [[P("<b>STT</b>", st_c), P("<b>Chỉ tiêu</b>", st_c), P("<b>Giá trị hàng hóa, dịch vụ (chưa có thuế GTGT)</b>", st_c), "",
             P("<b>Thuế giá trị gia tăng</b>", st_c), ""]]
    for stt, text, cv, ct, bold in TK_ROWS:
        sty, styr = (st_b, st_rb) if bold else (st, st_r)
        data.append([P(stt, ParagraphStyle("s", parent=sty, alignment=TA_CENTER)), P(text, sty),
                     P(f"[{cv}]", st_c) if cv else "", P(_fmt_vn(tk.get(cv)) if cv and cv != "21" else "", styr),
                     P(f"[{ct}]", st_c) if ct else "", P(_fmt_vn(tk.get(ct)) if ct else "", styr)])
    t = Table(data, colWidths=[10 * mm, 88 * mm, 11 * mm, 27 * mm, 11 * mm, 27 * mm], repeatRows=1)
    t.setStyle(grid)
    t.setStyle(TableStyle([("SPAN", (2, 0), (3, 0)), ("SPAN", (4, 0), (5, 0))]))
    story += [t, Spacer(1, 4 * mm),
              P("Tôi cam đoan số liệu khai trên là đúng và chịu trách nhiệm trước pháp luật về số liệu đã khai./.", st_i),
              Spacer(1, 3 * mm)]
    today = datetime.now()
    sign = Table([["", P(f"<i>Ngày {today:%d} tháng {today:%m} năm {today:%Y}</i><br/><b>NGƯỜI NỘP THUẾ hoặc<br/>ĐẠI DIỆN HỢP PHÁP CỦA NGƯỜI NỘP THUẾ</b><br/>"
                               f"<i>(Chữ ký, ghi rõ họ tên; chức vụ và đóng dấu (nếu có)/Ký điện tử)</i><br/><br/><br/><b>{signer}</b>", st_c)]],
                 colWidths=[90 * mm, 90 * mm])
    story += [sign, PageBreak()]

    story += [P("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b>Độc lập - Tự do - Hạnh phúc</b>", st_c), Spacer(1, 3 * mm),
              P("GIẢM THUẾ GIÁ TRỊ GIA TĂNG THEO NGHỊ QUYẾT SỐ 204/2025/QH15", st_t),
              P(f"(Kèm theo Tờ khai thuế GTGT Kỳ tính thuế: {ky})", ParagraphStyle("x2", parent=st_c, fontName="VN-I")),
              Spacer(1, 2 * mm), P(f"[01] Tên người nộp thuế: <b>{company['name']}</b>", st),
              P(f"[02] Mã số thuế: <b>{company['mst']}</b>", st), P("Đơn vị tiền: Đồng Việt Nam", st_r), Spacer(1, 2 * mm),
              P("<b>I. Hàng hóa, dịch vụ mua vào trong kỳ được áp dụng mức thuế suất thuế giá trị gia tăng 8%</b> "
                "(áp dụng cho người nộp thuế kê khai theo phương pháp khấu trừ thuế)", st), Spacer(1, 1 * mm)]
    d1 = [[P("<b>STT</b>", st_c), P("<b>Tên hàng hóa, dịch vụ</b>", st_c),
           P("<b>Giá trị hàng hóa, dịch vụ mua vào chưa có thuế GTGT được khấu trừ trong kỳ</b>", st_c),
           P("<b>Thuế GTGT của hàng hóa, dịch vụ mua vào được khấu trừ trong kỳ</b>", st_c)],
          [P("(1)", st_c), P("(2)", st_c), P("(3)", st_c), P("(4)", st_c)]]
    d1 += [[P(f"{n:,}".replace(",", "."), st_c), P(name, st), P(_fmt_vn(v), st_r), P(_fmt_vn(tx), st_r)]
           for n, (name, v, tx) in enumerate(pl1, start=1)]
    d1.append(["", P("<b>Tổng cộng</b>", st), P(_fmt_vn(sum(v for _, v, _ in pl1)), st_rb), P(_fmt_vn(pl3["06"]), st_rb)])
    t1 = LongTable(d1, colWidths=[13 * mm, 105 * mm, 32 * mm, 30 * mm], repeatRows=2)
    t1.setStyle(grid)
    story += [t1, Spacer(1, 4 * mm), P("<b>II. Hàng hóa, dịch vụ bán ra trong kỳ</b>", st), Spacer(1, 1 * mm)]
    d2 = [[P("<b>STT</b>", st_c), P("<b>Tên hàng hóa, dịch vụ</b>", st_c), P("<b>Giá trị hàng hóa, dịch vụ chưa có thuế GTGT</b>", st_c),
           P("<b>Thuế suất thuế GTGT theo quy định</b>", st_c), P("<b>Thuế suất thuế GTGT sau giảm</b>", st_c),
           P("<b>Thuế GTGT của hàng hóa, dịch vụ bán ra được giảm</b>", st_c)],
          [P("(1)", st_c), P("(2)", st_c), P("(3)", st_c), P("(4)", st_c), P("(5)=(4)x80%", st_c), P("(6)=(3)x[(4)-(5)]", st_c)]]
    d2 += [[P(f"{n:,}".replace(",", "."), st_c), P(name, st), P(_fmt_vn(v), st_r), P("10", st_c), P("8", st_c), P(_fmt_vn(tx), st_r)]
           for n, (name, v, tx) in enumerate(pl2, start=1)]
    d2.append(["", P("<b>Tổng cộng</b>", st), P(_fmt_vn(sum(v for _, v, _ in pl2)), st_rb), "", "", P(_fmt_vn(pl3["08"]), st_rb)])
    t2 = LongTable(d2, colWidths=[13 * mm, 83 * mm, 28 * mm, 18 * mm, 18 * mm, 24 * mm], repeatRows=2)
    t2.setStyle(grid)
    story += [t2, Spacer(1, 4 * mm),
              P(f"<b>III. Chênh lệch thuế GTGT của hàng hoá, dịch vụ bán ra và mua vào trong kỳ được áp dụng mức thuế suất "
                f"thuế giá trị gia tăng 8%: [09] = [08] - [06]: {_fmt_vn(pl3['09']) or '0'} đồng</b>", st),
              Spacer(1, 4 * mm), sign]
    SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=12 * mm, topMargin=16 * mm, bottomMargin=14 * mm,
                      title=f"To khai GTGT {ky} - BAN NHAP", author="doi_soat.py").build(story, onFirstPage=on_page, onLaterPages=on_page)
    return True


COMPARE_CODES = ["22", "23", "24", "25", "26", "27", "28", "29", "30", "31", "32", "33", "32a", "34", "35", "36",
                 "37", "38", "39a", "40a", "40b", "40", "41", "42", "43"]


def _ky_of(submitted):
    m = re.match(r"(\d)/(\d{4})", submitted.get("_ky", ""))
    return (int(m.group(1)), int(m.group(2))) if m else None


def tax_figures(active, cfg, d_from, submitted_list):
    """Số liệu tờ khai 01/GTGT + PL 142 (NQ 204), KHÔNG ghi file: [22] lấy từ CONFIG / [43] tờ khai kỳ trước / tờ khai cùng
    kỳ đã nộp; so sánh với tờ khai cùng kỳ đã nộp nếu có. Dùng chung cho CLI và KeToanBackEnd."""
    q_now = ((d_from.month - 1) // 3 + 1, d_from.year)
    q_prev = (4, q_now[1] - 1) if q_now[0] == 1 else (q_now[0] - 1, q_now[1])
    same = next((s for s in submitted_list if _ky_of(s) == q_now), None)
    prev = next((s for s in submitted_list if _ky_of(s) == q_prev), None)
    v22 = round_half_up(parse_num(cfg.get("thue_con_khau_tru_ky_truoc_22", 0)))
    if not v22 and prev is not None and "43" in prev:
        v22 = round_half_up(prev["43"])
        log(f"  [22] = {v22:,.0f} lấy từ [43] tờ khai kỳ trước ({prev['_file']})")
    elif not v22 and same is not None and "22" in same:
        v22 = round_half_up(same["22"])
        log(f"  [22] = {v22:,.0f} lấy từ tờ khai cùng kỳ đã nộp ({same['_file']})")
    elif not v22:
        log("  ! Chỉ tiêu [22] = 0 — nhập 'thue_con_khau_tru_ky_truoc_22' (= [43] tờ khai kỳ trước) trong CONFIG, "
            "hoặc chép XML tờ khai kỳ trước vào thư mục")
    ra_rows, vao_rows, tk, pl1, pl2, pl3, _ = compute_tax(active, v22)
    compare = None
    if same is not None:
        mine = dict(tk, pl1_gt=sum(v for _, v, _ in pl1), pl1_thue=pl3["06"], pl1_dong=len(pl1),
                    pl2_gt=sum(v for _, v, _ in pl2), pl2_giam=pl3["08"], pl2_dong=len(pl2), pl_09=pl3["09"])
        labels = {"pl1_gt": "PL142 I – tổng giá trị mua vào 8%", "pl1_thue": "PL142 I – tổng thuế [06]",
                  "pl1_dong": "PL142 I – số dòng mặt hàng", "pl2_gt": "PL142 II – tổng giá trị bán ra 8%",
                  "pl2_giam": "PL142 II – tổng thuế được giảm [08]", "pl2_dong": "PL142 II – số dòng mặt hàng",
                  "pl_09": "PL142 III – [09] = [08] − [06]"}
        compare = []
        for code in COMPARE_CODES + list(labels):
            a, b = same.get(code), mine.get(code)
            if a is None and not b:
                continue
            diff = (b or 0) - (a or 0)
            if abs(diff) < 0.5:
                verdict = "Khớp"
            elif code.endswith("_dong"):
                verdict = "LỆCH — khác số dòng mặt hàng"
            else:
                verdict = "Lệch làm tròn" if abs(diff) <= 5 else "LỆCH — xem lại"
            compare.append((f"[{code}]" if code in COMPARE_CODES else labels[code], a, b, diff, verdict))
        n_bad = sum(1 for r in compare if r[4].startswith("LỆCH"))
        log(f"  So với tờ khai đã nộp ({same['_file']}): {len(compare) - n_bad}/{len(compare)} chỉ tiêu khớp"
            + (f", {n_bad} LỆCH — xem sheet 'So sanh to khai da nop'" if n_bad else ""))
    return {"v22": v22, "ra_rows": ra_rows, "vao_rows": vao_rows, "tk": tk, "pl1": pl1, "pl2": pl2, "pl3": pl3,
            "compare": compare, "same": same, "prev": prev}


def export_tax_files(active, company, data_dir, cfg, ky, short, d_from, d_to, repl_adj, submitted_list):
    """Xuất bảng kê mua vào/bán ra + tờ khai 01/GTGT + PL 142 (NQ 204) (xlsx, pdf bản nháp)."""
    t = tax_figures(active, cfg, d_from, submitted_list)
    v22, tk, pl1, pl2, pl3, compare, same, prev = (t[k] for k in ("v22", "tk", "pl1", "pl2", "pl3", "compare", "same", "prev"))
    signer = cfg.get("nguoi_ky", "")
    files = []
    bk = os.path.join(data_dir, f"BK_MUAVAO_BANRA_{short}.xlsx")
    write_bk(bk, ky, company, t["ra_rows"], t["vao_rows"], v22, signer)
    files.append(bk)
    tkx = os.path.join(data_dir, f"TK_GTGT_{short}.xlsx")
    write_tk_xlsx(tkx, ky, company, tk, pl1, pl2, pl3, signer, repl_adj, compare, same["_file"] if same else "")
    files.append(tkx)
    if cfg.get("xuat_pdf_ban_nhap", True):
        pdf = os.path.join(data_dir, f"TK_GTGT_{short}_BAN_NHAP.pdf")
        if write_tk_pdf(pdf, ky, company, tk, pl1, pl2, pl3, signer):
            files.append(pdf)
    if cfg.get("xuat_xml_to_khai", True):
        ref = same or prev or (submitted_list[-1] if submitted_list else None)
        values = {code: tk.get(code, 0) for code in COMPARE_CODES + ["21", "23a", "24a"]}
        x = write_tk_xml(data_dir, "842", ref, d_from, d_to, values, signer, company, (pl1, pl2, pl3))
        if x:
            files.append(x)
    return files, tk, pl1, pl2, pl3, ky


# ============================================================ TỜ KHAI TNCN 05/KK-TNCN (từ bảng lương)

TNCN_ROWS = [  # (STT, chỉ tiêu, mã, đơn vị, in đậm) — mã "25_1" = [25.1] (tên thẻ XML ct25_1)
    ("1", "Tổng số người lao động:", "16", "Người", True),
    ("", "Trong đó: Cá nhân cư trú có hợp đồng lao động", "17", "Người", False),
    ("2", "Tổng số cá nhân đã khấu trừ thuế [18]=[19]+[20]", "18", "Người", True),
    ("2.1", "Cá nhân cư trú", "19", "Người", False),
    ("2.2", "Cá nhân không cư trú", "20", "Người", False),
    ("3", "Tổng thu nhập chịu thuế trả cho cá nhân [21]=[22]+[23]", "21", "VNĐ", True),
    ("3.1", "Cá nhân cư trú", "22", "VNĐ", False),
    ("3.2", "Cá nhân không cư trú", "23", "VNĐ", False),
    ("3.3", "Trong đó: Tổng thu nhập chịu thuế từ tiền phí mua bảo hiểm nhân thọ, bảo hiểm không bắt buộc khác của "
            "doanh nghiệp bảo hiểm không thành lập tại Việt Nam cho người lao động", "24", "VNĐ", False),
    ("4", "Trong đó tổng thu nhập chịu thuế được miễn theo quy định của Hợp đồng dầu khí", "25", "VNĐ", True),
    ("", "Thu nhập miễn thuế, thu nhập cá nhân theo nghị quyết", "25_1", "VNĐ", False),
    ("5", "Tổng thu nhập chịu thuế trả cho cá nhân thuộc diện phải khấu trừ thuế [26]=[27]+[28]", "26", "VNĐ", True),
    ("5.1", "Cá nhân cư trú", "27", "VNĐ", False),
    ("5.2", "Cá nhân không cư trú", "28", "VNĐ", False),
    ("6", "Tổng số thuế thu nhập cá nhân đã khấu trừ [29]=[30]+[31]", "29", "VNĐ", True),
    ("6.1", "Cá nhân cư trú", "30", "VNĐ", False),
    ("6.2", "Cá nhân không cư trú", "31", "VNĐ", False),
    ("6.3", "Trong đó: Tổng số thuế thu nhập cá nhân đã khấu trừ trên tiền phí mua bảo hiểm nhân thọ, bảo hiểm không "
            "bắt buộc khác của doanh nghiệp bảo hiểm không thành lập tại Việt Nam cho người lao động", "32", "VNĐ", False),
]
TNCN_COUNT_CODES = {"16", "17", "18", "19", "20"}
_MONTH_RE = re.compile(r"^(?:bang luong.*?)?thang\s*(\d{1,2})\s*(?:[/.\-]|nam)\s*(\d{4})")


def _tncn_code(code):
    return "[" + code.replace("_", ".") + "]"


def parse_submitted_tncn(data):
    """XML tờ khai 05/KK-TNCN đã nộp (HTKK, maTKhai 864) → {mã chỉ tiêu: số} + thông tin NNT; None nếu không phải."""
    root = ET.fromstring(data)
    ma = _desc(root, "maTKhai")
    if _local(root.tag) != "HSoThueDTu" or ma is None or (ma.text or "").strip() != "864":
        return None
    ky = _desc(root, "kyKKhai")
    out = {"_ky": (ky.text or "").strip() if ky is not None else ""}
    ct = _desc(root, "CTieuTKhaiChinh")
    for e in (ct.iter() if ct is not None else []):
        tag = _local(e.tag)
        if re.fullmatch(r"ct\d+(_\d+)?", tag) and e.text and re.fullmatch(r"-?\d+(\.\d+)?", e.text.strip()):
            out[tag[2:]] = float(e.text)
    nnt = _desc(root, "NNT")
    out["_nnt"] = {k: _text(nnt, k) for k in ("mst", "tenNNT", "dchiNNT", "tenXaNNT", "tenHuyenNNT", "tenTinhNNT",
                                               "dthoaiNNT", "faxNNT", "emailNNT")} if nnt is not None else {}
    nk = _desc(root, "nguoiKy")
    out["_nguoi_ky"] = (nk.text or "").strip() if nk is not None else ""
    return out


def find_payroll_files(data_dir, patterns):
    """File bảng lương .xlsx ở thư mục gốc (không phân biệt hoa thường), cũ → mới."""
    out = []
    for fn in os.listdir(data_dir):
        low = fn.lower()
        if low.startswith("~$") or not low.endswith(".xlsx") or not os.path.isfile(os.path.join(data_dir, fn)):
            continue
        if any(fnmatch.fnmatch(low, p.lower()) for p in patterns):
            out.append(os.path.join(data_dir, fn))
    return sorted(out, key=os.path.getmtime)


def _person_key(name):
    return " ".join(unicodedata.normalize("NFC", str(name or "")).split()).casefold()


def _is_stt(v):
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return v > 0 and float(v).is_integer()
    return isinstance(v, str) and v.strip().isdigit()


def _payroll_columns(head, sub):
    """head / sub: {cột: tiêu đề đã chuẩn hóa} của dòng tiêu đề và dòng tiêu đề phụ ngay dưới → {khóa: cột}."""
    def first(*prefixes, without=None):
        return next((c for c, h in sorted(head.items()) if h.startswith(prefixes) and not (without and without in h)), None)

    cols = {"stt": first("stt"), "ten": first("ho va ten", "ho ten"), "chuc_vu": first("chuc vu", "chuc danh"),
            "o": first("tong luong thuc te", "tong thu nhap"), "y": first("thuong", without="doanh so"),
            "z": first("giam tru gia canh"), "aa": first("thu nhap tinh thue"), "ab": first("thue tncn", "thue thu nhap ca nhan"),
            "ac": first("tam ung"), "ad": first("thuc linh", "thuc nhan"), "x": []}
    # BH bắt buộc phần NLĐ: nhóm tiêu đề "Bảo hiểm NLĐ" → cột "Cộng" ở dòng phụ (không có thì cộng các cột con)
    c0 = next((c for c, h in sorted(head.items()) if h.startswith("bao hiem") and ("nld" in h or "nguoi lao dong" in h)), None)
    if c0:
        end = min([c for c in head if c > c0] or [c0 + 1])
        span = range(c0, max(end, c0 + 1))
        total = next((c for c in span if sub.get(c) == "cong"), None)
        cols["x"] = [total] if total else ([c for c in span if sub.get(c)] or [c0])
    # các khoản cộng thành 'Tổng lương thực tế' — dùng khi ô không có công thức, để biết Thưởng đã nằm trong tổng chưa
    start = max([c for c, h in head.items() if h.startswith(("ngay lam viec", "ngay cong", "luong dong bhxh", "luong co ban"))] or [0])
    cols["thanh_phan"] = [c for c in sorted(head) if cols["o"] and start < c < cols["o"]]
    cols["thanh_phan_2"] = [c for c, h in sorted(head.items()) if h.startswith(("luong thuc nhan", "luong lam them"))]
    return cols


def load_payroll(path):
    """Bảng lương Excel: mỗi tháng 1 khối 'Tháng MM/YYYY' + dòng tiêu đề 'Họ và tên'
    → (dòng NV-tháng, cảnh báo, {name, mst, address} đọc ở đầu sheet đầu tiên)."""
    wb_v = load_workbook(path, data_only=True)
    wb_f = load_workbook(path, data_only=False)
    fname = os.path.basename(path)
    rows, warns, header = [], [], {}
    for ws in wb_v.worksheets:
        wf = wb_f[ws.title]
        grid = list(ws.iter_rows(values_only=True))
        month, cols = None, None
        for r_idx, vals in enumerate(grid, start=1):
            texts = {c: v for c, v in enumerate(vals, start=1) if isinstance(v, str) and v.strip()}
            if ws is wb_v.worksheets[0] and month is None and cols is None:
                for c, v in sorted(texts.items()):
                    h = norm_text(v)
                    if h.startswith("mst") and "mst" not in header:
                        digits = re.sub(r"\D", "", v) or re.sub(r"\D", "", str(next((x for x in vals[c:] if x not in (None, "")), "")))
                        header["mst"] = digits
                    elif "cong ty" in h and "name" not in header and len(v) < 150:
                        header["name"] = " ".join(v.split())
                    elif h.startswith(("so ", "dia chi")) and "address" not in header and "name" in header:
                        header["address"] = re.sub(r"^\s*Địa chỉ\s*:\s*", "", " ".join(v.split()), flags=re.I)
            for c, v in texts.items():
                m = _MONTH_RE.match(strip_accents(v).lower().strip()) if c <= 6 and len(v) <= 60 else None
                if m and 1 <= int(m.group(1)) <= 12:
                    month, cols = (int(m.group(2)), int(m.group(1))), None
                    break
            head = {c: norm_text(v) for c, v in texts.items()}
            if any(h.startswith(("ho va ten", "ho ten")) for h in head.values()):
                nxt = grid[r_idx] if r_idx < len(grid) else ()
                sub = {c: norm_text(v) for c, v in enumerate(nxt, start=1) if isinstance(v, str) and v.strip()}
                cols = _payroll_columns(head, sub)
                where = f"{fname} › {ws.title} dòng {r_idx}" + (f" (tháng {month[1]:02d}/{month[0]})" if month else "")
                if not cols["o"] or not cols["ten"]:
                    warns.append(f"{where}: không thấy cột 'Họ và tên' / 'Tổng lương thực tế' — bỏ qua khối này")
                    cols = None
                elif not cols["x"]:
                    warns.append(f"{where}: không thấy cột 'Bảo hiểm NLĐ' — tính BH NLĐ = 0")
                if month is None and cols:
                    warns.append(f"{where}: không thấy dòng 'Tháng MM/YYYY' phía trên — bỏ qua khối này")
                continue
            if month is None or cols is None:
                continue
            name = vals[cols["ten"] - 1]
            if not (isinstance(name, str) and name.strip()) or (cols["stt"] and not _is_stt(vals[cols["stt"] - 1])):
                continue

            def g(c):
                return parse_num(vals[c - 1]) if c and c <= len(vals) else 0.0
            o, y = g(cols["o"]), g(cols["y"])
            y_in_o = None
            if y:
                f = wf.cell(row=r_idx, column=cols["o"]).value
                if isinstance(f, str) and f.startswith("="):
                    y_in_o = bool(re.search(rf"(?<![A-Z$])\$?{_col_letter(cols['y'])}\$?\d", f.upper()))
                else:
                    for parts in (cols["thanh_phan"], cols["thanh_phan_2"]):
                        s = sum(g(c) for c in parts)
                        if parts and abs(o - s - y) < 1:
                            y_in_o = True
                            break
                        if parts and abs(o - s) < 1:
                            y_in_o = False
                            break
            rows.append({"thang": month, "ten": " ".join(name.split()), "chuc_vu": str(vals[cols["chuc_vu"] - 1] or "").strip()
                         if cols["chuc_vu"] else "", "o": o, "y": y, "y_in_o": y_in_o, "x": sum(g(c) for c in cols["x"]),
                         "z": g(cols["z"]), "aa": g(cols["aa"]), "ab": g(cols["ab"]), "ac": g(cols["ac"]), "ad": g(cols["ad"]),
                         "nguon": f"{ws.title}!{r_idx}"})
    if rows and all(r["o"] == 0 for r in rows):
        warns.append(f"{fname}: mọi ô 'Tổng lương thực tế' đều trống — mở file bằng Excel, bấm Lưu (để Excel ghi số đã tính) rồi chạy lại")
    return rows, warns, header


def _months_between(d_from, d_to):
    out, y, m = [], d_from.year, d_from.month
    while (y, m) <= (d_to.year, d_to.month):
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def compute_tncn(rows, months, cfg):
    """Chỉ tiêu 05/KK-TNCN từ bảng lương các tháng trong kỳ → (tk, danh sách người, dòng NV-tháng, thông tin)."""
    tru_bh = cfg.get("tncn_tru_bh_nld", True)
    by_last = cfg.get("tncn_dem_so_nguoi", "ca_quy") == "thang_cuoi"
    no_hd = {_person_key(n) for n in (cfg.get("tncn_khong_hdld") or [])}
    q_rows = sorted((r for r in rows if r["thang"] in months), key=lambda r: (r["thang"], r["nguon"]))
    for r in q_rows:
        # Thưởng chỉ cộng thêm khi chắc chắn chưa nằm trong 'Tổng lương thực tế'
        r["thuong_cong"] = r["y"] if r["y"] and r["y_in_o"] is False else 0.0
        r["tnct"] = round_half_up(r["o"] + r["thuong_cong"] - (r["x"] if tru_bh else 0))
        r["co_luong"] = r["o"] + r["thuong_cong"] > 0
    paid = [r for r in q_rows if r["co_luong"]]
    last = max((r["thang"] for r in paid), default=None)
    persons = {}
    for r in paid:
        p = persons.setdefault(_person_key(r["ten"]), {"ten": r["ten"], "chuc_vu": r["chuc_vu"], "thang": [], "o": 0.0,
                                                       "thuong_cong": 0.0, "x": 0.0, "tnct": 0, "aa": 0.0, "ab": 0.0})
        if r["thang"] not in p["thang"]:
            p["thang"].append(r["thang"])
        for k in ("o", "thuong_cong", "x", "tnct", "aa", "ab"):
            p[k] += r[k]
    for k, p in persons.items():
        p["thang_cuoi"] = last in p["thang"]
        p["dem_16"] = p["thang_cuoi"] if by_last else True
        p["hdld"] = k not in no_hd
    counted = [p for p in persons.values() if p["dem_16"]]
    tk = {code: 0 for _, _, code, _, _ in TNCN_ROWS}
    tk["16"] = len(counted)
    tk["17"] = sum(1 for p in counted if p["hdld"])
    tk["18"] = tk["19"] = sum(1 for p in persons.values() if p["ab"] > 0)
    tk["21"] = tk["22"] = sum(r["tnct"] for r in paid)
    tk["26"] = tk["27"] = sum(r["tnct"] for r in paid if r["ab"] > 0)
    tk["29"] = tk["30"] = round_half_up(sum(r["ab"] for r in q_rows))
    seen = Counter((r["thang"], _person_key(r["ten"])) for r in q_rows)
    info = {"n_quy": len(persons), "n_cuoi": sum(1 for p in persons.values() if p["thang_cuoi"]), "thang_cuoi": last,
            "by_last": by_last, "tru_bh": tru_bh, "thang": months,
            "thieu_thang": [m for m in months if m not in {r["thang"] for r in q_rows}],
            "y_khong_ro": [r for r in q_rows if r["y"] and r["y_in_o"] is None],
            "trung": [k for k, n in seen.items() if n > 1]}
    return tk, sorted(persons.values(), key=lambda p: (not p["dem_16"], -len(p["thang"]))), q_rows, info


def _tncn_how(code, info):
    m = info["thang_cuoi"]
    last = f"{m[1]:02d}/{m[0]}" if m else "—"
    return {
        "16": (f"Số người có lương tháng cuối kỳ ({last})" if info["by_last"] else
               f"Số người có lương trong kỳ, mỗi người đếm 1 lần (riêng tháng {last}: {info['n_cuoi']} người)"),
        "17": "= [16] trừ NV ghi trong CONFIG tncn_khong_hdld (không ký HĐLĐ từ 3 tháng trở lên)",
        "18": "Số người có 'Thuế TNCN' > 0 trên bảng lương",
        "19": "= [18]",
        "21": "Σ các tháng: Tổng lương thực tế + Thưởng (nếu chưa cộng vào tổng)" + (" − BH bắt buộc phần NLĐ" if info["tru_bh"] else ""),
        "22": "= [21] (toàn bộ là cá nhân cư trú)",
        "26": "Thu nhập chịu thuế của các tháng có 'Thuế TNCN' > 0",
        "27": "= [26]",
        "29": "Σ cột 'Thuế TNCN' trên bảng lương",
        "30": "= [29]",
    }.get(code, "")


def write_tncn_xlsx(path, ky, company, tk, persons, q_rows, info, signer, src_name, compare=None, compare_file=""):
    wb = Workbook()
    ws = wb.active
    ws.title = "To khai 05-KK-TNCN"
    widths = [6, 70, 9, 8, 18, 62]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[_col_letter(i)].width = w
    _xl_title(ws, 1, "TỜ KHAI THUẾ THU NHẬP CÁ NHÂN — Mẫu số 05/KK-TNCN (TT 80/2021/TT-BTC)", 6, 14)
    _xl_title(ws, 2, "BẢN NHÁP ĐỐI CHIẾU — số liệu tính từ bảng lương; nộp chính thức qua HTKK / eTax", 6, 10, bold=False, italic=True)
    ws.cell(row=2, column=1).font = Font(italic=True, color="C00000")
    head = [f"[01] Kỳ tính thuế: {ky}", f"[04] Tên người nộp thuế: {company.get('name', '')}",
            f"[05] Mã số thuế: {company.get('mst', '')}", f"[06] Địa chỉ: {company.get('address', '')}",
            f"[09] Điện thoại: {company.get('phone', '')}    [11] Email: {company.get('email', '')}", "Đơn vị tiền: Đồng Việt Nam"]
    for i, line in enumerate(head, start=4):
        ws.cell(row=i, column=1, value=line)
    _xl_header(ws, 11, ["STT", "Chỉ tiêu", "Mã chỉ tiêu", "Đơn vị tính", "Số người / Số tiền", "Cách tính (script)"], widths)
    row = 13
    for stt, text, code, unit, bold in TNCN_ROWS:
        ws.cell(row=row, column=1, value=stt).font = Font(bold=bold)
        ws.cell(row=row, column=2, value=text).font = Font(bold=bold)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        ws.cell(row=row, column=3, value=_tncn_code(code)).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=4, value=unit).alignment = Alignment(horizontal="center")
        c = ws.cell(row=row, column=5, value=tk.get(code, 0))
        c.number_format = VN_NUM
        c.font = Font(bold=bold)
        ws.cell(row=row, column=6, value=_tncn_how(code, info)).font = Font(italic=True, color="595959")
        ws.cell(row=row, column=6).alignment = Alignment(wrap_text=True)
        row += 1
    months_txt = ", ".join(f"{m:02d}/{y}" for y, m in info["thang"])
    notes = ["Ghi chú tính số (bản nháp):",
             f"• Nguồn: {src_name} — các tháng {months_txt}. Chi tiết: sheet 'Theo nguoi', 'Theo thang'.",
             "• Thưởng chỉ cộng thêm khi công thức ô 'Tổng lương thực tế' chưa gồm cột Thưởng (cột 'Thưởng đã trong tổng?' ở sheet 'Theo thang').",
             "• [21] đang trừ BH bắt buộc phần NLĐ (giống tờ khai đã nộp). Theo luật, BH bắt buộc là khoản giảm trừ khi tính "
             "thu nhập TÍNH thuế; nếu kế toán muốn [21] là thu nhập chịu thuế trước giảm trừ → CONFIG tncn_tru_bh_nld = False."
             if info["tru_bh"] else "• [21] chưa trừ BH bắt buộc phần NLĐ (CONFIG tncn_tru_bh_nld = False).",
             "• NV không ký HĐLĐ từ 3 tháng trở lên (vd lương khoán theo ngày): ghi họ tên vào CONFIG tncn_khong_hdld để không "
             "tính vào [17]; khoản trả cho người này từ 2.000.000 đ/lần trở lên phải khấu trừ 10% (trừ khi có cam kết 08/CK-TNCN)."]
    if info["thieu_thang"]:
        notes.append("• THIẾU bảng lương tháng: " + ", ".join(f"{m:02d}/{y}" for y, m in info["thieu_thang"]) + " — số liệu chưa đủ kỳ.")
    if info["y_khong_ro"]:
        notes.append(f"• {len(info['y_khong_ro'])} dòng có Thưởng nhưng không rõ đã nằm trong 'Tổng lương thực tế' chưa → "
                     "tạm KHÔNG cộng thêm; xem cột 'Thưởng đã trong tổng?' ở sheet 'Theo thang'.")
    if info["trung"]:
        notes.append(f"• {len(info['trung'])} người xuất hiện 2 lần trong cùng 1 tháng — kiểm lại bảng lương (sheet 'Theo thang').")
    row += 1
    for line in notes:
        ws.cell(row=row, column=2, value=line).font = Font(italic=True, bold=line.endswith(":"))
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 30 if len(line) > 110 else 16
        row += 1
    _xl_signature(ws, row + 1, 6, signer)

    def table(sheet, headers, widths_, data, money_cols, total_cols=()):
        for i, (h, w) in enumerate(zip(headers, widths_), start=1):
            c = sheet.cell(row=1, column=i, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.column_dimensions[_col_letter(i)].width = w
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "A2"
        for r_i, vals in enumerate(data, start=2):
            for c_i, v in enumerate(vals, start=1):
                cell = sheet.cell(row=r_i, column=c_i, value=v)
                if c_i in money_cols:
                    cell.number_format = VN_NUM
        if total_cols and data:
            tr = len(data) + 2
            sheet.cell(row=tr, column=2, value="Tổng cộng").font = Font(bold=True)
            for c_i in total_cols:
                col = _col_letter(c_i)
                cell = sheet.cell(row=tr, column=c_i, value=f"=SUM({col}2:{col}{tr - 1})")
                cell.font = Font(bold=True)
                cell.number_format = VN_NUM
        sheet.auto_filter.ref = f"A1:{_col_letter(len(headers))}{len(data) + 1}"

    yn = {True: "Có", False: "Không"}
    w2 = wb.create_sheet("Theo nguoi")
    table(w2, ["STT", "Họ và tên", "Chức vụ", "Các tháng có lương", "Có lương tháng cuối kỳ", "Tính vào [16]", "Có HĐLĐ ([17])",
               "Tổng lương thực tế", "Thưởng cộng thêm", "BH bắt buộc NLĐ", "Thu nhập chịu thuế", "Thu nhập tính thuế",
               "Thuế TNCN đã khấu trừ"],
          [6, 28, 26, 22, 12, 10, 10, 16, 14, 14, 16, 14, 14],
          [(i, p["ten"], p["chuc_vu"], ", ".join(f"{m:02d}/{y}" for y, m in sorted(p["thang"])), yn[p["thang_cuoi"]],
            yn[p["dem_16"]], yn[p["hdld"]], p["o"], p["thuong_cong"], p["x"], p["tnct"], p["aa"], p["ab"])
           for i, p in enumerate(persons, start=1)],
          {8, 9, 10, 11, 12, 13}, (8, 9, 10, 11, 12, 13))
    w3 = wb.create_sheet("Theo thang")
    y_label = {True: "Có", False: "Chưa → cộng thêm", None: "Không rõ → chưa cộng"}
    table(w3, ["Tháng", "Họ và tên", "Chức vụ", "Có lương", "Tổng lương thực tế", "Thưởng", "Thưởng đã trong tổng?",
               "Thưởng cộng thêm", "BH bắt buộc NLĐ", "Thu nhập chịu thuế", "Giảm trừ gia cảnh", "Thu nhập tính thuế",
               "Thuế TNCN", "Tạm ứng", "Thực lĩnh", "Ô trong bảng lương (sheet!dòng)"],
          [9, 28, 26, 8, 16, 12, 18, 14, 14, 16, 14, 14, 12, 12, 14, 22],
          [(f"{r['thang'][1]:02d}/{r['thang'][0]}", r["ten"], r["chuc_vu"], yn[r["co_luong"]], r["o"], r["y"],
            y_label[r["y_in_o"]] if r["y"] else "", r["thuong_cong"], r["x"], r["tnct"] if r["co_luong"] else 0, r["z"],
            r["aa"], r["ab"], r["ac"], r["ad"], r["nguon"]) for r in q_rows],
          {5, 6, 8, 9, 10, 11, 12, 13, 14, 15}, (5, 6, 8, 9, 10, 13))

    if compare:
        w4 = wb.create_sheet("So sanh to khai da nop")
        for i, w in enumerate([12, 18, 18, 18, 80], start=1):
            w4.column_dimensions[_col_letter(i)].width = w
        _xl_title(w4, 1, f"SO SÁNH VỚI TỜ KHAI 05/KK-TNCN ĐÃ NỘP — {ky}", 5, 13)
        _xl_title(w4, 2, f"File: {compare_file}", 5, 9, bold=False, italic=True, align="left")
        w4.append([])
        w4.append(["Chỉ tiêu", "Đã nộp (XML)", "Script tính", "Chênh (script − đã nộp)", "Kết quả"])
        for c in w4[4]:
            c.font = Font(bold=True)
        for label, a, b, diff, verdict in compare:
            w4.append([label, a, b, diff, verdict])
            r = w4.max_row
            for ci in (2, 3, 4):
                w4.cell(row=r, column=ci).number_format = "#,##0;(#,##0);0"
            if verdict.startswith("LỆCH"):
                w4.cell(row=r, column=5).font = Font(bold=True, color="C00000")
    wb.save(path)


def write_tncn_pdf(path, ky, company, tk, signer):
    _ensure_reportlab()
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    if not _pdf_fonts():
        return False
    st = ParagraphStyle("n", fontName="VN", fontSize=9.5, leading=12)
    st_b = ParagraphStyle("b", parent=st, fontName="VN-B")
    st_c = ParagraphStyle("c", parent=st, alignment=TA_CENTER)
    st_t = ParagraphStyle("t", parent=st_c, fontName="VN-B", fontSize=13, leading=16)
    st_i = ParagraphStyle("i", parent=st, fontName="VN-I")
    st_ic = ParagraphStyle("ic", parent=st_c, fontName="VN-I")
    st_r = ParagraphStyle("r", parent=st, alignment=TA_RIGHT)
    st_rb = ParagraphStyle("rb", parent=st_r, fontName="VN-B")
    P = Paragraph
    story = [Table([[P("<b>CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM</b><br/><b>Độc lập - Tự do - Hạnh phúc</b>", st_c),
                     P("<b>Mẫu số: 05/KK-TNCN</b><br/>(Ban hành kèm theo Thông tư số 80/2021/TT-BTC ngày 29 tháng 9 năm "
                       "2021 của Bộ trưởng Bộ Tài chính)", st_c)]], colWidths=[110 * mm, 70 * mm]),
             Spacer(1, 4 * mm), P("TỜ KHAI THUẾ THU NHẬP CÁ NHÂN", st_t),
             P("(Áp dụng đối với tổ chức, cá nhân trả các khoản thu nhập từ tiền lương, tiền công)", st_ic), Spacer(1, 3 * mm),
             P(f"[01] Kỳ tính thuế: <b>{ky}</b>", st), P("[02] Lần đầu: X &nbsp;&nbsp;&nbsp; [03] Bổ sung lần thứ:", st),
             P(f"[04] Tên người nộp thuế: <b>{company.get('name', '')}</b>", st),
             P(f"[05] Mã số thuế: <b>{company.get('mst', '')}</b>", st), P(f"[06] Địa chỉ: {company.get('address', '')}", st),
             P(f"[07] Quận/Huyện: {company.get('ward', '')} &nbsp;&nbsp;&nbsp; [08] Tỉnh/Thành phố: {company.get('province', '')}", st),
             P(f"[09] Điện thoại: {company.get('phone', '')} &nbsp;&nbsp;&nbsp; [10] Fax: {company.get('fax', '')} "
               f"&nbsp;&nbsp;&nbsp; [11] Email: {company.get('email', '')}", st),
             P("[12] Tên đại lý thuế (nếu có):", st), P("[13] Mã số thuế:", st), P("[14] Hợp đồng đại lý thuế: Số: &nbsp;&nbsp;&nbsp; Ngày:", st),
             P("[15] Phân bổ thuế do có đơn vị hạch toán phụ thuộc tại địa bàn cấp tỉnh khác nơi có trụ sở chính:", st),
             P("Đơn vị tiền: Đồng Việt Nam", st_r), Spacer(1, 2 * mm)]
    data = [[P("<b>STT</b>", st_c), P("<b>Chỉ tiêu</b>", st_c), P("<b>Mã chỉ tiêu</b>", st_c), P("<b>Đơn vị tính</b>", st_c),
             P("<b>Số người/ Số tiền</b>", st_c)]]
    for stt, text, code, unit, bold in TNCN_ROWS:
        v = tk.get(code, 0)
        shown = (f"{v:,}".replace(",", ".") if v else "") if code in TNCN_COUNT_CODES else _fmt_vn(v)
        data.append([P(stt, st_c), P(text, st_b if bold else st), P(_tncn_code(code), st_c), P(unit, st_c),
                     P(shown, st_rb if bold else st_r)])
    t = Table(data, colWidths=[11 * mm, 104 * mm, 17 * mm, 17 * mm, 31 * mm], repeatRows=1)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.black), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                           ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
    today = datetime.now()
    sign = Table([[P("NHÂN VIÊN ĐẠI LÝ THUẾ<br/>Họ và tên:<br/>Chứng chỉ hành nghề số:", st),
                   P(f"<i>Ngày {today:%d} tháng {today:%m} năm {today:%Y}</i><br/><b>NGƯỜI NỘP THUẾ hoặc<br/>ĐẠI DIỆN HỢP PHÁP CỦA "
                     f"NGƯỜI NỘP THUẾ</b><br/><i>(Chữ ký, ghi rõ họ tên; chức vụ và đóng dấu (nếu có)/Ký điện tử)</i>"
                     f"<br/><br/><br/><b>{signer}</b>", st_c)]], colWidths=[80 * mm, 100 * mm])
    story += [t, Spacer(1, 4 * mm),
              P("Tôi cam đoan số liệu khai trên là đúng và chịu trách nhiệm trước pháp luật về những số liệu đã khai./.", st_i),
              Spacer(1, 3 * mm), sign]
    on_page = _pdf_draft_page("BẢN NHÁP ĐỐI CHIẾU — tính từ bảng lương, KHÔNG DÙNG ĐỂ NỘP (nộp chính thức qua HTKK / eTax)")
    SimpleDocTemplate(path, pagesize=A4, leftMargin=15 * mm, rightMargin=12 * mm, topMargin=16 * mm, bottomMargin=14 * mm,
                      title=f"To khai TNCN {ky} - BAN NHAP", author="doi_soat.py").build(story, onFirstPage=on_page, onLaterPages=on_page)
    return True


def tncn_figures(paths, cfg, period, company, submitted_list):
    """Số liệu tờ khai 05/KK-TNCN từ các file bảng lương (paths, cũ → mới), KHÔNG ghi file; so với XML tờ khai đã nộp nếu
    có. period = (từ ngày, đến ngày, tên kỳ, tên ngắn) của tờ khai GTGT, hoặc None. Bảng lương chưa có tháng nào của kỳ đó
    (và kỳ không đặt cứng trong CONFIG) → lấy quý gần nhất có lương. Không có số liệu → None.
    Dùng chung cho CLI và KeToanBackEnd."""
    loaded = []
    for p in reversed(paths):  # mới → cũ
        try:
            rs, warns, header = load_payroll(p)
        except Exception as e:
            log(f"  ! Không đọc được bảng lương {os.path.basename(p)}: {type(e).__name__}: {e}")
            continue
        for w in warns:
            log(f"  ! {w}")
        if rs:
            loaded.append((p, rs, header))
    if not loaded:
        log(f"  ! Không đọc được dòng lương nào trong {', '.join(os.path.basename(p) for p in paths)} → không xuất tờ khai TNCN "
            "(cần dòng 'Tháng MM/YYYY' và dòng tiêu đề có 'Họ và tên', 'Tổng lương thực tế')")
        return None
    chosen = None
    if period:
        months = _months_between(period[0], period[1])
        chosen = next((x for x in loaded if any(r["thang"] in months for r in x[1])), None)
        if chosen is None and cfg.get("tu_ngay") and cfg.get("den_ngay"):
            log(f"  ! Bảng lương không có tháng nào thuộc {period[2]} (kỳ đặt trong CONFIG) → không xuất tờ khai TNCN")
            return None
    if chosen is None:
        chosen = loaded[0]
        y, m = max(r["thang"] for r in chosen[1])
        q = (m - 1) // 3 + 1
        if period:
            log(f"  ! Bảng lương chưa có tháng nào của {period[2]} (kỳ của hóa đơn) → tờ khai TNCN lấy quý gần nhất có lương: "
                f"Quý {q} năm {y}")
        period = (*_quarter_bounds(y, q), f"Quý {q} năm {y}", f"Q{q}-{y}")
    d_from, d_to, ky, short = period
    months = _months_between(d_from, d_to)
    path, rows, header = chosen
    log(f"  Đang xuất tờ khai TNCN {ky} từ {os.path.basename(path)}...")
    tk, persons, q_rows, info = compute_tncn(rows, months, cfg)
    if info["thieu_thang"]:
        log("  ! Bảng lương THIẾU tháng " + ", ".join(f"{m:02d}/{y}" for y, m in info["thieu_thang"]))
    q_now = ((d_from.month - 1) // 3 + 1, d_from.year)
    same = next((s for s in submitted_list if _ky_of(s) == q_now), None)
    ref = same or (submitted_list[-1] if submitted_list else None)
    nnt = (ref or {}).get("_nnt", {})
    company = dict(company, name=nnt.get("tenNNT") or company.get("name") or header.get("name", ""),
                   mst=nnt.get("mst") or company.get("mst") or header.get("mst", ""),
                   address=nnt.get("dchiNNT") or company.get("address") or header.get("address", ""),
                   ward=nnt.get("tenHuyenNNT") or nnt.get("tenXaNNT", ""),
                   province=nnt.get("tenTinhNNT", ""), phone=nnt.get("dthoaiNNT", ""), fax=nnt.get("faxNNT", ""),
                   email=nnt.get("emailNNT", ""))
    signer = cfg.get("nguoi_ky", "") or (ref or {}).get("_nguoi_ky", "")
    compare = None
    if same is not None:
        compare = []
        for _, _, code, _, _ in TNCN_ROWS:
            a, b = same.get(code), tk.get(code, 0)
            if a is None and not b:
                continue
            diff = b - (a or 0)
            if abs(diff) < 0.5:
                verdict = "Khớp"
            elif code in TNCN_COUNT_CODES:
                verdict = "LỆCH — khác số người"
                if code in ("16", "17"):
                    verdict += (f" (cả kỳ {info['n_quy']} người, riêng tháng cuối kỳ {info['n_cuoi']} người — "
                                "đổi cách đếm ở CONFIG tncn_dem_so_nguoi / tncn_khong_hdld)")
            else:
                verdict = "Lệch làm tròn" if abs(diff) <= 5 else "LỆCH — xem lại"
            compare.append((_tncn_code(code), a, b, diff, verdict))
        n_bad = sum(1 for r in compare if r[4].startswith("LỆCH"))
        log(f"  So với tờ khai TNCN đã nộp ({same['_file']}): {len(compare) - n_bad}/{len(compare)} chỉ tiêu khớp"
            + (f", {n_bad} LỆCH — xem sheet 'So sanh to khai da nop'" if n_bad else ""))
    log(f"  Tờ khai TNCN {ky}: [16] {tk['16']} người · [17] {tk['17']} · [21] {tk['21']:,.0f} · [26] {tk['26']:,.0f} · "
        f"[29] {tk['29']:,.0f}")
    return {"period": period, "tk": tk, "persons": persons, "q_rows": q_rows, "info": info, "company": company,
            "signer": signer, "ref": ref, "same": same, "compare": compare, "src_name": os.path.basename(path)}


def export_tncn(data_dir, cfg, period, company, submitted_list):
    """Tờ khai 05/KK-TNCN (xlsx + pdf bản nháp + XML) từ bảng lương ở thư mục gốc — số liệu từ tncn_figures()."""
    paths = find_payroll_files(data_dir, cfg.get("file_bang_luong") or ["*luong*.xlsx"])
    if not paths:
        log("  (Không thấy bảng lương *luong*.xlsx ở thư mục gốc → không xuất tờ khai TNCN 05/KK-TNCN)")
        return []
    t = tncn_figures(paths, cfg, period, company, submitted_list)
    if t is None:
        return []
    d_from, d_to, ky, short = t["period"]
    tk, company, signer, ref, same = t["tk"], t["company"], t["signer"], t["ref"], t["same"]
    files = []
    xlsx = os.path.join(data_dir, f"TK_TNCN_{short}.xlsx")
    write_tncn_xlsx(xlsx, ky, company, tk, t["persons"], t["q_rows"], t["info"], signer, t["src_name"], t["compare"],
                    same["_file"] if same else "")
    files.append(xlsx)
    if cfg.get("xuat_pdf_ban_nhap", True):
        pdf = os.path.join(data_dir, f"TK_TNCN_{short}_BAN_NHAP.pdf")
        if write_tncn_pdf(pdf, ky, company, tk, signer):
            files.append(pdf)
    if cfg.get("xuat_xml_to_khai", True):
        x = write_tk_xml(data_dir, "864", ref, d_from, d_to, dict(tk, **{"15": 0}), signer, company)
        if x:
            files.append(x)
    return files


# ============================================================ XML TỜ KHAI (định dạng HTKK)

TK_XML_NS = "http://kekhaithue.gdt.gov.vn/TKhaiThue"
_TK_XML_HEAD = ('<HSoThueDTu xmlns="' + TK_XML_NS + '"><HSoKhaiThue id="ID_1"><TTinChung><TTinDVu><maDVu>HTKK</maDVu>'
                '<tenDVu>HỖ TRỢ KÊ KHAI THUẾ</tenDVu><pbanDVu>5.8.1</pbanDVu><ttinNhaCCapDVu/></TTinDVu><TTinTKhaiThue><TKhaiThue>'
                '<maTKhai>{ma}</maTKhai><tenTKhai>{ten}</tenTKhai><moTaBMau>(Ban hành kèm theo Thông tư số 80/2021/TT-BTC ngày 29 '
                'tháng 9 năm 2021 của Bộ trưởng Bộ Tài chính)</moTaBMau><pbanTKhaiXML>{pban}</pbanTKhaiXML><loaiTKhai>C</loaiTKhai>'
                '<soLan>0</soLan><KyKKhaiThue><kieuKy>Q</kieuKy><kyKKhai/><kyKKhaiTuNgay/><kyKKhaiDenNgay/><kyKKhaiTuThang/>'
                '<kyKKhaiDenThang/></KyKKhaiThue><maCQTNoiNop/><tenCQTNoiNop/><ngayLapTKhai/><GiaHan><maLyDoGiaHan/><lyDoGiaHan/>'
                '</GiaHan><nguoiKy/><ngayKy/><nganhNgheKD/></TKhaiThue><NNT>{nnt}</NNT></TTinTKhaiThue></TTinChung>')
# Khung XML theo đúng thứ tự thẻ của HTKK 5.8.1 (01/GTGT pban 2.8.3, 05/KK-TNCN pban 2.9.3) — dùng khi thư mục không có
# XML tờ khai đã nộp làm mẫu; có mẫu thì phần TTinChung (CQT, ngành nghề, NNT...) lấy từ mẫu.
TK_XML_SKELETON = {
    "842": _TK_XML_HEAD.format(
        ma="842", ten="TỜ KHAI THUẾ GIÁ TRỊ GIA TĂNG (Mẫu số 01/GTGT)", pban="2.8.3",
        nnt="".join(f"<{t}/>" for t in ("mst", "tenNNT", "dchiNNT", "phuongXa", "maHuyenNNT", "tenHuyenNNT", "maTinhNNT",
                                         "tenTinhNNT", "dthoaiNNT", "faxNNT", "emailNNT")))
    + "<CTieuTKhaiChinh><ma_NganhNghe>00</ma_NganhNghe><ten_NganhNghe>Hoạt động sản xuất kinh doanh thông thường"
      "</ten_NganhNghe><tieuMucHachToan>1701</tieuMucHachToan><Header><ct09/><ct10/><DiaChiHDSXKDKhacTinhNDTSC>"
      "<ct11a_phuongXa_ma/><ct11a_phuongXa_ten/><ct11b_quanHuyen_ma/><ct11b_quanHuyen_ten/><ct11c_tinhTP_ma/><ct11c_tinhTP_ten/>"
      "</DiaChiHDSXKDKhacTinhNDTSC></Header><ct21>0</ct21><ct22>0</ct22><GiaTriVaThueGTGTHHDVMuaVao><ct23>0</ct23><ct24>0</ct24>"
      "</GiaTriVaThueGTGTHHDVMuaVao><HangHoaDichVuNhapKhau><ct23a>0</ct23a><ct24a>0</ct24a></HangHoaDichVuNhapKhau><ct25>0</ct25>"
      "<ct26>0</ct26><HHDVBRaChiuThueGTGT><ct27>0</ct27><ct28>0</ct28></HHDVBRaChiuThueGTGT><ct29>0</ct29><HHDVBRaChiuTSuat5>"
      "<ct30>0</ct30><ct31>0</ct31></HHDVBRaChiuTSuat5><HHDVBRaChiuTSuat10><ct32>0</ct32><ct33>0</ct33></HHDVBRaChiuTSuat10>"
      "<ct32a>0</ct32a><TongDThuVaThueGTGTHHDVBRa><ct34>0</ct34><ct35>0</ct35></TongDThuVaThueGTGTHHDVBRa><ct36>0</ct36>"
      "<ct37>0</ct37><ct38>0</ct38><ct39a>0</ct39a><ct40a>0</ct40a><ct40b>0</ct40b><ct40>0</ct40><ct41>0</ct41><ct42>0</ct42>"
      "<ct43>0</ct43></CTieuTKhaiChinh><PLuc><PL_NQ142_GTGT><HH_DV_MuaVaoTrongKy>"
      "<tongCongGiaTriHHDVMuaVao>0</tongCongGiaTriHHDVMuaVao><tongCongThueGTGTHHDV>0</tongCongThueGTGTHHDV>"
      "</HH_DV_MuaVaoTrongKy><HH_DV_BanRaTrongKy><tongCongGiaTriHHDV>0</tongCongGiaTriHHDV>"
      "<tongCongThueGTGTDuocGiam>0</tongCongThueGTGTDuocGiam></HH_DV_BanRaTrongKy><ChenhLech><ct9>0</ct9></ChenhLech>"
      "</PL_NQ142_GTGT></PLuc></HSoKhaiThue></HSoThueDTu>",
    "864": _TK_XML_HEAD.format(
        ma="864", ten="TK khấu trừ thuế thu nhập cá nhân Mẫu 05/KK-TNCN (TT80/2021)", pban="2.9.3",
        nnt="".join(f"<{t}/>" for t in ("mst", "tenNNT", "dchiNNT", "tenXaNNT", "maXaNNT", "maHuyenNNT", "tenHuyenNNT",
                                         "maTinhNNT", "tenTinhNNT", "dthoaiNNT", "faxNNT", "emailNNT")))
    + "<CTieuTKhaiChinh><mst_cu/>"
    + "".join(f"<ct{c}>0</ct{c}>" for c in ("15", "16", "17", "18", "19", "20", "21", "22", "23", "24", "25", "25_1",
                                            "26", "27", "28", "29", "30", "31", "32"))
    + "</CTieuTKhaiChinh></HSoKhaiThue></HSoThueDTu>",
}
_XML_BAD_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xset(el, tag, value):
    e = _desc(el, tag)
    if e is not None:
        e.text = None if value in (None, "") else _XML_BAD_CHARS.sub("", str(value))


def write_tk_xml(data_dir, ma, ref, d_from, d_to, values, signer, company, pl=None):
    """XML tờ khai 01/GTGT (ma 842, kèm PL 142 NQ204) hoặc 05/KK-TNCN (ma 864) theo định dạng HTKK.
    ref = tờ khai cùng loại đã nộp (dict có '_raw') để lấy CQT / ngành nghề / NNT; None = dùng khung + thông tin công ty."""
    q = (d_from.month - 1) // 3 + 1
    if (d_from, d_to) != _quarter_bounds(d_from.year, q):
        log(f"  ! Kỳ {d_from:%d/%m/%Y}–{d_to:%d/%m/%Y} không phải 1 quý → không xuất XML tờ khai {ma}")
        return None
    ET.register_namespace("", TK_XML_NS)
    root = ET.fromstring(TK_XML_SKELETON[ma])
    hs = _child(root, "HSoKhaiThue")
    ct = _child(hs, "CTieuTKhaiChinh")
    ref_root = None
    if ref and ref.get("_raw"):
        try:
            ref_root = ET.fromstring(ref["_raw"])
        except ET.ParseError:
            ref_root = None
    ref_tc = _desc(ref_root, "TTinChung") if ref_root is not None else None
    if ref_tc is not None:
        # khung thẻ giữ theo HTKK; chỉ lấy GIÁ TRỊ (CQT, ngành nghề, NNT, người ký...) từ tờ khai mẫu theo tên thẻ
        skel_ver = _text(_desc(root, "TKhaiThue"), "pbanTKhaiXML")
        ref_ver = _text(_desc(ref_tc, "TKhaiThue"), "pbanTKhaiXML")
        if ref_ver and ref_ver != skel_ver:
            log(f"  ! XML mẫu {ref.get('_file', '')} là phiên bản {ref_ver}, script xuất theo phiên bản {skel_ver} — "
                "kiểm lại khi nhập vào HTKK/eTax")
        ct_ref = _desc(ref_root, "CTieuTKhaiChinh")
        for part, src_root in ((_child(hs, "TTinChung"), ref_tc), (_child(ct, "Header"), _desc(ct_ref, "Header"))):
            for e in (part.iter() if part is not None and src_root is not None else []):
                src = _desc(src_root, _local(e.tag)) if not len(e) else None
                if src is not None and _local(e.tag) != "pbanTKhaiXML":
                    e.text = src.text
        for tag in ("ma_NganhNghe", "ten_NganhNghe", "tieuMucHachToan", "mst_cu"):
            src, dst = _child(ct_ref, tag), _child(ct, tag)
            if src is not None and dst is not None:
                dst.text = src.text
    else:
        log(f"  ! Không có XML tờ khai {'01/GTGT' if ma == '842' else '05/KK-TNCN'} đã nộp làm mẫu → XML thiếu mã CQT nơi nộp, "
            "ngành nghề, điện thoại... (chép 1 XML tờ khai cùng loại kỳ trước vào thư mục rồi chạy lại)")
        for tag, key in (("mst", "mst"), ("tenNNT", "name"), ("dchiNNT", "address"), ("mst_cu", "mst")):
            _xset(root, tag, company.get(key, ""))
    today = date.today().isoformat()
    for tag, v in (("ttinNhaCCapDVu", ""), ("loaiTKhai", "C"), ("soLan", "0"), ("kieuKy", "Q"), ("kyKKhai", f"{q}/{d_from.year}"),
                   ("kyKKhaiTuNgay", f"{d_from:%d/%m/%Y}"), ("kyKKhaiDenNgay", f"{d_to:%d/%m/%Y}"), ("kyKKhaiTuThang", ""),
                   ("kyKKhaiDenThang", ""), ("ngayLapTKhai", today), ("ngayKy", today)):
        _xset(root, tag, v)
    if signer:
        _xset(root, "nguoiKy", signer)
    for code, v in values.items():
        _xset(ct, f"ct{code}", round_half_up(v or 0))

    if pl is not None:
        pl1, pl2, pl3 = pl
        pluc = _child(hs, "PLuc")
        if not pl1 and not pl2:
            hs.remove(pluc)
        else:
            def add_rows(parent, rows, fields):
                for i, vals in enumerate(rows, start=1):
                    row = ET.Element(f"{{{TK_XML_NS}}}BangKeTenHHDV", {"ID": f"ID_{i}"})
                    for f, v in zip(fields, vals):
                        ET.SubElement(row, f"{{{TK_XML_NS}}}{f}").text = _XML_BAD_CHARS.sub("", str(v))
                    parent.insert(i - 1, row)
            rows1 = [(n, round_half_up(v), t) for n, v, t in pl1]
            rows2 = [(n, round_half_up(v), 10, 8, t) for n, v, t in pl2]
            mv, br = _desc(pluc, "HH_DV_MuaVaoTrongKy"), _desc(pluc, "HH_DV_BanRaTrongKy")
            add_rows(mv, rows1, ("tenHHDVMuaVao", "giaTriHHDVMuaVao", "thueGTGTHHDV"))
            add_rows(br, rows2, ("tenHHDV", "giaTriHHDV", "thueSuatTheoQuyDinh", "thueSuatSauGiam", "thueGTGTDuocGiam"))
            _xset(mv, "tongCongGiaTriHHDVMuaVao", sum(r[1] for r in rows1))
            _xset(mv, "tongCongThueGTGTHHDV", pl3["06"])
            _xset(br, "tongCongGiaTriHHDV", sum(r[1] for r in rows2))
            _xset(br, "tongCongThueGTGTDuocGiam", pl3["08"])
            _xset(pluc, "ct9", pl3["09"])

    if hasattr(ET, "indent"):
        ET.indent(root, space="  ")
    body = ET.tostring(root, encoding="unicode").replace(
        f'<HSoThueDTu xmlns="{TK_XML_NS}">',
        f'<HSoThueDTu xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns="{TK_XML_NS}">', 1)
    mst = re.sub(r"\D", "", _text(_desc(root, "NNT"), "mst") or company.get("mst", ""))
    mst13 = mst + "000" if len(mst) == 10 else mst
    name = {"842": "01_GTGT", "864": "05_KK_TNCN"}[ma]
    out_dir = os.path.join(data_dir, XML_OUT_DIR)
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{mst13}-{name}_TT80-Q{q}{d_from.year}-L00.xml")
    with open(path, "w", encoding="utf-8-sig", newline="\r\n") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n' + body)
    return path


if __name__ == "__main__":
    _args = parse_args()
    code = 1
    try:
        code = main(_args)
    except KeyboardInterrupt:
        log("Đã dừng.")
        code = 130
    except Exception:
        import traceback
        traceback.print_exc()
        log("\n!! Lỗi không mong muốn — chụp màn hình này gửi cho người hỗ trợ.")
    if not _args.khong_dung and os.name == "nt":
        input("\nNhấn Enter để đóng cửa sổ...")
    sys.exit(code)
