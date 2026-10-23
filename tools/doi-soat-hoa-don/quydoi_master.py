"""Danh mục Vietda (QUYDOI.xlsx) → chỉ giữ ĐVT nhỏ nhất (ĐVT master KiotViet). Ghi FILE MỚI, không sửa file gốc.
Mỗi sản phẩm (= 1 mã master KiotViet, ghép theo tên): có mã ở ĐVT master → giữ mã đó, xóa mã ở ĐVT lớn hơn;
chưa có → giữ mã có hệ số nhỏ nhất, đổi ĐVT thành ĐVT master. ĐVT không nhận ra → không đụng, ghi 'Can xem'.
Dùng: python quydoi_master.py <doi_soat.py> <QUYDOI.xlsx> <bang_quy_doi_kiotviet.json> <file ra .xlsx>"""
import importlib.util
import json
import sys
from collections import Counter, defaultdict

sys.stdout.reconfigure(encoding="utf-8")
spec = importlib.util.spec_from_file_location("ds", sys.argv[1])
ds = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ds)
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

SRC, CAT, OUT = sys.argv[2], sys.argv[3], sys.argv[4]
cat = ds.Catalog(CAT)
by_name = defaultdict(set)   # tên KiotViet (name, fullName — chuẩn hóa) → mã master
import re  # noqa: E402

TAG_RE = re.compile(r"^\s*\[[^\]]*\]\s*")


def strip_tag(s):
    """Bỏ nhãn thương hiệu '[CP] ', '[SCA Vinamilk] ' ở đầu tên (KiotViet hay có, tên Vietda theo HĐ NCC thường không)."""
    return TAG_RE.sub("", s or "")


by_name_notag = defaultdict(set)   # như by_name nhưng tên đã bỏ nhãn [..] ở đầu
kv_group = {}                      # mã master → (tên chung chuẩn hóa, tên chung) của nhóm hàng cùng loại KiotViet
group_info = {}                    # khóa nhóm "NHOM:<tên>" → {"code", "name", "unit"} (dùng như cat.masters[mk])
for p in json.load(open(CAT, encoding="utf-8"))["products"]:
    if p.get("isDeleted") or p.get("isClone") or not p.get("isActive", True):
        continue
    mk = ds.norm_code(p.get("masterCode") or p.get("code"))
    if mk in cat.masters:
        if p.get("name"):   # tên chung của nhóm "hàng hóa cùng loại" (các thuộc tính cùng 'name')
            kv_group.setdefault(mk, (ds.norm_text(p["name"]), p["name"]))
        full, unit = p.get("fullName") or "", p.get("unit") or ""
        # fullName = tên + " - thuộc tính" + " (đơn vị)" → thêm bản bỏ "(đơn vị)" để khớp tên Vietda có ghi thuộc tính
        no_unit = full[:-len(unit) - 2].strip() if unit and full.endswith(f"({unit})") else ""
        for nm in (p.get("name"), full, no_unit):
            if nm:
                by_name[ds.norm_text(nm)].add(mk)
                by_name_notag[ds.norm_text(strip_tag(nm))].add(mk)

wb = load_workbook(SRC)
ws = wb.worksheets[0]
hdr_row = col = None
for r in range(1, 30):
    vals = {c: ds.norm_text(ws.cell(r, c).value) for c in range(1, ws.max_column + 1) if isinstance(ws.cell(r, c).value, str)}
    if any(v.startswith("ma hang") for v in vals.values()) and any(v.startswith("ten hang") for v in vals.values()):
        hdr_row = r
        col = {"stt": next((c for c, v in vals.items() if v == "stt"), None),
               "ma": next(c for c, v in vals.items() if v.startswith("ma hang")),
               "ten": next(c for c, v in vals.items() if v.startswith("ten hang")),
               "dvt": next(c for c, v in vals.items() if v in ("dvt", "don vi tinh"))}
        break
assert hdr_row, "không thấy dòng tiêu đề"


def unit_conv(g, dvt):
    """Hệ số của ĐVT Vietda theo các ĐVT KiotViet của mặt hàng; None = không nhận ra (không đoán)."""
    mu = g["unit"]
    if not dvt or ds.norm_unit(dvt) == ds.norm_unit(mu):
        return 1.0, ""
    conv, how, label, cands = ds.match_unit(g, dvt, None)
    if conv is None:
        return None, ""
    note = f"'{dvt}' khớp nhiều ĐVT KiotViet ({', '.join(f'{c[0]}={c[1]:g}' for c in cands)}) → '{label}'" \
        if how == "keyword" and len(cands) > 1 else ""
    return conv, note


groups = defaultdict(list)   # mã master → [dòng Vietda đã nhận ra ĐVT]
review, n_rows = [], 0
for r in range(hdr_row + 1, ws.max_row + 1):
    ma, ten, dvt = ws.cell(r, col["ma"]).value, ws.cell(r, col["ten"]).value, ws.cell(r, col["dvt"]).value
    if not (isinstance(ma, str) and ma.strip() and isinstance(ten, str) and ten.strip()):
        continue
    if len(ma.strip()) == 1 and len(ten.strip()) == 1:
        continue
    n_rows += 1
    dvt = str(dvt or "").strip()
    masters, how = by_name.get(ds.norm_text(ten), set()), "Tên Vietda = tên KiotViet"
    if not masters:
        masters, how = by_name_notag.get(ds.norm_text(strip_tag(ten)), set()), "Tên giống (bỏ nhãn [..] ở đầu)"
    if not masters:
        masters = {mk for mk, g in cat.masters.items() if ds.base_name(ds.norm_text(ten)) in g["names"]}
        how = "Tên gốc (bỏ quy cách) = tên KiotViet"
    if not masters:
        ranked = cat.by_name(ds.base_name(ds.norm_text(ten)))
        review.append((r, ma, ten, dvt, "", "", "Không tìm thấy tên trên KiotViet",
                       " | ".join(cat.masters[m]["name"] for s, m in ranked[:2])))
        continue
    if len(masters) > 1:
        names = sorted(cat.masters[m]["name"] for m in masters)
        ms = sorted(masters)
        same_group = len({kv_group.get(m, (m,))[0] for m in ms}) == 1
        same_unit = len({ds.norm_unit(cat.masters[m]["unit"]) for m in ms}) == 1
        convs = [unit_conv(cat.masters[m], dvt) for m in ms]
        same_conv = all(c is not None for c, _ in convs) and len({c for c, _ in convs}) == 1
        if same_group and same_unit and same_conv:
            # KiotViet gộp các thuộc tính dưới 1 tên chung, cùng ĐVT nhỏ nhất + cùng hệ số → coi cả nhóm là 1 sản phẩm
            key = "NHOM:" + kv_group[ms[0]][0]
            group_info.setdefault(key, {"code": ", ".join(cat.masters[m]["code"] for m in ms),
                                        "name": f"{kv_group[ms[0]][1]} (nhóm {len(ms)} thuộc tính)",
                                        "unit": cat.masters[ms[0]]["unit"]})
            groups[key].append({"r": r, "ma": ma.strip(), "ten": ten, "dvt": dvt, "conv": convs[0][0],
                                "how": "Tên chung nhóm hàng cùng loại KiotViet", "note": convs[0][1]})
            continue
        why = ("khác ĐVT nhỏ nhất / hệ số giữa các thuộc tính" if same_group else "khác nhóm hàng")
        review.append((r, ma, ten, dvt, "", "", f"Tên chung của {len(masters)} sản phẩm KiotViet ({why}) — "
                       "chọn đúng sản phẩm / tách mã", " | ".join(names)))
        continue
    mk = next(iter(masters))
    g = cat.masters[mk]
    conv, note = unit_conv(g, dvt)
    if conv is None:
        review.append((r, ma, ten, dvt, g["code"], g["unit"],
                       f"ĐVT '{dvt}' không có trên KiotViet (có: {ds._units_text(g)}) — không đoán hệ số", ""))
        continue
    groups[mk].append({"r": r, "ma": ma.strip(), "ten": ten, "dvt": dvt, "conv": conv, "how": how, "note": note})

changed, deleted, kept_dup = [], [], []
to_delete = []
for mk, rows in groups.items():
    g = group_info.get(mk) or cat.masters[mk]
    rows.sort(key=lambda x: (x["conv"], x["r"]))
    masters_rows = [x for x in rows if x["conv"] == 1]
    keep = masters_rows[0] if masters_rows else rows[0]
    if len(masters_rows) > 1:   # nhiều mã cùng ở ĐVT master → không tự chọn, giữ tất cả, báo
        kept_dup.append((g["code"], g["name"], g["unit"], ", ".join(f"{x['ma']} ({x['dvt']})" for x in masters_rows)))
    if keep["conv"] != 1:   # cùng đơn vị (hệ số 1) chỉ khác cách viết 'Gói'/'gói' → giữ chữ của Vietda
        notes = [x for x in (keep["note"], f"TỒN KHO mã {keep['ma']} trong Vietda × {keep['conv']:g} "
                              f"(1 {keep['dvt']} = {keep['conv']:g} {g['unit']})" if keep["conv"] != 1 else "") if x]
        ws.cell(keep["r"], col["dvt"]).value = g["unit"]
        ws.cell(keep["r"], col["dvt"]).fill = PatternFill("solid", fgColor="FFF2CC")
        changed.append((keep["r"], keep["ma"], keep["ten"], keep["dvt"], g["unit"], keep["conv"], g["code"], g["name"],
                        keep["how"], "; ".join(notes)))
    for x in rows:
        if x is keep or (x["conv"] == 1 and len(masters_rows) > 1):
            continue
        ratio = x["conv"] / keep["conv"] if keep["conv"] else x["conv"]
        deleted.append((x["r"], x["ma"], x["ten"], x["dvt"], keep["ma"], g["unit"], x["conv"], g["code"],
                        f"Chuyển tồn kho {x['ma']} sang {keep['ma']}: SL × {x['conv']:g} (1 {x['dvt']} = {x['conv']:g} {g['unit']})"
                        + (f"; {x['note']}" if x["note"] else "")))
        to_delete.append(x["r"])

for r in sorted(to_delete, reverse=True):
    ws.delete_rows(r)
if col["stt"]:   # đánh lại STT
    n = 0
    for r in range(hdr_row + 1, ws.max_row + 1):
        v = ws.cell(r, col["stt"]).value
        if isinstance(v, (int, float)) and ws.cell(r, col["ma"]).value:
            n += 1
            ws.cell(r, col["stt"]).value = n


def sheet(title, headers, rows, widths):
    w = wb.create_sheet(title)
    w.append(headers)
    for c in w[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        w.append(list(row))
    for i, wd in enumerate(widths, start=1):
        w.column_dimensions[ds._col_letter(i)].width = wd
    w.freeze_panes = "A2"
    w.auto_filter.ref = w.dimensions


sheet("Da xoa", ["Dòng (file gốc)", "Mã đã xóa", "Tên hàng", "ĐVT", "Mã giữ lại", "ĐVT giữ lại", "Hệ số",
                 "Mã master KiotViet", "Việc cần làm trong Vietda"], deleted, [8, 10, 45, 10, 10, 10, 7, 16, 90])
sheet("Doi DVT", ["Dòng (file gốc)", "Mã hàng", "Tên hàng", "ĐVT cũ", "ĐVT mới (master)", "Hệ số (1 ĐVT cũ = ? ĐVT mới)",
                  "Mã master KiotViet", "Tên KiotViet", "Cách ghép", "Ghi chú"], changed, [8, 10, 45, 10, 12, 12, 16, 45, 24, 90])
sheet("Trung ma DVT master", ["Mã master KiotViet", "Tên KiotViet", "ĐVT master", "Các mã Vietda cùng ở ĐVT master (đều giữ — kế toán chọn)"],
      kept_dup, [18, 45, 10, 60])
sheet("Can xem", ["Dòng (file gốc)", "Mã hàng", "Tên hàng", "ĐVT", "Mã master KiotViet", "ĐVT master", "Lý do (giữ nguyên)",
                  "Gợi ý gần nhất"], review, [8, 10, 45, 10, 16, 10, 70, 80])
wb.save(OUT)
print(f"Mã Vietda: {n_rows:,} | ghép được sản phẩm KiotViet: {sum(len(v) for v in groups.values()):,} mã / {len(groups):,} sản phẩm")
print(f"  XÓA (đơn vị lớn, đã có mã đơn vị nhỏ nhất): {len(deleted)} | ĐỔI ĐVT về master: {len(changed)} "
      f"(trong đó hệ số ≠ 1: {sum(1 for x in changed if x[5] != 1)}) | còn lại sau xử lý: {n_rows - len(deleted):,}")
print(f"  nhiều mã cùng ĐVT master (giữ hết, báo): {len(kept_dup)} sản phẩm | cần xem (không đụng): {len(review)} — "
      + ", ".join(f"{k}: {v}" for k, v in Counter(x[6].split(" (")[0].split(" — ")[0][:45] for x in review).most_common(4)))
for x in [d for d in deleted if "larue" in d[2].lower()][:3]:
    print("  vd xóa:", x[1], x[2][:40], x[3], "→ giữ", x[4], x[5], "|", x[8][:90])
for x in changed[:4]:
    print("  vd đổi:", x[1], x[2][:40], f"{x[3]} → {x[4]} ×{x[5]:g}", "|", x[9][:90])
for x in deleted[:4]:
    print("  vd xóa:", x[1], x[2][:40], x[3], "→ giữ", x[4], x[5])
